from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from templyfier.smart import build_smart_toplines, inspect_smart_package, proposal_to_row


def _multistage_export() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    stages = (
        ("NEAT", "N", "Q-01-NEAT-Overall fragrance liking", 6.2, 6.6),
        ("WET", "W", "Q-02-WET-Overall fragrance liking", 5.8, 6.1),
    )
    for index, (stage, suffix, question, benchmark, candidate) in enumerate(stages, 1):
        sheet = workbook.create_sheet(f"Table_{index} 2_TAILED")
        sheet["A1"] = f"Stage: {stage}"
        sheet["A2"] = "Search: TOTAL"
        sheet["A5"], sheet["B5"] = "Variable", "Metric"
        sheet["C3"], sheet["D3"] = "Benchmark", "Candidate"
        sheet["C4"], sheet["D4"] = f"100 - B1{suffix}", f"100 - C1{suffix}"
        sheet["C5"], sheet["D5"] = "A", "B"
        metrics = (
            ("Mean", benchmark, candidate),
            ("Top Box", 0.20, 0.30),
            ("Top 2 Boxes", 0.50, 0.65),
            ("Bottom 2 Boxes", 0.12, 0.08),
        )
        for row, (metric, left, right) in enumerate(metrics, 6):
            sheet.cell(row, 1, question)
            sheet.cell(row, 2, metric)
            sheet.cell(row, 3, left)
            sheet.cell(row, 4, right)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class MultiStageTests(unittest.TestCase):
    def test_all_stage_questions_are_detected_and_written(self):
        source = _multistage_export()
        files = [("DataViz_CLT_NEAT_WET.xlsx", source)]

        info = inspect_smart_package(files)
        self.assertEqual(set(info.stages), {"NEAT", "WET"})
        self.assertEqual(
            {question.question_id for question in info.questions},
            {
                "Q-01-NEAT-Overall fragrance liking",
                "Q-02-WET-Overall fragrance liking",
            },
        )
        self.assertEqual(
            {question.question_id: question.stages for question in info.questions},
            {
                "Q-01-NEAT-Overall fragrance liking": ("NEAT",),
                "Q-02-WET-Overall fragrance liking": ("WET",),
            },
        )

        rows = [proposal_to_row(question) for question in info.questions]
        output, report = build_smart_toplines(
            files,
            rows,
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
        )
        workbook = load_workbook(BytesIO(output), data_only=False)
        self.assertEqual(len(workbook.sheetnames), 2)
        neat = next(workbook[name] for name in workbook.sheetnames if name.startswith("NEAT"))
        wet = next(workbook[name] for name in workbook.sheetnames if name.startswith("WET"))
        neat_values = {cell.value for row in neat.iter_rows() for cell in row if cell.value}
        wet_values = {cell.value for row in wet.iter_rows() for cell in row if cell.value}
        self.assertTrue(any(str(value).endswith("Overall fragrance liking") for value in neat_values))
        self.assertTrue(any(str(value).endswith("Overall fragrance liking") for value in wet_values))
        self.assertEqual(report["data_rows_written"], 8)

    def test_unknown_source_stage_is_preserved_for_manual_review(self):
        workbook = load_workbook(BytesIO(_multistage_export()))
        workbook.remove(workbook["Table_2 2_TAILED"])
        sheet = workbook["Table_1 2_TAILED"]
        sheet["A1"] = "Stage: PRE-WASH"
        for row in range(6, 10):
            sheet.cell(row, 1, "Q-01-Overall fragrance liking")
        buffer = BytesIO()
        workbook.save(buffer)

        info = inspect_smart_package([("DataViz_CLT_PRE_WASH.xlsx", buffer.getvalue())])
        self.assertEqual(info.stages, ("PRE-WASH",))
        self.assertEqual(info.questions[0].stages, ("PRE-WASH",))
        self.assertEqual(proposal_to_row(info.questions[0])["Stage"], "PRE-WASH")

    def test_unassigned_stage_never_drops_the_question(self):
        workbook = load_workbook(BytesIO(_multistage_export()))
        workbook.remove(workbook["Table_2 2_TAILED"])
        sheet = workbook["Table_1 2_TAILED"]
        sheet["A1"] = None
        for row in range(6, 10):
            sheet.cell(row, 1, "Q-01-Overall fragrance liking")
        buffer = BytesIO()
        workbook.save(buffer)
        files = [("DataViz_CLT_UNKNOWN_STAGE.xlsx", buffer.getvalue())]

        info = inspect_smart_package(files)
        self.assertEqual(len(info.questions), 1)
        question_row = proposal_to_row(info.questions[0])
        self.assertEqual(question_row["Stage"], "Unassigned")
        output, report = build_smart_toplines(
            files,
            [question_row],
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
        )
        self.assertTrue(output)
        self.assertEqual(report["data_rows_written"], 4)


if __name__ == "__main__":
    unittest.main()
