from io import BytesIO
from copy import copy
import unittest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from tests.test_editor import source_bytes, prepared, AGG
from templyfier.editor_model import export_rows
from templyfier.smart import build_smart_toplines, _split_detection, _friendly_split_label, _stage_name, _study_question_priority
from templyfier.core import TemplyfierError


def fixture():
    wb=load_workbook(BytesIO(source_bytes()))
    ws=wb.active
    for c in range(5,8):
        ws.cell(3,c,f'Product {c-2}')
        ws.cell(4,c,f'{40+c} - P{c-2}')
        ws.cell(5,c,chr(62+c))
        for r in range(8,ws.max_row+1):
            ws.cell(r,c,c+.01*r if ws.cell(r,2).value=='Mean' else .1*c)
            ws.cell(r,c).fill=PatternFill('solid',fgColor='FF0000' if c%2 else '00FF00')
    b=BytesIO();wb.save(b);return b.getvalue()


class BenchmarkColumnsTests(unittest.TestCase):
    def test_wet_only_clt_starts_with_ofo_then_strength_then_descriptors(self):
        rows = [
            (_study_question_priority('CLT','Q-02-OFL','Overall Fragrance Opinion','FRAGRANCE EVALUATION',20,('WET',)),'OFO'),
            (_study_question_priority('CLT','Q-03-Strength','Strength','FRAGRANCE STRENGTH',30,('WET',)),'Strength'),
            (_study_question_priority('CLT','Q-05-Attributes','Olfactive Attributes','FRAGRANCE CHARACTERISTICS',50,('WET',)),'Descriptors'),
        ]
        self.assertEqual([label for _,label in sorted(rows)],['OFO','Strength','Descriptors'])

    def test_split_filter_can_move_from_a2_and_use_semicolon_separators(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']=None
        ws['C3']='Filters: S-1-REGION: North; Search: S-2-USAGE: Heavy users'
        self.assertEqual(_split_detection(ws,'export.xlsx'),('NORTH · HEAVY USERS','High confidence'))

    def test_french_filter_can_move_deeper_in_header(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']=None
        ws['J12']='Filtres: S-1-REGION: North; Recherche: S-2-USAGE: Heavy users'
        self.assertEqual(_split_detection(ws,'export.xlsx'),('NORTH · HEAVY USERS','High confidence'))

    def test_separate_filter_cells_are_combined_without_losing_a_dimension(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']='Search: S-1-REGION: North'
        ws['B2']='Filter: S-2-USAGE: Heavy users'
        self.assertEqual(_split_detection(ws,'export.xlsx'),('NORTH · HEAVY USERS','High confidence'))

    def test_stage_can_move_and_use_a_french_label(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A1']='Converted export'
        ws['F4']='Stade : DRY'
        self.assertEqual(_stage_name(ws),'DRY')

    def test_unrecognised_nonempty_filter_is_never_silently_total(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']='Search: malformed filter without value'
        self.assertEqual(_split_detection(ws,'split_unknown.xlsx'),('split_unknown','Review recommended'))

    def test_explicit_total_is_stable_across_benchmark_filenames(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']='Search: TOTAL'
        for name in ('bench-a.xlsx','bench-b.xlsx','bench-c.xlsx'):
            self.assertEqual(_split_detection(ws,name),('TOTAL','High confidence'))

    def test_native_total_without_search_prefix_is_explicit(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']='TOTAL'
        self.assertEqual(_split_detection(ws,'timestamped-export.xlsx'),('TOTAL','High confidence'))

    def test_mo_brand_buckets_do_not_collapse_distinct_splits(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        cases={
            'S-14-MO brands: SURF':'SURF MO',
            'S-14-MO brands: BOLD, SURF':'BOLD MO',
            "S-14-MO brands: ALDI ALMAT, ASDA, LIDL FORMIL, MARKS & SPENCER (M&S), MORRISONS, SAINSBURY'S, TESCO, WAITROSE & PARTNERS":'RETAILER BRANDS MO',
            'S-14-MO brands: ARIEL, BIO-D, DAZ, DETTOL, ECOVER, FAIRY, KINFILL, METHOD, MINIML, PERSIL, SMOL, TALLOW + ASH, WILTON LONDON, WOOLITE, Other brand: Please specify':'NATIONAL BRANDS MO',
        }
        for raw,expected in cases.items():
            ws['A2']='Search: '+raw
            self.assertEqual(_friendly_split_label(ws),expected)

    def test_exact_total_filename_is_not_flagged_when_filter_is_absent(self):
        wb=load_workbook(BytesIO(fixture()));ws=wb.active
        ws['A2']=None
        self.assertEqual(_split_detection(ws,'TOTAL.xlsx'),('TOTAL','High confidence'))

    def build(self,mode,gaps,benchmarks=(0,1,2),splits=('TOTAL',)):
        raw=fixture();_,rows=prepared(raw)
        data,report=build_smart_toplines([(f'{i}.xlsx',raw) for i in range(len(splits))],export_rows(rows),
            split_names=list(splits),benchmark_positions=benchmarks,standard_metrics=AGG,
            include_screeners=False,include_sections=False,benchmark_sheet_mode=mode,
            show_monadic_gaps=gaps)
        return load_workbook(BytesIO(data)),report

    def check_panels(self,gaps,benchmarks=(0,1,2)):
        combined,report=self.build('benchmark_columns',gaps,benchmarks)
        separate,reference=self.build('separate',gaps,benchmarks)
        self.assertEqual(len(combined.sheetnames),1)
        self.assertEqual(len(separate.sheetnames),len(benchmarks))
        self.assertEqual(len(report['benchmark_column_panels']),len(benchmarks))
        self.assertFalse(report['missing_panel_metrics'])
        ws=combined.active
        self.assertEqual(ws['A5'].value,'Variable')
        self.assertEqual(ws['B5'].value,'Metric')
        for panel,source in zip(report['benchmark_column_panels'],separate.worksheets):
            start=panel['first_column']
            for row in range(6,source.max_row+1):
                self.assertEqual(ws.cell(row,1).value,source.cell(row,1).value)
                self.assertEqual(ws.cell(row,2).value,source.cell(row,2).value)
                for col in range(3,source.max_column+1):
                    actual=ws.cell(row,start+col-3); expected=source.cell(row,col)
                    if expected.data_type!='f':self.assertEqual(actual.value,expected.value)
                    else:
                        self.assertEqual(actual.data_type,'f')
                        self.assertNotIn('#REF!',actual.value)
                        from openpyxl.formula.translate import Translator
                        self.assertEqual(actual.value,Translator(expected.value,origin=expected.coordinate).translate_formula(actual.coordinate))
                    self.assertEqual(copy(actual.fill),copy(expected.fill))
                    self.assertEqual(actual.number_format,expected.number_format)
        return combined

    def test_three_benchmarks_with_gaps_preserve_values_styles_and_formula_references(self):
        self.check_panels(True)

    def test_three_benchmarks_without_gaps_keep_each_comparison_independent(self):
        wb=self.check_panels(False)
        self.assertFalse(any(c.data_type=='f' for row in wb.active for c in row))

    def test_four_benchmarks_supported(self):
        self.check_panels(True,(0,1,2,3))

    def test_two_splits_produce_two_combined_or_six_separate_sheets(self):
        combined,_=self.build('benchmark_columns',False,splits=('TOTAL','GROUP A'))
        separate,_=self.build('separate',False,splits=('TOTAL','GROUP A'))
        self.assertEqual(len(combined.sheetnames),2)
        self.assertEqual(len(separate.sheetnames),6)

    def test_auto_exports_keep_different_source_scores_and_bases(self):
        inputs=[]
        for i,code in enumerate(('B1','C1','P3')):
            wb=load_workbook(BytesIO(fixture()));ws=wb.active
            ws.cell(5,8,code)
            ws.cell(4,7,f'{60+i} - P5')
            for r in range(8,ws.max_row+1):ws.cell(r,7,10+i)
            b=BytesIO();wb.save(b);inputs.append((f'{i}.xlsx',b.getvalue()))
        _,rows=prepared(inputs[0][1])
        options=dict(split_names=['TOTAL']*3,benchmark_positions=[0,1,2],standard_metrics=AGG,
            include_screeners=False,include_sections=False,show_monadic_gaps=False)
        data,report=build_smart_toplines(inputs,export_rows(rows),benchmark_sheet_mode='auto_columns',**options)
        ref,_=build_smart_toplines(inputs,export_rows(rows),benchmark_sheet_mode='auto_exports',**options)
        wb=load_workbook(BytesIO(data));separate=load_workbook(BytesIO(ref))
        self.assertEqual(len(wb.sheetnames),1)
        self.assertEqual(len(report['benchmark_column_panels']),3)
        for panel,source in zip(report['benchmark_column_panels'],separate.worksheets):
            for row in (4,6,7):
                for col in range(3,source.max_column+1):
                    self.assertEqual(wb.active.cell(row,panel['first_column']+col-3).value,source.cell(row,col).value)

    def test_auto_exports_pair_every_split_with_three_benchmarks(self):
        inputs=[]
        for split in ('TOTAL','GROUP A'):
            for code in ('B1','C1','P3'):
                wb=load_workbook(BytesIO(fixture()));ws=wb.active
                ws.cell(5,8,code)
                payload=BytesIO();wb.save(payload)
                inputs.append((f'{split}-{code}.xlsx',payload.getvalue()))
        _,rows=prepared(inputs[0][1])
        options=dict(split_names=['TOTAL']*3+['GROUP A']*3,
            benchmark_positions=[0,1,2],standard_metrics=AGG,
            include_screeners=False,include_sections=False,show_monadic_gaps=False)
        combined,report=build_smart_toplines(inputs,export_rows(rows),
            benchmark_sheet_mode='auto_columns',**options)
        separate,separate_report=build_smart_toplines(inputs,export_rows(rows),
            benchmark_sheet_mode='auto_exports',**options)
        self.assertEqual(load_workbook(BytesIO(combined)).sheetnames,['TOTAL','GROUP A'])
        self.assertEqual(len(report['benchmark_column_panels']),6)
        self.assertEqual(len(load_workbook(BytesIO(separate)).sheetnames),6)
        self.assertEqual(separate_report['missing_benchmark_exports'],[])

    def test_missing_metric_is_aligned_as_blank_not_zero(self):
        wb=load_workbook(BytesIO(fixture()));wb.active.cell(5,8,'B1')
        first=BytesIO();wb.save(first);raw=first.getvalue()
        wb=load_workbook(BytesIO(raw));ws=wb.active
        ws.cell(5,8,'C1')
        ws.delete_rows(next(c.row for c in ws['B'] if c.value=='Mean'))
        b=BytesIO();wb.save(b)
        _,rows=prepared(raw)
        data,report=build_smart_toplines([('first.xlsx',raw),('second.xlsx',b.getvalue())],export_rows(rows),
            split_names=['TOTAL','TOTAL'],benchmark_positions=[0,1],standard_metrics=AGG,
            include_screeners=False,include_sections=False,benchmark_sheet_mode='auto_columns')
        self.assertTrue(report['missing_panel_metrics'])
        wb=load_workbook(BytesIO(data))
        comments=[cell for row in wb.active for cell in row if cell.comment]
        self.assertTrue(comments)
        self.assertTrue(all(cell.value is None for cell in comments))

    def test_kpi_readings_unchanged_by_layout_and_optional_gap_applies_to_details(self):
        raw=fixture();_,rows=prepared(raw)
        for row in rows:
            row['KPI Summary']=row['Question ID']=='Q-1-Overall liking'
            if row['KPI Summary']:row['Summary label']='Autres'
        options=dict(split_names=['TOTAL'],benchmark_positions=[0,1,2],standard_metrics=AGG,
            include_screeners=False,summary_scope='total')
        combined,report=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),
            benchmark_sheet_mode='benchmark_columns',show_monadic_gaps=False,**options)
        separate,ref=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),
            benchmark_sheet_mode='separate',show_monadic_gaps=False,**options)
        a=load_workbook(BytesIO(combined));b=load_workbook(BytesIO(separate))
        for name in report['summary_sheets']:
            self.assertEqual(list(a[name].values),list(b[name].values))
        details=a['KPI Details'];headers=[c.value for c in details[1]]
        self.assertIn('Product score',headers)
        self.assertIn('Favourable direction',headers)
        self.assertNotIn('Gap',headers)
        self.assertEqual(details.cell(2,headers.index('KPI')+1).value,'Autres')
        gap,_=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),
            benchmark_sheet_mode='benchmark_columns',show_monadic_gaps=True,**options)
        self.assertIn('Gap',[c.value for c in load_workbook(BytesIO(gap))['KPI Details'][1]])

    def test_duplicate_manual_split_names_cannot_create_ambiguous_panels(self):
        with self.assertRaisesRegex(TemplyfierError,'Duplicate readings'):
            self.build('benchmark_columns',False,splits=('TOTAL','TOTAL'))

    def test_paired_delta_option_removes_score_difference_columns(self):
        raw=source_bytes(paired=True);_,rows=prepared(raw)
        for enabled in (False,True):
            data,report=build_smart_toplines([('paired.xlsx',raw)],export_rows(rows),split_names=['TOTAL'],
                benchmark_positions=[0],standard_metrics=AGG,test_type='Paired',include_deltas=enabled,
                include_screeners=False)
            wb=load_workbook(BytesIO(data))
            formulas=[c for s in wb for row in s for c in row if c.data_type=='f']
            self.assertEqual(bool(formulas),enabled)
