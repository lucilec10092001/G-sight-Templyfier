"""Controlled bilingual fixtures, not a substitute for real French-export acceptance."""
from io import BytesIO
import unittest
from openpyxl import load_workbook
from templyfier.language import analysis_normal
from templyfier.smart import _classify, _metric_key, _normal, default_metric_selection, build_smart_toplines
from templyfier.editor_model import export_rows, set_metric_selection
from .test_editor import source_bytes, prepared, AGG

FR_AGG=['Moyenne','Case supérieure','2 cases supérieures','3 cases supérieures','Case inférieure','2 cases inférieures','3 cases inférieures']
CASES=[
 ('Overall product opinion','Opinion globale sur le produit',['1-Hate','2-Neutral','3-Love'],['1-Déteste','2-Neutre','3-Adore'],'Standard'),
 ('Strength','Intensité',['Too Weak','Just About Right','Too Strong'],['Trop faible','Juste comme il faut','Trop forte'],'Strength'),
 ('Olfactive attributes','Attributs olfactifs',['1-No','2-Yes'],['1-Non','2-Oui'],'Attribute'),
 ('Preferred product','Produit préféré',['1-Test product','2-Usual product'],['1-Produit testé','2-Produit habituel'],'Preference'),
 ('Bipolar attributes','Différentiel sémantique attributs',['1-Very feminine','2-Feminine','3-Neutral','4-Masculine','5-Very masculine'],['1-Très féminin','2-Féminin','3-Neutre','4-Masculin','5-Très masculin'],'Bipolaire'),
 ('Listing colours','Liste couleurs',['1-Red','2-Blue','3-Green','4-Yellow'],['1-Rouge','2-Bleu','3-Vert','4-Jaune'],'Listing'),
 ('Compared to current','Comparaison avec le produit habituel',['1-Better','2-Same','3-Worse'],['1-Meilleur','2-Identique','3-Moins bon'],'Autres'),
]


class LanguageTests(unittest.TestCase):
    def test_reviewed_types_match(self):
        for en,fr,em,fm,expected in CASES:
            for wording,metrics in [(en,em),(fr,fm)]:
                aggregates=[] if expected=='Listing' else (AGG if wording==en else FR_AGG)
                self.assertEqual(_classify('Q-1',wording,metrics+aggregates)[0],expected,wording)
    def test_default_metric_positions_match(self):
        for en,fr,em,fm,expected in CASES:
            a=em+AGG;b=fm+FR_AGG
            self.assertEqual([a.index(m) for m in default_metric_selection(expected,a)],[b.index(m) for m in default_metric_selection(expected,b)],fr)
    def test_french_explicit_boxes_are_aggregates(self):
        self.assertEqual([_metric_key(m) for m in FR_AGG],[_metric_key(m) for m in AGG])
        for m in ['1-Couleur supérieure','2-Aspect supérieur','3-Très haut']:
            self.assertNotIn(_metric_key(m),{'top_1','top_2','top_3'})
    def test_analysis_does_not_change_legacy_normalization(self):
        source='Q-4-Intention d’achat après séchage'
        self.assertIn('purchase intent',analysis_normal(source))
        self.assertNotIn('purchase intent',_normal(source))
        self.assertEqual(analysis_normal('Xylophone mystérieux'),'xylophone mysterieux')
    def test_cata_presets_exclude_numbered_french_aggregates(self):
        from templyfier.review import metric_presets
        row={'Type':'CATA','Available metric list':['1-Non','2-Oui']+FR_AGG,'Selected metrics':['2-Oui']}
        presets=metric_presets(row)
        self.assertEqual(presets['Toutes les modalités'],['1-Non','2-Oui'])
        self.assertEqual(presets['CATA : réponses 2-'],['2-Oui'])
        self.assertEqual(presets['CATA : les deux réponses'],['1-Non','2-Oui'])
    def test_numeric_formulas_and_fills_match_monadic_and_paired(self):
        for paired in (False,True):
            outputs=[]
            for french in (False,True):
                questions={f'Q-{i}-{case[1 if french else 0]}':case[3 if french else 2]+(FR_AGG if french else AGG) for i,case in enumerate(CASES,1)}
                raw=source_bytes(paired,questions)
                wb=load_workbook(BytesIO(raw));ws=wb.active
                for row in range(8,ws.max_row+1):
                    mean=_metric_key(ws.cell(row,2).value)=='mean'
                    ws.cell(row,3,3.1 if mean else .2);ws.cell(row,4,3.5 if mean else .3)
                buffer=BytesIO();wb.save(buffer);raw=buffer.getvalue()
                info,rows=prepared(raw)
                by_id={r['Question ID']:r for r in rows}
                ordered=[]
                for i,(qid,metrics) in enumerate(questions.items(),1):
                    row=by_id[qid];self.assertEqual(row['Question ID'],qid)
                    row.update({'Keep':True,'Order':i,'Section':'Common','Display label':f'Question {i}','Metric label':'','KPI Summary':False})
                    selected=default_metric_selection(CASES[i-1][4],metrics)
                    set_metric_selection(row,selected,{m:f'Metric {metrics.index(m)}' for m in selected});ordered.append(row)
                data,_=build_smart_toplines([('test.xlsx',raw)],export_rows(ordered),split_names=['TOTAL'],benchmark_positions=[0],standard_metrics=AGG,test_type='Paired' if paired else 'Monadic',include_screeners=False,include_sections=False)
                sheet=load_workbook(BytesIO(data)).worksheets[0]
                outputs.append([(c.coordinate,c.value,c.number_format,c.fill.patternType,c.fill.fgColor.rgb) for line in sheet.iter_rows(min_row=6) for c in line if isinstance(c.value,(int,float)) or (isinstance(c.value,str) and c.value.startswith('='))])
            self.assertEqual(outputs[0],outputs[1])
