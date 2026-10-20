from io import BytesIO
import unittest

from openpyxl import Workbook

from templyfier.core import TemplyfierError
from templyfier.smart import (
    _cmr_match_score,
    _embedded_data_sheets,
    CmrProduct,
    cmr_benchmark_positions,
    inspect_smart_package,
    match_cmr_products,
    parse_cmr_products,
)


class CmrMatchingTests(unittest.TestCase):
    @staticmethod
    def _cmr_bytes(rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "CMR code", "Formula code", "Fantasy name", "Formula description",
            "Formula type", "Fr-Land ID", "Brand", "Line Variant",
        ])
        for row in rows:
            sheet.append(row)
        payload = BytesIO()
        workbook.save(payload)
        return payload.getvalue()

    def test_composite_gsight_headers_use_explicit_cmr_benchmark_type(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "CMR code", "Formula code", "Fantasy name", "Formula description",
            "Formula type", "Fr-Land ID",
        ])
        sheet.append(["A26", "", "Benchmark A", "", "Benchmark", 356897])
        sheet.append(["B35", "", "Benchmark B", "", "benchmark", "355433.0"])
        sheet.append(["C91", "UAH132BUA", "Candidate C", "", "Candidate", ""])
        payload = BytesIO()
        workbook.save(payload)

        matches = match_cmr_products(
            ("A26356897", "B35355433", "C9108UAH132BUA"),
            ("A26 356897", "B35 355433", "C91 0.8% UAH132BUA"),
            payload.getvalue(),
        )

        self.assertEqual([item.score for item in matches], [99, 99, 99])
        self.assertEqual(cmr_benchmark_positions(matches), (0, 1))

    def test_exact_duplicate_cmr_rows_are_safely_deduplicated(self):
        row = ["A26", "ABC123", "Benchmark A", "Description", "Benchmark", 356897]
        self.assertEqual(len(parse_cmr_products(self._cmr_bytes([row, row]))), 1)

    def test_conflicting_cmr_code_is_blocked_instead_of_chosen_by_row_order(self):
        rows = [
            ["A26", "ABC123", "Benchmark A", "Description A", "Benchmark", 356897],
            ["A26", "XYZ789", "Candidate X", "Description X", "Candidate", 999999],
        ]
        with self.assertRaisesRegex(TemplyfierError, "CMR ambiguity"):
            parse_cmr_products(self._cmr_bytes(rows))

    def test_shared_fr_land_id_across_pair_specific_cmr_codes_is_allowed(self):
        rows = [
            ["H2N", "BNCK-325430", "", "Skip Active Clean", "Benchmark", 325430],
            ["P4J", "BNCK-325430", "", "Skip Active Clean", "Benchmark", 325430],
            ["B1V", "BNCK-325430", "", "Skip Active Clean", "Benchmark", 325430],
        ]
        products = parse_cmr_products(self._cmr_bytes(rows))
        self.assertEqual([item.cmr_code for item in products], ["H2N", "P4J", "B1V"])
        matches = match_cmr_products(
            ("H2N", "P4J", "B1V"),
            ("H2N", "P4J", "B1V"),
            self._cmr_bytes(rows),
        )
        self.assertEqual([item.score for item in matches], [100, 100, 100])

    def test_repeated_gsight_position_can_reuse_one_exact_cmr_match(self):
        cmr = self._cmr_bytes([
            ["H2N", "BNCK-325430", "", "Skip Active Clean", "Benchmark", 325430],
        ])
        matches = match_cmr_products(
            ("H2N", "H2N"),
            ("Skip Active Clean", "Skip Active Clean"),
            cmr,
        )
        self.assertEqual([item.score for item in matches], [100, 100])
        self.assertEqual([item.cmr_code for item in matches], ["H2N", "H2N"])

    def test_pair_suffix_and_fantasy_name_can_reuse_one_cmr_product(self):
        cmr = self._cmr_bytes([
            ["K1M", "ANL941CNX", "APOPHIS P3P 64D", "APOPHIS 601DO MOD P3P", "Candidate", ""],
        ])
        matches = match_cmr_products(
            ("44 k1m1", "44 apophis p3p 64d"),
            ("44 - K1M1", "44 - APOPHIS P3P 64D"),
            cmr,
        )
        self.assertEqual([item.cmr_code for item in matches], ["K1M", "K1M"])
        self.assertEqual([item.score for item in matches], [97, 84])

    def test_blank_question_cells_carry_down_across_strength_metrics(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table_1_2_TAILED"
        sheet["A5"], sheet["B5"] = "Variable", "Metric"
        sheet["C3"], sheet["C4"], sheet["C5"] = "Product A", "44 - A1", "A1"
        sheet["E3"], sheet["E4"], sheet["E5"] = "Product B", "44 - B1", "B1"
        sheet["A6"], sheet["B6"], sheet["C6"], sheet["E6"] = "Q-16-1-Intensity by stages-Wet", "Too Strong", 10, 20
        sheet["B7"], sheet["C7"], sheet["E7"] = "Just about right", 80, 70
        sheet["B8"], sheet["C8"], sheet["E8"] = "Too Weak", 10, 10
        payload = BytesIO()
        workbook.save(payload)

        info = inspect_smart_package((("paired.xlsx", payload.getvalue()),))

        self.assertEqual(len(info.questions), 1)
        self.assertEqual(info.questions[0].question_type, "Strength")
        self.assertEqual(
            info.questions[0].metrics,
            ("Too Strong", "Just about right", "Too Weak"),
        )

    def test_country_word_in_description_is_not_treated_as_a_product_code(self):
        product = CmrProduct(
            cmr_code="Z4A",
            formula_code="BNCK-360387",
            fantasy_name="",
            formula_description="Ariel Alpine Machine Wash Regular France",
            formula_type="benchmark",
            fr_land_id="360387",
            brand="Ariel",
            line_variant="Alpine",
        )
        score, reason = _cmr_match_score(
            "skip active clean france 2024",
            "Skip Active Clean France 2024",
            product,
        )
        self.assertEqual((score, reason), (0, "Aucun rapprochement fiable"))

    def test_complete_brand_and_variant_match_is_strong_enough_for_benchmark_detection(self):
        product = CmrProduct(
            cmr_code="H2N",
            formula_code="BNCK-325430",
            fantasy_name="",
            formula_description="Skip Active Clean Machine Wash Regular France",
            formula_type="benchmark",
            fr_land_id="325430",
            brand="Skip",
            line_variant="Active Clean",
        )
        score, reason = _cmr_match_score(
            "skip active clean france 2024",
            "Skip Active Clean France 2024",
            product,
        )
        self.assertEqual((score, reason), (96, "Marque / variante benchmark exacte"))

    def test_pair_specific_codes_for_same_benchmark_are_not_ambiguous(self):
        rows = [
            ["H2N", "BNCK-325430", "", "Skip Active Clean France", "Benchmark", 325430],
            ["P4J", "BNCK-325430", "", "Skip Active Clean France", "Benchmark", 325430],
            ["B1V", "BNCK-325430", "", "Skip Active Clean France", "Benchmark", 325430],
            ["W3C", "BNCK-325430", "", "Skip Active Clean France", "Benchmark", 325430],
        ]
        matches = match_cmr_products(
            ("skip active clean france 2024",) * 4,
            ("Skip Active Clean France 2024",) * 4,
            self._cmr_bytes(rows),
        )
        self.assertEqual([item.score for item in matches], [92, 92, 92, 92])
        self.assertEqual(len({item.cmr_code for item in matches}), 4)

    def test_pair_specific_benchmark_rows_with_distinct_technical_ids_are_not_ambiguous(self):
        rows = [
            ["H2N", "BNCK-100001", "", "Skip formula one", "Benchmark", 100001, "Skip", "Active Clean"],
            ["P4J", "BNCK-100002", "", "Skip formula two", "Benchmark", 100002, "Skip", "Active Clean"],
            ["B1V", "BNCK-100003", "", "Skip formula three", "Benchmark", 100003, "Skip", "Active Clean"],
            ["W3C", "BNCK-100004", "", "Skip formula four", "Benchmark", 100004, "Skip", "Active Clean"],
        ]
        matches = match_cmr_products(
            ("skip active clean france 2024",) * 4,
            ("Skip - <No Line> - Active Clean - France - 2024",) * 4,
            self._cmr_bytes(rows),
        )
        self.assertEqual([item.score for item in matches], [96, 96, 96, 96])
        self.assertEqual(len({item.cmr_code for item in matches}), 4)

    def test_equal_best_cmr_matches_are_blocked(self):
        rows = [
            ["A26", "SAME123", "Product A", "", "Candidate", ""],
            ["B35", "SAME123", "Product B", "", "Candidate", ""],
        ]
        with self.assertRaisesRegex(TemplyfierError, "several CMR rows match"):
            match_cmr_products(("UNKNOWN",), ("SAME123 0.8%",), self._cmr_bytes(rows))

    def test_mismatched_gsight_product_vectors_are_blocked(self):
        with self.assertRaisesRegex(TemplyfierError, "do not have the same length"):
            match_cmr_products(("A26", "B35"), ("A26",), self._cmr_bytes([
                ["A26", "", "Product A", "", "Candidate", ""],
            ]))

    def test_stage_named_two_tailed_sheet_excludes_companion_views(self):
        workbook = Workbook()
        two_tailed = workbook.active
        two_tailed.title = "NEAT 2_TAILED"
        one_tailed = workbook.create_sheet("NEAT 1_TAILED")
        delta = workbook.create_sheet("NEAT DELTA")
        for sheet in (two_tailed, one_tailed, delta):
            sheet["A5"], sheet["B5"] = "Variable", "Metric"
            sheet["C3"], sheet["C4"], sheet["C5"] = "Product A", "100 - A", "A"
            sheet["E3"], sheet["E4"], sheet["E5"] = "Product B", "100 - B", "B"
            for row, metric in enumerate(("Mean", "Top Box", "Top 2 Boxes"), 6):
                sheet.cell(row, 1, "Q1")
                sheet.cell(row, 2, metric)
                sheet.cell(row, 3, 5.0)
                sheet.cell(row, 5, 5.5)

        self.assertEqual(
            [sheet.title for sheet in _embedded_data_sheets(workbook)],
            ["NEAT 2_TAILED"],
        )


if __name__ == "__main__":
    unittest.main()

