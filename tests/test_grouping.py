from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from templyfier.editor_model import export_rows, group_rows, prepare_rows
from templyfier.grouping import (accept_suggestion, create_group, dismiss_suggestion,
    remember_edit, reset_suggestions, separate_group, stage_for, suggest_groups, undo_edit)
from templyfier.smart import (apply_profile, build_smart_toplines, profile_from_json,
                             profile_to_json, proposal_to_row)
from tests.test_editor import AGG, prepared, source_bytes


def row(qid, label, section='PRODUCT EVALUATION', kind='Standard', metrics=None):
    return {'Question ID':qid,'Display label':label,'Metric label':'','Group ID':'',
            'Section':section,'Type':kind,'Keep':True,'Order':1,
            'Available metric list':metrics or ['1-No','2-Yes','Mean','Top Box'],
            'Selected metrics':['Mean'],'Metric labels':{},'Dismissed groups':[]}


def battery(prefix='Packaging experience', family=42):
    return [row(f'Q-{family}-{i}-{prefix}-{item}',f'{prefix}-{item}')
            for i,item in enumerate(['Easy to open','Easy to hold','Easy to store'],1)]


class GroupingTests(unittest.TestCase):
    def test_arbitrary_prefix_is_not_limited_to_known_cmi_vocabulary(self):
        for prefix in ['Packaging experience','Sustainability promises','Aftercare services']:
            rows=battery(prefix)
            before=deepcopy(rows)
            g=suggest_groups(rows)[0]
            self.assertEqual(g['label'],prefix)
            self.assertEqual(list(g['items'].values()),['Easy to open','Easy to hold','Easy to store'])
            self.assertEqual(rows,before,'A proposal must not mutate the editor')

    def test_word_prefix_without_separator(self):
        rows=[row(f'Q-50-{i}-Packaging {item}',f'Packaging {item}') for i,item in enumerate(['opening','grip','storage'],1)]
        self.assertEqual(suggest_groups(rows)[0]['label'],'Packaging')

    def test_source_prefix_survives_legacy_humanized_labels(self):
        rows=battery('Product benefit')
        for r in rows:r['Display label']=r['Display label'].split('-')[-1]
        g=suggest_groups(rows)[0]
        self.assertEqual(g['label'],'Product benefit')
        self.assertEqual(list(g['items'].values()),[r['Display label'] for r in rows])

    def test_custom_item_label_is_not_erased(self):
        rows=battery()
        rows[0]['Display label']='Opening comfort, client wording'
        g=suggest_groups(rows)[0]
        accept_suggestion(rows,g)
        self.assertEqual(rows[0]['Metric label'],'Opening comfort, client wording')

    def test_family_without_shared_words_gets_provisional_name(self):
        rows=[row(f'Q-60-{i}-{item}',item) for i,item in enumerate(['Floral','Woody','Citrus'],1)]
        g=suggest_groups(rows)[0]
        self.assertEqual(g['label'],'Battery Q-60')
        self.assertEqual(g['confidence'],'Nom à préciser')
        import hashlib
        from templyfier.grouping import _normal
        legacy=hashlib.sha1(('\0'.join(sorted(r['Question ID'] for r in rows))+'\0'+_normal('Batterie Q-60')).encode()).hexdigest()[:16]
        self.assertEqual(g['id'],legacy)
        rows[0]['Dismissed groups']=[legacy]
        self.assertEqual(suggest_groups(rows),[])

    def test_no_group_from_type_or_technical_q_prefix_alone(self):
        rows=[row('Q-U2-Loads','Q-U2-Loads'),row('Q-U9-Temperature','Q-U9-Temperature'),
              row('Q-2-Freshness','Freshness'),row('Q-3-Softness','Softness')]
        self.assertEqual(suggest_groups(rows),[])

    def test_sections_stages_types_and_scales_stay_separate(self):
        base=battery()
        other=battery()
        for field,value in [('Section','SECOND SECTION'),('Type','Listing'),
                            ('Available metric list',['1-Low','2-Mid','3-High','Mean'])]:
            changed=deepcopy(other)
            for r in changed:
                r['Question ID'] += ' extra'
                r[field]=value
            proposals=suggest_groups(base+changed)
            self.assertEqual(sorted(len(g['members']) for g in proposals),[3,3])
        wet=[row(f'Q-42-{i}-WET-Feeling-{item}',f'WET-Feeling-{item}') for i,item in enumerate(['A','B','C'],1)]
        dry=[row(f'Q-42-{i+3}-DRY-Feeling-{item}',f'DRY-Feeling-{item}') for i,item in enumerate(['A','B','C'],1)]
        self.assertEqual({g['stage'] for g in suggest_groups(wet+dry)},{'WET','DRY'})

    def test_cmi_defined_stages_are_recognised_and_kept_separate(self):
        after_application = [
            row(f'Q-42-{i}-After application-Feeling-{item}', f'Feeling-{item}')
            for i, item in enumerate(['A', 'B', 'C'], 1)
        ]
        rinse = [
            row(f'Q-42-{i+3}-Rinse-Feeling-{item}', f'Feeling-{item}')
            for i, item in enumerate(['A', 'B', 'C'], 1)
        ]
        stages = ('After application', 'Rinse')
        self.assertEqual(stage_for(after_application[0], stages), 'After application')
        self.assertEqual(stage_for(rinse[0], stages), 'Rinse')
        proposals = suggest_groups(after_application + rinse, stage_names=stages)
        self.assertEqual({group['stage'] for group in proposals}, set(stages))
        self.assertEqual(sorted(len(group['members']) for group in proposals), [3, 3])

    def test_long_custom_stage_wins_over_embedded_default_stage_word(self):
        item = row('Q-8-Skin_dry down-Overall liking', 'Overall liking')
        self.assertEqual(stage_for(item, ('Skin dry-down',)), 'Skin dry-down')

    def test_manual_unknown_stage_mapping_is_used_directly(self):
        item = row('Q-8-Overall liking', 'Overall liking')
        item['Stage'] = 'PRE-WASH'
        self.assertEqual(stage_for(item), 'PRE-WASH')
        item['Stage'] = 'PRE-WASH ; AFTER RINSE'
        self.assertEqual(stage_for(item), 'PRE-WASH/AFTER RINSE')

    def test_extra_aggregates_do_not_fragment_response_battery(self):
        rows=battery()
        rows[1]['Available metric list']=['1-No','2-Yes','Top Box']
        self.assertEqual(len(suggest_groups(rows)),1)
        self.assertEqual(len(suggest_groups(rows)[0]['members']),3)

    def test_different_batteries_with_same_name_are_not_joined(self):
        groups=suggest_groups(battery(family=42)+battery(family=43))
        self.assertEqual(len(groups),2)
        self.assertNotEqual(groups[0]['id'],groups[1]['id'])

    def test_dismissal_is_stable_and_reversible(self):
        rows=battery();g=suggest_groups(rows)[0]
        dismiss_suggestion(rows,g)
        self.assertEqual(suggest_groups(rows),[])
        self.assertEqual(suggest_groups(list(reversed(rows))),[])
        reset_suggestions(rows)
        self.assertEqual(suggest_groups(rows)[0]['id'],g['id'])

    def test_group_acceptance_is_contiguous_and_preserves_other_order(self):
        members=battery();a=row('a','Independent A');b=row('b','Independent B')
        rows=[a,members[0],b,members[1],members[2]]
        metrics={r['Question ID']:deepcopy(r['Selected metrics']) for r in rows}
        accept_suggestion(rows,suggest_groups(rows)[0],label='Pack')
        self.assertEqual([r['Question ID'] for r in rows],['a']+[r['Question ID'] for r in members]+['b'])
        self.assertEqual([r['Order'] for r in rows],[1,2,3,4,5])
        self.assertEqual(metrics,{r['Question ID']:r['Selected metrics'] for r in rows})
        self.assertEqual(len(group_rows(rows)),1)

    def test_manual_group_requires_explicit_valid_members_and_context(self):
        rows=battery();ids=[r['Question ID'] for r in rows]
        for chosen in [ids[:1],ids+['missing'],ids+[ids[0]]]:
            before=deepcopy(rows)
            with self.assertRaises(ValueError):create_group(rows,chosen,'Group')
            self.assertEqual(rows,before)
        rows[1]['Section']='Other'
        with self.assertRaises(ValueError):create_group(rows,ids,'Group')

    def test_separate_keeps_latest_labels_and_blocks_new_suggestions(self):
        rows=battery();accept_suggestion(rows,suggest_groups(rows)[0],label='Pack')
        gid=rows[0]['Group ID'];rows[0]['Metric label']='Client item'
        separate_group(rows,gid)
        self.assertEqual(rows[0]['Display label'],'Pack · Client item')
        self.assertTrue(all(not r['Group ID'] and not r['Metric label'] for r in rows))
        self.assertEqual(suggest_groups(rows),[])
        self.assertTrue(all(r['Selected metrics']==['Mean'] for r in rows))

    def test_undo_restores_order_labels_and_choices_and_has_bounded_history(self):
        rows=battery();model={'rows':deepcopy(rows),'revision':0}
        changed=deepcopy(rows);accept_suggestion(changed,suggest_groups(changed)[0])
        remember_edit(model,changed,'Group created')
        self.assertTrue(undo_edit(model));self.assertEqual(model['rows'],rows)
        self.assertEqual(model['revision'],2);self.assertFalse(undo_edit(model))
        self.assertFalse(remember_edit(model,deepcopy(rows),'No change'))
        for i in range(25):
            changed=deepcopy(model['rows']);changed[0]['Display label']=str(i)
            remember_edit(model,changed,str(i))
        self.assertEqual(len(model['history']),20)

    def test_profile_restores_accepted_dismissed_and_separated_choices(self):
        raw=source_bytes(overrides={f'Q-42-{i}-Packaging-{item}':['1-No','2-Yes']+AGG
            for i,item in enumerate(['Opening','Grip','Storage'],1)})
        info,rows=prepared(raw)
        g=suggest_groups(rows)[0]
        for action in ['dismiss','accept','separate']:
            changed=deepcopy(rows)
            if action=='dismiss':dismiss_suggestion(changed,g)
            else:
                accept_suggestion(changed,g)
                if action=='separate':separate_group(changed,g['id'])
            profile=profile_from_json(profile_to_json(export_rows(changed),{}))
            restored=prepare_rows(apply_profile(info.questions,profile),info.questions)
            self.assertEqual(suggest_groups(restored),[])
            self.assertEqual([r['Display label'] for r in restored],[r['Display label'] for r in changed])

    def test_two_same_named_groups_remain_distinct_in_excel_without_sections(self):
        for paired in [False,True]:
            raw=source_bytes(paired=paired,overrides={f'Q-{family}-{i}-Packaging-{item}':['1-No','2-Yes']+AGG
                for family in [42,43] for i,item in enumerate(['Opening','Grip','Storage'],1)})
            info,rows=prepared(raw)
            for g in suggest_groups(rows):accept_suggestion(rows,g)
            data,_=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),split_names=['TOTAL'],
                benchmark_positions=[0],standard_metrics=AGG,test_type='Paired' if paired else 'Monadic',
                include_screeners=False,include_sections=False)
            ws=load_workbook(BytesIO(data)).active
            self.assertEqual(sum(ws.cell(r,1).value=='Packaging' for r in range(6,ws.max_row+1)),2)


class RealGroupingTests(unittest.TestCase):
    def test_real_non_colour_proposals_and_lossless_exports(self):
        root=Path(__file__).resolve().parents[1]/'upload'
        path=root/'DataViz_test_2026-08-24 14_12_41(1).xlsx'
        if not path.exists():self.skipTest('Real export unavailable')
        raw=path.read_bytes();info,rows=prepared(raw)
        proposals=suggest_groups(rows)
        self.assertTrue({'WET','NEAT','DRY','Emotions','Most'}.issubset({g['label'] for g in proposals}))
        self.assertEqual(len(next(g for g in proposals if g['label']=='WET')['members']),29)
        before={r['Question ID']:(deepcopy(r['Selected metrics']),deepcopy(r['Metric labels']),r['Type']) for r in rows}
        # Source values, formats and significance must survive every proposed group.
        def signatures(config):
            data,_=build_smart_toplines([(path.name,raw)],export_rows(config),split_names=['TOTAL'],
                benchmark_positions=[0],standard_metrics=AGG,include_screeners=False,include_sections=False)
            sheet=load_workbook(BytesIO(data)).active
            def value(cell):
                if cell.data_type=='f':
                    return Translator(cell.value,origin=cell.coordinate).translate_formula(f'{cell.column_letter}1')
                return cell.value
            return sorted(repr([(value(c),c.number_format,c.fill.patternType,c.fill.fgColor.type,str(c.fill.fgColor.rgb))
                for c in line[2:]]) for line in list(sheet.iter_rows())[5:])
        original=signatures(rows)
        for g in proposals:accept_suggestion(rows,g)
        self.assertEqual(before,{r['Question ID']:(r['Selected metrics'],r['Metric labels'],r['Type']) for r in rows})
        self.assertEqual(original,signatures(rows))


if __name__=='__main__':unittest.main()
