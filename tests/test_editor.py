from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from templyfier.core import detect_layout
from templyfier.editor_model import (change_type, export_rows, group_rows, matching_selection,
                                     prepare_rows, rename_group, reorder_rows, set_metric_selection)
from templyfier.smart import (STANDARD_METRICS, _classify, _metric_key, apply_profile,
    build_smart_toplines, configured_source_rows, default_metric_selection,
    inspect_smart_package, profile_to_json, profile_from_json, proposal_to_row)

AGG=['Mean','Top Box','Top 2 Boxes','Top 3 Boxes','Bottom Box','Bottom 2 Boxes','Bottom 3 Boxes']
QUESTIONS={
    'Q-1-Overall liking':['1-Hate','2-Dislike','3-Neutral','4-Like','5-Love']+AGG,
    'Q-2-Strength':['Too Weak','Just About Right','Too Strong']+AGG,
    'Q-3-Color-A1':['1-No','2-Yes']+AGG,
    'Q-4-Color-A2':['1-No','2-Yes']+AGG,
    'Q-5-Preferred product':['1-Test Product','2-Usual blue fabric conditioner']+AGG,
    'Q-6-Most soft':['1-Test Product','2-Usual product']+AGG,
    'Q-7-Bipolar attributes':['1-Very feminine','2-Feminine','3-Neutral','4-Masculine','5-Very masculine']+AGG,
    'Q-8-Listing colours':['1-Red','2-Blue','3-Green','4-Yellow']+AGG,
    'Q-9-Unusual question':['First option','Second option'],
}


def source_bytes(paired=False, overrides=None):
    wb=Workbook();ws=wb.active;ws.title='Table_1 2_TAILED'
    ws['A1']='Stage: Paired' if paired else 'Stage: USE'
    ws['A5']='Variable';ws['B5']='Metric'
    for col,name,sample,letter in [(3,'Benchmark','50 - B1','A'),(4,'Candidate','50 - C1','B')]:
        ws.cell(3,col,name);ws.cell(4,col,sample);ws.cell(5,col,letter)
    row=8
    for q,metrics in (overrides or QUESTIONS).items():
        for m in metrics:
            ws.cell(row,1,q);ws.cell(row,2,m)
            ws.cell(row,3,3.1 if m=='Mean' else .2)
            ws.cell(row,4,3.5 if m=='Mean' else .3)
            ws.cell(row,4).fill=PatternFill('solid',fgColor='00FF00')
            row+=1
    buf=BytesIO();wb.save(buf);return buf.getvalue()


def prepared(raw):
    info=inspect_smart_package([('test.xlsx',raw)])
    rows=prepare_rows([proposal_to_row(p) for p in info.questions],info.questions)
    return info,rows


class EditorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw=source_bytes();cls.info,cls.rows=prepared(cls.raw)

    def test_classification_keeps_comparative_answers(self):
        for wording in ['Strongest','Which product','Compared to current']:
            self.assertEqual(_classify(wording,'',['1-Test product','2-Usual product']+AGG)[0],'Autres')
        for wording in ['Preference','Pref-softness','Most fresh','Which product do you prefer','Liked the most']:
            self.assertEqual(_classify(wording,'',['1-Test product','2-Usual product']+AGG)[0],'Preference')

    def test_no_yes_remains_cata(self):
        self.assertEqual(_classify('Q-1-Floral','',['1-No','2-Floral']+AGG)[0],'Attribute')

    def test_bipolar_and_listing(self):
        self.assertEqual(_classify('Q-7-Bipolar attributes','',QUESTIONS['Q-7-Bipolar attributes'])[0],'Bipolaire')
        self.assertEqual(_classify('Q-8-Listing colours','',QUESTIONS['Q-8-Listing colours'])[0],'Listing')

    def test_numbered_top_word_is_not_an_aggregate(self):
        self.assertNotEqual(_metric_key('1-On top of the bottle'),_metric_key('Top Box'))

    def test_standard_default_and_t3b_only(self):
        row=deepcopy(self.rows[0]);row['Available metric list']=QUESTIONS['Q-1-Overall liking']
        self.assertEqual(default_metric_selection('Standard',row['Available metric list']),['Mean','Top Box','Top 2 Boxes','Bottom 2 Boxes'])
        set_metric_selection(row,['Top 3 Boxes'])
        self.assertFalse(row['Mean']);self.assertTrue(row['Top 3 Boxes'])

    def test_strength_selection_can_be_overridden(self):
        wb=load_workbook(BytesIO(self.raw));ws=wb.active
        cfg={'Question ID':'Q-2-Strength','Type':'Strength','Selected metrics':['Mean']}
        found=configured_source_rows(ws,detect_layout(ws),cfg,AGG)
        self.assertEqual([ws.cell(r,2).value for r in found],['Mean'])

    def test_cata_negative_or_both(self):
        ws=load_workbook(BytesIO(self.raw)).active
        for chosen in [['1-No'],['1-No','2-Yes']]:
            cfg={'Question ID':'Q-3-Color-A1','Type':'CATA','Selected metrics':chosen}
            found=configured_source_rows(ws,detect_layout(ws),cfg,AGG)
            self.assertEqual([ws.cell(r,2).value for r in found],chosen)

    def test_empty_selection_does_not_fall_back_to_all(self):
        ws=load_workbook(BytesIO(self.raw)).active
        cfg={'Question ID':'Q-5-Preferred product','Type':'Autres','Selected metrics':[]}
        self.assertEqual(configured_source_rows(ws,detect_layout(ws),cfg,AGG),[])

    def test_listing_has_no_aggregates_by_default(self):
        self.assertEqual(default_metric_selection('Listing',QUESTIONS['Q-8-Listing colours']),['1-Red','2-Blue','3-Green','4-Yellow'])

    def test_preference_keeps_all_choices_without_aggregates(self):
        self.assertEqual(default_metric_selection('Preference',QUESTIONS['Q-5-Preferred product']),['1-Test Product','2-Usual blue fabric conditioner'])

    def test_group_rename_and_order_are_lossless(self):
        rows=deepcopy(self.rows);groups=group_rows(rows)
        color=next(g for g in groups if g['label']=='Color')
        rename_group(rows,color['id'],'Colour')
        ids=[r['Question ID'] for r in rows]
        moved=reorder_rows(rows,list(reversed(ids)))
        self.assertEqual([r['Order'] for r in moved],list(range(1,len(rows)+1)))
        self.assertTrue(all(r['Display label']=='Colour' for r in moved if r.get('Group ID')==color['id']))
        before={r['Question ID']:r['Selected metrics'] for r in rows}
        self.assertEqual(before,{r['Question ID']:r['Selected metrics'] for r in moved})

    def test_reject_duplicate_or_missing_order(self):
        ids=[r['Question ID'] for r in self.rows]
        for broken in [ids[:-1],ids[:-1]+[ids[0]],ids+['unknown']]:
            with self.assertRaises(ValueError):reorder_rows(self.rows,broken)

    def test_apply_metrics_across_cata_group_by_code(self):
        self.assertEqual(matching_selection(['1-No','2-Floral'],['1-No','2-Fruity','Top Box']),['1-No','2-Fruity'])

    def test_unknown_metric_rejected(self):
        with self.assertRaises(ValueError):set_metric_selection(deepcopy(self.rows[0]),['invented'])

    def test_profile_round_trip_preserves_selection_labels_and_order(self):
        rows=deepcopy(self.rows)
        row=next(r for r in rows if r['Question ID']=='Q-5-Preferred product')
        set_metric_selection(row,['2-Usual blue fabric conditioner'],{'2-Usual blue fabric conditioner':'Usual product'})
        restored=apply_profile(self.info.questions,profile_from_json(profile_to_json(export_rows(rows),{})))
        found=next(r for r in restored if r['Question ID']==row['Question ID'])
        self.assertEqual(found['Selected metrics'],row['Selected metrics']);self.assertEqual(found['Metric labels'],row['Metric labels'])

    def test_grouped_single_metric_can_also_be_renamed(self):
        from templyfier.smart import grouped_metric_label
        config={'Display label':'Colour','Metric label':'A1','Metric labels':{'2-Yes':'Chosen'}}
        self.assertEqual(grouped_metric_label(config,'2-Yes',1),'A1 · Chosen')

    def test_partial_profile_new_rows_do_not_create_nan_groups(self):
        import pandas as pd
        partial=deepcopy(self.rows[0]);partial['Group ID']='saved'
        records=[partial]+[proposal_to_row(p) for p in self.info.questions[1:]]
        prepared_rows=prepare_rows(pd.DataFrame(records).to_dict('records'),self.info.questions)
        self.assertTrue(all(isinstance(r['Group ID'],str) for r in prepared_rows))

    def test_old_preference_profile_is_migrated(self):
        q='Q-5-Preferred product'
        restored=apply_profile(self.info.questions,{'questions':[{'Question ID':q,'Type':'Attribute'}]})
        self.assertEqual(next(r for r in restored if r['Question ID']==q)['Type'],'Preference')

    def test_type_change_reproposes_then_allows_override(self):
        row=deepcopy(self.rows[0]);row['Available metric list']=QUESTIONS['Q-7-Bipolar attributes']
        row['Type']='Standard';change_type(row,'Bipolaire')
        self.assertEqual(len(row['Selected metrics']),5)
        set_metric_selection(row,['Top Box','Bottom 3 Boxes']);self.assertEqual(len(row['Selected metrics']),2)

    def test_union_retains_metrics_from_equal_length_split_lists(self):
        a=source_bytes(overrides={'Q-1-Overall liking':['Mean','Top Box']})
        b=source_bytes(overrides={'Q-1-Overall liking':['Mean','Top 3 Boxes']})
        info=inspect_smart_package([('a.xlsx',a),('b.xlsx',b)])
        self.assertEqual(set(info.questions[0].metrics),{'Mean','Top Box','Top 3 Boxes'})

    def test_monadic_and_paired_export_selected_answers_renames_and_colours(self):
        for paired in [False,True]:
            raw=source_bytes(paired);info,rows=prepared(raw)
            for r in rows:r['Keep']=False
            row=next(r for r in rows if r['Question ID']=='Q-5-Preferred product')
            row['Keep']=True;set_metric_selection(row,['1-Test Product','2-Usual blue fabric conditioner'],{'1-Test Product':'Tested product'})
            data,_=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),split_names=['TOTAL'],benchmark_positions=[0],
                standard_metrics=AGG,test_type='Paired' if paired else 'Monadic',include_screeners=False,include_sections=False)
            ws=load_workbook(BytesIO(data)).worksheets[0]
            metrics=[ws.cell(r,2).value for r in range(6,ws.max_row+1)]
            self.assertEqual(metrics,['Tested product','2-Usual blue fabric conditioner'])
            self.assertTrue(any(c.value==.3 for cells in ws.iter_rows(min_row=6) for c in cells))
            # Paired sources carry significance directly on candidate values.
            if paired:self.assertTrue(any(c.fill.fgColor.rgb=='0000FF00' for cells in ws.iter_rows(min_row=6) for c in cells))


class RealDataTests(unittest.TestCase):
    def test_coded_colour_battery_is_grouped_once_in_excel(self):
        path=Path(__file__).resolve().parents[1]/'upload'/'DataViz_test_2026-08-24 14_12_41(1).xlsx'
        if not path.exists():self.skipTest('Real export unavailable')
        info,rows=prepared(path.read_bytes())
        group=next(g for g in group_rows(rows) if g['label']=='Color')
        self.assertGreater(len(group['members']),10)
        rename_group(rows,group['id'],'Colour')
        chosen=[r for r in rows if r['Question ID'] in group['members']]
        data,_=build_smart_toplines([(path.name,path.read_bytes())],export_rows(chosen),split_names=['TOTAL'],
            benchmark_positions=[0],standard_metrics=AGG,include_sections=False,include_screeners=False)
        ws=load_workbook(BytesIO(data)).active
        self.assertEqual(sum(ws.cell(r,1).value=='Colour' for r in range(6,ws.max_row+1)),1)
        self.assertEqual(ws.cell(6,2).value,'A1')

    def test_original_two_tailed_regression_with_new_editor_and_hidden_gaps(self):
        path=Path(__file__).resolve().parents[1]/'upload'/'DataViz_test_2026-08-24 14_12_41.xlsx'
        if not path.exists():self.skipTest('Real export unavailable')
        info,rows=prepared(path.read_bytes())
        row=next(r for r in rows if r['Question ID']=='Q-1-1-Overall opinion')
        data,_=build_smart_toplines([(path.name,path.read_bytes())],[row],split_names=['TOTAL'],benchmark_positions=[0],
            standard_metrics=AGG,include_screeners=False,show_monadic_gaps=True)
        ws=load_workbook(BytesIO(data)).active
        self.assertEqual(ws['N7'].fill.fgColor.rgb.upper(),'0000FF00')
        self.assertEqual(ws['O7'].fill.fgColor.rgb.upper(),'0000FF00')
        self.assertFalse(ws['I7'].fill.patternType)
        data,_=build_smart_toplines([(path.name,path.read_bytes())],[row],split_names=['TOTAL'],benchmark_positions=[0],
            standard_metrics=AGG,include_screeners=False,show_monadic_gaps=False)
        hidden=load_workbook(BytesIO(data)).active
        def colours(sheet):return [(c.fill.patternType,c.fill.fgColor.rgb) for c in sheet[7] if isinstance(c.value,(int,float))]
        self.assertEqual(colours(ws),colours(hidden))

    def test_real_exports_accept_all_new_editor_settings(self):
        root=Path(__file__).resolve().parents[1]/'upload'
        files=list(root.glob('DataViz_test*.xlsx'))+list(root.glob('DATAVI*.XLS'))+[root/'DA754D~1.XLS']
        files=[p for p in files if p.exists()]
        if not files:self.skipTest('Real exports unavailable')
        for path in files:
            with self.subTest(file=path.name):
                info,rows=prepared(path.read_bytes())
                for row in rows:
                    self.assertTrue(set(row['Selected metrics']).issubset(set(row['Available metric list'])))
                data,report=build_smart_toplines([(path.name,path.read_bytes())],export_rows(rows),
                    split_names=[info.inputs[0].split_name],benchmark_positions=[0],standard_metrics=AGG,include_screeners=False)
                wb=load_workbook(BytesIO(data));self.assertGreater(wb.worksheets[0].max_row,10)


if __name__=='__main__':unittest.main()
