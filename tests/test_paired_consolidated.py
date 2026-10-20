from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from templyfier.smart import (
    audit_input_plan,
    build_smart_toplines,
    inspect_smart_package,
    proposal_to_row,
)


class ConsolidatedPairedTests(unittest.TestCase):
    @staticmethod
    def source_bytes():
        workbook = Workbook()
        total = workbook.active
        total.title = "TOTAL"
        total["A5"], total["B5"] = "Variable", "Metric"
        products = [
            (3, "Benchmark 1", "50 - B1", "A"),
            (4, "Candidate 1", "50 - C1", "B"),
            (5, "Benchmark 2", "50 - B2", "C"),
            (6, "Candidate 2", "50 - C2", "D"),
        ]
        for col, name, sample, letter in products:
            total.cell(3, col).value = name
            total.cell(4, col).value = sample
            total.cell(5, col).value = letter
        for row, metric, values in (
            (8, "Mean", [5.0, 5.5, 4.8, 5.1]),
            (9, "Top Box", [0.2, 0.3, 0.1, 0.2]),
            (10, "Top 2 Boxes", [0.5, 0.7, 0.4, 0.6]),
            (11, "Bottom 2 Boxes", [0.1, 0.0, 0.2, 0.1]),
        ):
            total.cell(row, 1).value = "Q-1-Overall liking"
            total.cell(row, 2).value = metric
            for offset, value in enumerate(values, 3):
                total.cell(row, offset).value = value
        boost = workbook.copy_worksheet(total)
        boost.title = "BOOST"
        boost["E4"] = "0 - B2"
        boost["F4"] = "0 - C2"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_one_workbook_expands_sheets_and_accepts_variable_pair_counts(self):
        raw = self.source_bytes()
        info = inspect_smart_package([("ALL SPLITS.xlsx", raw)])
        self.assertEqual(info.study_format, "HUT / in-use")
        self.assertTrue(info.suggested_test_type.startswith("Paired"))
        self.assertEqual([item.split_name for item in info.inputs], ["TOTAL", "BOOST"])
        self.assertEqual([len(item.product_keys) for item in info.inputs], [4, 2])
        audit = audit_input_plan(
            info, ["TOTAL", "BOOST"], (), test_type="Paired", automatic_exports=False
        )
        self.assertTrue(audit["ready"], audit["blockers"])

        rows = [proposal_to_row(item) for item in info.questions]
        data, report = build_smart_toplines(
            [("ALL SPLITS.xlsx", raw)], rows,
            split_names=["TOTAL", "BOOST"], benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False, test_type="Paired",
        )
        output = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(output.sheetnames, ["TOTAL", "BOOST"])
        self.assertEqual(report["pairs_by_split"], {"TOTAL": 2, "BOOST": 1})
        self.assertEqual(output["TOTAL"]["E7"].value, '=IF(COUNT(D7,C7)<2,"",D7-C7)')
        self.assertEqual(output["BOOST"]["E7"].value, '=IF(COUNT(D7,C7)<2,"",D7-C7)')

    def test_manual_mapping_supports_odd_products_and_a_shared_benchmark(self):
        workbook = load_workbook(BytesIO(self.source_bytes()))
        workbook.remove(workbook["BOOST"])
        total = workbook["TOTAL"]
        total["F4"] = "0 - C2"
        buffer = BytesIO()
        workbook.save(buffer)
        raw = buffer.getvalue()
        filename = "odd paired.xlsx"
        info = inspect_smart_package([(filename, raw)])
        self.assertEqual(info.study_format, "HUT / in-use")
        self.assertTrue(info.suggested_test_type.startswith("Paired"))
        source_key = f"{filename}::{info.inputs[0].source_sheet}"
        names = [info.inputs[0].split_name]
        automatic = audit_input_plan(
            info, names, (), test_type="Paired", automatic_exports=False
        )
        self.assertFalse(automatic["ready"])
        mapping = {source_key: [[0, 1], [0, 2]]}
        corrected = audit_input_plan(
            info, names, (), test_type="Paired", automatic_exports=False,
            paired_mappings=mapping,
        )
        self.assertTrue(corrected["ready"], corrected["blockers"])

        rows = [proposal_to_row(item) for item in info.questions]
        data, report = build_smart_toplines(
            [(filename, raw)], rows, split_names=names, benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False, test_type="Paired", paired_mappings=mapping,
        )
        output = load_workbook(BytesIO(data), data_only=False)[names[0]]
        self.assertEqual(report["pairs_by_split"][names[0]], 2)
        self.assertEqual((output["C7"].value, output["D7"].value), (5.0, 5.5))
        self.assertEqual((output["G7"].value, output["H7"].value), (5.0, 4.8))
        self.assertEqual(output["I7"].value, '=IF(COUNT(H7,G7)<2,"",H7-G7)')

    def test_repeated_product_code_in_distinct_paired_positions_is_allowed(self):
        workbook = load_workbook(BytesIO(self.source_bytes()))
        workbook.remove(workbook["BOOST"])
        total = workbook["TOTAL"]
        total["C4"] = "50 - SHARED"
        total["E4"] = "50 - SHARED"
        payload = BytesIO()
        workbook.save(payload)
        raw = payload.getvalue()
        info = inspect_smart_package([("repeated paired.xlsx", raw)])
        self.assertEqual(info.inputs[0].product_keys, ("SHARED", "C1", "SHARED", "C2"))
        rows = [proposal_to_row(item) for item in info.questions]
        data, report = build_smart_toplines(
            [("repeated paired.xlsx", raw)], rows,
            split_names=[info.inputs[0].split_name], benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False, test_type="Paired",
        )
        output = load_workbook(BytesIO(data), data_only=False).active
        self.assertEqual(report["pairs"], 2)
        self.assertEqual((output["C7"].value, output["G7"].value), (5.0, 4.8))

    def test_generic_two_tailed_tabs_are_splits_and_companion_views_are_ignored(self):
        workbook = load_workbook(BytesIO(self.source_bytes()))
        total = workbook["TOTAL"]
        workbook.remove(workbook["BOOST"])
        total.title = "Table_1_2_TAILED"
        total["A1"] = "Stage: OVERALL - IN USE"

        age = workbook.copy_worksheet(total)
        age.title = "Table_2_2_TAILED"
        age["A2"] = "Search: S-03-Age: 18 - 30 years, 31 - 40 years"
        younger = workbook.copy_worksheet(total)
        younger.title = "Table_3_2_TAILED"
        younger["A2"] = "Search: S-03-Age: Under 18 years, 18 - 30 years"
        duplicate_total = workbook.copy_worksheet(total)
        duplicate_total.title = "Table_4_2_TAILED"
        one_tailed = workbook.copy_worksheet(age)
        one_tailed.title = "Table_2_1_TAILED"
        delta = workbook.copy_worksheet(age)
        delta.title = "Table_2_2_TAILED_DELTA"
        legend = workbook.create_sheet("Table_2_Legend")
        legend["A1"] = "2-TAILED"
        payload = BytesIO()
        workbook.save(payload)

        info = inspect_smart_package((("consolidated.xlsx", payload.getvalue()),))

        self.assertEqual(len(info.inputs), 3)
        self.assertEqual(
            [item.source_sheet for item in info.inputs],
            ["Table_1_2_TAILED", "Table_2_2_TAILED", "Table_3_2_TAILED"],
        )
        self.assertEqual(
            [item.split_name for item in info.inputs],
            ["TOTAL", "18-40 YO", "UNDER 18 YEARS, 18 - 30 YEARS"],
        )

        rows = [proposal_to_row(item) for item in info.questions]
        names = [item.split_name for item in info.inputs]
        data, report = build_smart_toplines(
            (("consolidated.xlsx", payload.getvalue()),), rows,
            split_names=names, benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False, test_type="Paired",
        )
        output = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(output.sheetnames, names)
        self.assertEqual(report["data_rows_written"], 12)


if __name__ == "__main__":
    unittest.main()

