from copy import deepcopy
from io import BytesIO
import json
import unittest

from openpyxl import load_workbook

from templyfier.core import TemplyfierError
from templyfier.editor_model import export_rows, group_rows, rename_group
from templyfier.grouping import remember_edit, undo_edit, redo_edit
from templyfier.review import (audit_questions, editor_view_key, metric_presets, plan_metric_change,
    preview_rows, profile_comparison, safe_metric_match, validate_profile)
from templyfier.smart import build_smart_toplines, profile_from_json
from tests.test_editor import AGG, prepared, source_bytes
from tests.test_grouping import battery, row


class ReviewTests(unittest.TestCase):
    def test_missing_metric_blocks_but_does_not_rewrite_selection(self):
        rows=battery();rows[0]['Selected metrics']=['Top 3 Boxes']
        before=deepcopy(rows)
        audit=audit_questions(rows)
        self.assertFalse(audit['ready'])
        self.assertEqual(audit['blockers'][0]['Code'],'missing_metrics')
        self.assertEqual(rows,before)

    def test_empty_selection_label_and_invalid_type_are_objective_blockers(self):
        for field,value,code in [('Selected metrics',[],'no_metrics'),('Display label',' ','empty_label'),('Type','Unknown','unknown_type')]:
            rows=battery();rows[0][field]=value
            self.assertIn(code,{i['Code'] for i in audit_questions(rows)['blockers']})

    def test_removed_questions_do_not_block_export(self):
        rows=battery();rows[0]['Selected metrics']=[];rows[0]['Keep']=False
        self.assertTrue(audit_questions(rows)['ready'])
        for r in rows:r['Keep']=False
        self.assertFalse(audit_questions(rows)['ready'])

    def test_split_filter_excluding_everything_blocks_with_exact_question(self):
        rows=battery();rows[0]['Included splits']='Young; Boost'
        audit=audit_questions(rows,split_names=['TOTAL'])
        self.assertEqual(audit['blockers'][0]['Question ID'],rows[0]['Question ID'])
        self.assertTrue(audit_questions(rows,split_names=['TOTAL','young'])['ready'])
        rows[0]['Included splits']='Tous les splits'
        self.assertTrue(audit_questions(rows,split_names=['TOTAL'])['ready'])

    def test_warning_is_advisory_and_identifies_each_duplicate_item(self):
        rows=battery()
        for r in rows:r.update({'Group ID':'g','Metric label':'Same','Confidence':'Faible'})
        audit=audit_questions(rows)
        self.assertTrue(audit['ready'])
        self.assertEqual({r['Question ID'] for r in rows},{i['Question ID'] for i in audit['issues'] if i['Code']=='duplicate_item'})

    def test_presets_use_available_metrics_and_respect_listing_and_cata(self):
        r=row('q','listing',kind='Listing',metrics=['1-Red','2-Blue','Top Box','Mean'])
        presets=metric_presets(r)
        self.assertEqual(presets['Proposition du type'],['1-Red','2-Blue'])
        self.assertEqual(presets['Boxes disponibles, sans Mean'],['Top Box'])
        r=row('q','Floral',kind='CATA',metrics=['1-No','2-Floral','Top Box'])
        presets=metric_presets(r)
        self.assertEqual(presets['CATA : réponses 1-No'],['1-No'])
        self.assertEqual(presets['CATA : les deux réponses'],['1-No','2-Floral'])
        self.assertNotIn('Mean uniquement',presets)

    def test_answer_code_is_not_sufficient_for_comparisons_or_bipolar_scales(self):
        for kind in ['Autres','Bipolaire','Listing']:
            source=row('source','a',kind=kind,metrics=['1-Test product','2-Current product'])
            target=row('target','b',kind=kind,metrics=['1-Current product','2-Test product'])
            matched,missing=safe_metric_match(source,target,['1-Test product'])
            self.assertEqual(matched,[])
            self.assertEqual(missing,['1-Test product'])

    def test_positive_cata_codes_match_across_attributes_and_negative_does_not_flip(self):
        source=row('source','Floral',kind='CATA',metrics=['1-No','2-Floral'])
        target=row('target','Fruity',kind='CATA',metrics=['1-No','2-Fruity'])
        self.assertEqual(safe_metric_match(source,target,['2-Floral']),(['2-Fruity'],[]))
        self.assertEqual(safe_metric_match(source,target,['1-No']),(['1-No'],[]))
        target['Available metric list']=['1-Fruity','2-No']
        self.assertEqual(safe_metric_match(source,target,['1-No']),([],['1-No']))

    def test_partial_match_keeps_entire_target_recipe_and_preview_explains(self):
        rows=battery()
        rows[0]['Available metric list']=['Mean','Top Box','Top 3 Boxes']
        rows[1]['Available metric list']=['Mean','Top Box']
        rows[1]['Selected metrics']=['Top Box']
        before=deepcopy(rows)
        changed,preview=plan_metric_change(rows,rows[0]['Question ID'],['Mean','Top 3 Boxes'],{},'Toutes les questions de ce type')
        self.assertEqual(rows,before)
        self.assertEqual(changed[1]['Selected metrics'],['Top Box'])
        self.assertEqual(preview[1]['Sans correspondance sûre'],'Top 3 Boxes')
        self.assertTrue(preview[1]['Résultat'].startswith('Inchangée'))

    def test_scope_preserves_removed_and_unrelated_questions_and_labels(self):
        rows=battery();rows[0]['Group ID']='g';rows[1]['Group ID']='g'
        rows[1]['Keep']=False;rows[2]['Group ID']='other'
        changed,preview=plan_metric_change(rows,rows[0]['Question ID'],['Top Box'],{'Top Box':'TB'},'Tout le groupe')
        self.assertEqual(len(preview),1)
        self.assertEqual(changed[1:],rows[1:])
        self.assertEqual(changed[0]['Metric labels'],{'Top Box':'TB'})

    def test_empty_bulk_selection_is_rejected_without_touching_rows(self):
        rows=battery();before=deepcopy(rows)
        with self.assertRaises(ValueError):plan_metric_change(rows,rows[0]['Question ID'],[],{},'Toutes les questions de ce type')
        self.assertEqual(rows,before)

    def test_redo_restores_and_new_edit_discards_redo_branch(self):
        rows=battery();model={'rows':deepcopy(rows),'revision':0}
        first=deepcopy(rows);first[0]['Display label']='First'
        remember_edit(model,first,'Rename')
        undo_edit(model);self.assertTrue(redo_edit(model));self.assertEqual(model['rows'],first)
        undo_edit(model)
        second=deepcopy(rows);second[0]['Display label']='Second'
        remember_edit(model,second,'Different rename')
        self.assertFalse(redo_edit(model));self.assertEqual(model['rows'],second)

    def test_filter_keys_change_even_when_two_views_have_same_row_count(self):
        rows=battery()
        a=editor_view_key(rows[:1],'a','Toutes')
        b=editor_view_key(rows[1:2],'b','Toutes')
        self.assertNotEqual(a,b)
        self.assertEqual(a,editor_view_key(deepcopy(rows[:1]),'a','Toutes'))

    def test_profile_rejects_malformed_and_ambiguous_input_with_friendly_error(self):
        valid={'version':1,'questions':[{'Question ID':'Q-1','Type':'Standard','Keep':True}]}
        self.assertEqual(validate_profile(valid),valid)
        bad=[[],{'version':1,'questions':[None]},dict(valid,settings=[]),
             dict(valid,questions=valid['questions']*2)]
        for field,value in [('Type',[]),('Keep','False'),('Order',-1),('Selected metrics','Mean'),('Metric labels',[])]:
            item=deepcopy(valid);item['questions'][0][field]=value;bad.append(item)
        for payload in bad:
            with self.subTest(payload=payload):
                with self.assertRaises(TemplyfierError):profile_from_json(json.dumps(payload).encode())
        with self.assertRaises(TemplyfierError):profile_from_json(b'{broken')

    def test_profile_comparison_reports_new_absent_and_unavailable_metrics(self):
        info,rows=prepared(source_bytes())
        profile={'questions':[{'Question ID':rows[0]['Question ID'],'Selected metrics':['Invented']},
                              {'Question ID':'Q-999-Old'}]}
        report=profile_comparison(info.questions,profile)
        self.assertEqual((report['matched'],report['new'],report['absent']),(1,len(rows)-1,1))
        self.assertEqual({d['Statut'] for d in report['details']},{'Nouvelle question','Absente des exports','Métriques à adapter'})

    def test_structure_preview_matches_export_labels_and_source_metric_order(self):
        for paired in [False,True]:
            raw=source_bytes(paired=paired);info,rows=prepared(raw)
            for r in rows:
                if r['Question ID']=='Q-1-Overall liking':r['Selected metrics']=['Top Box','Mean']
            rename_group(rows,next(g for g in group_rows(rows) if g['label']=='Color')['id'],'Colour')
            preview=preview_rows(rows)
            data,_=build_smart_toplines([('test.xlsx',raw)],export_rows(rows),split_names=['TOTAL'],
                benchmark_positions=[0],standard_metrics=AGG,test_type='Paired' if paired else 'Monadic',
                include_screeners=False,include_sections=False)
            ws=load_workbook(BytesIO(data)).active
            actual=[(ws.cell(r,1).value or '',ws.cell(r,2).value or '') for r in range(6,ws.max_row+1)]
            self.assertEqual([(r['Variable clean'],r['Item / métrique']) for r in preview],actual)

    def test_metric_availability_is_indexed_per_export_and_not_saved_in_profiles(self):
        a=source_bytes(overrides={'Q-1-Overall liking':['Mean','Top Box']})
        b=source_bytes(overrides={'Q-1-Overall liking':['Mean','Top 3 Boxes']})
        from templyfier.smart import inspect_smart_package, proposal_to_row
        from templyfier.editor_model import prepare_rows
        info=inspect_smart_package([('a.xlsx',a),('b.xlsx',b)])
        question=info.questions[0]
        presence=dict(question.metric_availability)
        self.assertEqual(presence['Mean'],('a.xlsx','b.xlsx'))
        self.assertEqual(presence['Top Box'],('a.xlsx',))
        self.assertEqual(presence['Top 3 Boxes'],('b.xlsx',))
        rows=prepare_rows([proposal_to_row(question)],[question])
        self.assertEqual(rows[0]['Result export count'],2)
        self.assertNotIn('Metric availability',export_rows(rows)[0])
        self.assertNotIn('Result export count',export_rows(rows)[0])

    def test_partial_metric_availability_is_advisory_and_visible_in_preview(self):
        rows=battery();rows[0]['Metric availability']={'Mean':['a.xlsx'],'Top Box':['a.xlsx','b.xlsx']}
        rows[0]['Result export count']=2;rows[0]['Selected metrics']=['Mean','Top Box']
        audit=audit_questions(rows)
        self.assertTrue(audit['ready'])
        self.assertEqual([i['Code'] for i in audit['issues']],['partial_metrics'])
        matrix={rows[0]['Question ID']:{'Mean':['a.xlsx'],'Top Box':['a.xlsx','b.xlsx']}}
        preview=preview_rows(rows,metric_availability_by_id=matrix,result_export_count=3)
        mean=next(p for p in preview if p['Question ID']==rows[0]['Question ID'] and p['Métrique source']=='Mean')
        self.assertEqual(mean['Présence'],'1/3 exports')


if __name__=='__main__':unittest.main()
