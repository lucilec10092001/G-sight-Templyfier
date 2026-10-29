import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from templyfier.smart import build_smart_toplines, inspect_smart_package, proposal_to_row


class KpiSummaryGapRegressionTests(unittest.TestCase):
    @staticmethod
    def _significant_monadic_export() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table_1 2_TAILED"
        sheet["A5"], sheet["B5"] = "Variable", "Metric"
        sheet["C3"], sheet["C4"], sheet["C5"] = "Benchmark", "100 - B1", "A"
        sheet["D3"], sheet["D4"], sheet["D5"] = "Candidate", "100 - C1", "B"
        rows = (
            (8, "Mean", 5.0, 5.5),
            (9, "Top Box", 0.20, 0.30),
            (10, "Top 2 Boxes", 0.50, 0.70),
            (11, "Bottom 2 Boxes", 0.10, 0.05),
        )
        for row, metric, benchmark, candidate in rows:
            sheet.cell(row, 1).value = "Q-1-Overall liking"
            sheet.cell(row, 2).value = metric
            sheet.cell(row, 3).value = benchmark
            sheet.cell(row, 4).value = candidate
            sheet.cell(row, 5).value = "A"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _composite_header_export() -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "NEAT 2_TAILED"
        sheet["A5"], sheet["B5"], sheet["C5"] = "Mapped KPI/attribute", "Original question", "Metric"
        sheet["D3"], sheet["D4"], sheet["D5"] = "O1P 25% VE00971346", "155", "A"
        sheet["F3"], sheet["F4"], sheet["F5"] = "M2U 25% VE00987301", "155", "B"
        sheet["H5"] = "O1P"
        for row, metric, benchmark, candidate in (
            (6, "Mean", 4.3, 5.0),
            (7, "Top Box", 0.20, 0.30),
            (8, "Top 2 Boxes", 0.50, 0.70),
            (9, "Bottom 2 Boxes", 0.10, 0.05),
        ):
            sheet.cell(row, 1).value = "Overall Opinion"
            sheet.cell(row, 2).value = "Overall Opinion"
            sheet.cell(row, 3).value = metric
            sheet.cell(row, 4).value = benchmark
            sheet.cell(row, 6).value = candidate
            sheet.cell(row, 7).value = "A"
            sheet.cell(row, 8).value = benchmark
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_summary_uses_candidate_significance_when_gap_columns_are_visible(self):
        raw = self._significant_monadic_export()
        info = inspect_smart_package([("significant.xlsx", raw)])
        row = proposal_to_row(info.questions[0])
        row["KPI Summary"] = True
        data, _ = build_smart_toplines(
            [("significant.xlsx", raw)],
            [row],
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            benchmark_labels=("Benchmark",),
            standard_metrics=("Mean",),
            include_screeners=False,
            show_monadic_gaps=True,
            summary_scope="total",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        topline = workbook["TOTAL"]
        summary = workbook["KPI Summary"]
        self.assertFalse(topline["E6"].fill.patternType)
        self.assertEqual(summary["B5"].value, "")

    def test_composite_product_headers_still_identify_repeated_benchmark_code(self):
        raw = self._composite_header_export()
        info = inspect_smart_package([("composite.xlsx", raw)])
        self.assertEqual(info.inputs[0].product_keys, ("O1P", "M2U"))
        self.assertEqual(info.inputs[0].comparison_codes, ("O1P",))
        self.assertEqual(info.suggested_test_type, "Monadic")


if __name__ == "__main__":
    unittest.main()

