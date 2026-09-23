from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from templyfier.core import build_toplines, detect_data_sheet, detect_layout, inspect_package
from templyfier.preferences import (
    delete_cmi_profile,
    has_seen_onboarding,
    mark_onboarding_seen,
    read_preferences,
    save_cmi_profile,
    saved_cmi_profiles,
)
from templyfier.smart import (
    _active_layout,
    _summary_higher_is_better,
    audit_input_plan,
    build_smart_toplines,
    cmr_product_label,
    humanize_question,
    inspect_smart_package,
    match_cmr_products,
    proposal_to_row,
)


ROOT = Path(__file__).resolve().parents[1]
UPLOAD = ROOT / "upload"
REFERENCE = UPLOAD / "Clean file - v2.xlsx"
RAW_FILES = sorted(UPLOAD.glob("DataViz_WAVE*.xlsx"))
PAIRED_FILES = sorted(UPLOAD.glob("DataViz_Raw Data Einstein LC_*.xlsx"))
CMR_FILE = UPLOAD / "CMR Export 202529880AR0151.xlsx"
WHITE_ALL_SPLITS = UPLOAD / "G-sight output - ALL SPLITS.xlsx"
DATAVIZ_TEST = UPLOAD / "DataViz_test_2026-08-24 14_12_41.xlsx"
SKIP_CLT_FILES = [
    UPLOAD / "DATAVI~1(1).XLS",
    UPLOAD / "DATAVI~2(1).XLS",
    UPLOAD / "DATAVI~3(1).XLS",
    UPLOAD / "DATAVI~4(1).XLS",
    UPLOAD / "DA754D~1.XLS",
]

if not REFERENCE.exists() or not RAW_FILES:
    raise unittest.SkipTest("Fichiers d’exemple absents.")


class TemplyfierRegressionTests(unittest.TestCase):
    @staticmethod
    def _paired_example_bytes():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table_1 2_TAILED"
        sheet["A1"] = "Stage: Paired"
        products = [
            (3, "Benchmark 1", "50 - B1", "A"),
            (4, "Candidate 1", "50 - C1", "B"),
            (5, "Benchmark 2", "50 - B2", "C"),
            (6, "Candidate 2", "50 - C2", "D"),
        ]
        sheet["A5"], sheet["B5"] = "Variable", "Metric"
        for col, name, sample, letter in products:
            sheet.cell(3, col).value = name
            sheet.cell(4, col).value = sample
            sheet.cell(5, col).value = letter
        rows = [
            (8, "Q-1-Overall liking", "Mean", [5.0, 5.5, 4.8, 5.1]),
            (9, "Q-1-Overall liking", "Top Box", [0.2, 0.3, 0.1, 0.2]),
            (10, "Q-1-Overall liking", "Top 2 Boxes", [0.5, 0.7, 0.4, 0.6]),
            (11, "Q-1-Overall liking", "Bottom 2 Boxes", [0.1, 0.0, 0.2, 0.1]),
        ]
        for row, question, metric, values in rows:
            sheet.cell(row, 1).value = question
            sheet.cell(row, 2).value = metric
            for offset, value in enumerate(values, 3):
                sheet.cell(row, offset).value = value
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _paired_significant_example_bytes():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table_1 2_TAILED"
        sheet["A1"] = "Stage: Paired"
        sheet["A5"], sheet["B5"] = "Variable", "Metric"
        for col, name, sample, letter in (
            (3, "Benchmark", "50 - B1", "A"),
            (4, "Candidate", "50 - C1", "B"),
        ):
            sheet.cell(3, col).value = name
            sheet.cell(4, col).value = sample
            sheet.cell(5, col).value = letter
        rows = (
            (8, "Mean", 5.0, 5.5),
            (9, "Bottom 2 Boxes", 0.20, 0.10),
        )
        for row, metric, benchmark, candidate in rows:
            sheet.cell(row, 1).value = "Q-1-Overall liking"
            sheet.cell(row, 2).value = metric
            sheet.cell(row, 3).value = benchmark
            sheet.cell(row, 4).value = candidate
            # G-Sight comparison letter: uppercase = significant at 95%.
            sheet.cell(row, 5).value = "A"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_detects_all_reference_split_names(self):
        reference_sheet, infos = inspect_package(
            REFERENCE,
            [(path.name, path) for path in RAW_FILES],
        )
        self.assertEqual(reference_sheet, "TOTAL")
        self.assertEqual(
            {info.split_name for info in infos},
            {
                "TOTAL",
                "YOUNG",
                "OLD",
                "ARIEL MO",
                "SKIP MO",
                "SKIP ACTIVE CLEAN MO",
            },
        )

    def test_onboarding_is_persisted_once_and_survives_a_new_session(self):
        with TemporaryDirectory() as directory:
            preference_file = Path(directory) / "preferences.json"
            self.assertFalse(has_seen_onboarding(preference_file))
            self.assertTrue(mark_onboarding_seen(preference_file))
            self.assertTrue(has_seen_onboarding(preference_file))
            self.assertTrue(read_preferences(preference_file)["onboarding_seen"])

    def test_corrupted_onboarding_preferences_fail_safely(self):
        with TemporaryDirectory() as directory:
            preference_file = Path(directory) / "preferences.json"
            preference_file.write_text("not-json", encoding="utf-8")
            self.assertFalse(has_seen_onboarding(preference_file))
            self.assertTrue(mark_onboarding_seen(preference_file))
            self.assertTrue(has_seen_onboarding(preference_file))

    def test_cmi_profiles_are_saved_only_on_explicit_request(self):
        with TemporaryDirectory() as directory:
            preference_file = Path(directory) / "preferences.json"
            profile = {"version": 1, "settings": {"test_type": "HUT"}, "questions": []}
            self.assertEqual(saved_cmi_profiles(preference_file), {})
            self.assertTrue(save_cmi_profile("Unilever HUT", profile, preference_file))
            self.assertEqual(saved_cmi_profiles(preference_file)["Unilever HUT"], profile)
            self.assertTrue(delete_cmi_profile("Unilever HUT", preference_file))
            self.assertEqual(saved_cmi_profiles(preference_file), {})

    def test_tutorial_uses_plain_language_and_explicit_upload_instructions(self):
        tutorial = (ROOT / "onboarding.py").read_text(encoding="utf-8")
        interface = (ROOT / "smart_ui.py").read_text(encoding="utf-8")
        self.assertIn("5 splits face à 2 benchmarks = 10 exports G-Sight", tutorial)
        self.assertIn("chaque combinaison split × benchmark", interface)
        self.assertNotIn("profil Templyfier JSON", tutorial)
        self.assertNotIn("rapport JSON", interface)
        self.assertNotIn("MaxDiff", tutorial)
        self.assertNotIn('icon="✓"', interface)
        self.assertNotIn("Lecture CMI proposée", interface)
        self.assertNotIn("cœur parfum", interface)

    def test_monadic_gaps_and_significance_colours_are_enabled_by_default(self):
        interface = (ROOT / "smart_ui.py").read_text(encoding="utf-8")
        self.assertIn('settings.get("show_monadic_gaps", True)', interface)
        self.assertIn("leurs couleurs de significativité (recommandé)", interface)

    def test_detects_arbitrary_split_values_from_gsight_search(self):
        workbook = load_workbook(BytesIO(self._paired_example_bytes()))
        workbook.active["A2"] = "Search: S-1-REGION: North, Search: S-2-USAGE: Heavy users"
        buffer = BytesIO()
        workbook.save(buffer)
        info = inspect_smart_package([("arbitrary_split.xlsx", buffer.getvalue())])
        self.assertEqual(info.inputs[0].split_name, "NORTH · HEAVY USERS")

    def test_builds_formulas_and_matching_values(self):
        _, infos = inspect_package(REFERENCE, [(path.name, path) for path in RAW_FILES])
        data, report = build_toplines(
            REFERENCE,
            [(path.name, path) for path in RAW_FILES],
            split_names=[info.split_name for info in infos],
            benchmark_positions=(0,),
            product_labels=tuple(f"Fantasy {index + 1}" for index in range(len(infos[0].product_names))),
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertGreater(report["mapped_rows"], 200)
        self.assertIn("TOTAL", workbook.sheetnames)
        self.assertEqual(workbook["TOTAL"]["C7"].value, 5.37)
        self.assertEqual(workbook["TOTAL"]["F7"].value, "=E7-C7")
        self.assertEqual(workbook["YOUNG"]["C7"].value, 5.28)
        self.assertEqual(workbook["ARIEL MO"]["C7"].value, 5.29)
        total_layout = detect_layout(workbook["TOTAL"])
        self.assertEqual(workbook["TOTAL"].cell(total_layout.product_name_row, total_layout.product_cols[0]).value, "Fantasy 1")
        self.assertTrue(report["product_labels_updated"])

    def test_smart_mode_classifies_and_builds_without_template(self):
        smart_files = [(path.name, path) for path in RAW_FILES[:2]]
        info = inspect_smart_package(smart_files)
        types = {proposal.question_type for proposal in info.questions}
        self.assertEqual(types, {"Standard", "Strength", "Attribute"})
        data, report = build_smart_toplines(
            smart_files,
            [proposal_to_row(proposal) for proposal in info.questions],
            split_names=[item.split_name for item in info.inputs],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(report["mode"], "Nouveau projet intelligent")
        self.assertEqual(workbook["TOTAL"]["C7"].value, 5.37)
        self.assertEqual(workbook["TOTAL"]["F7"].value, "=E7-C7")

    def test_standard_metrics_can_vary_by_question(self):
        smart_files = [(RAW_FILES[0].name, RAW_FILES[0])]
        info = inspect_smart_package(smart_files)
        rows = []
        chosen = next(item for item in info.questions if item.question_type == "Standard" and "Mean" in item.metrics)
        row = proposal_to_row(chosen)
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric == "Mean"
        rows.append(row)
        data, report = build_smart_toplines(
            smart_files,
            rows,
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes"),
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(report["data_rows_written"], 1)
        self.assertEqual(workbook["TOTAL"]["B7"].value, "Mean")

    def test_detects_clt_multistage_and_embedded_splits(self):
        files = [
            ("DataViz_CLT Comfort Yellow UK_NEAT.xlsx", UPLOAD / "DataViz_CLT Comfort Yellow UK_NEAT.xlsx"),
            ("DataViz_CLT Comfort Yellow UK_WET.xlsx", UPLOAD / "DataViz_CLT Comfort Yellow UK_WET.xlsx"),
            ("DataViz_Comparaison bench_NEAT.xlsx", UPLOAD / "DataViz_Comparaison bench_NEAT.xlsx"),
        ]
        info = inspect_smart_package(files)
        self.assertEqual(info.study_format, "CLT")
        self.assertEqual(set(info.stages), {"NEAT", "DAMP-WET"})
        self.assertEqual(len(info.embedded_splits), 6)
        self.assertEqual(len(info.product_names), 8)
        self.assertIn("Comparaison benchmark", {item.role for item in info.inputs})

        chosen = next(item for item in info.questions if item.question_type == "Standard" and "Mean" in item.metrics)
        row = proposal_to_row(chosen)
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric == "Mean"
        data, report = build_smart_toplines(
            files[:2],
            [row],
            split_names=["NEAT", "WET"],
            benchmark_positions=(0,),
            standard_metrics=("Mean",),
            include_screeners=False,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(len(workbook.sheetnames), 12)
        self.assertEqual(report["products"], 8)

    def test_clt_batteries_are_grouped_and_follow_the_reference_story(self):
        files = [
            ("DataViz_CLT Comfort Yellow UK_NEAT.xlsx", UPLOAD / "DataViz_CLT Comfort Yellow UK_NEAT.xlsx"),
            ("DataViz_CLT Comfort Yellow UK_WET.xlsx", UPLOAD / "DataViz_CLT Comfort Yellow UK_WET.xlsx"),
        ]
        info = inspect_smart_package(files)
        olfactive = [item for item in info.questions if "Olfactive attributes" in item.question_id]
        freshness = [item for item in info.questions if "Freshness type" in item.question_id]
        colors = [item for item in info.questions if "WET-Color" in item.question_id]
        self.assertTrue(olfactive and freshness and colors)
        self.assertEqual({item.display_label for item in olfactive}, {"Olfactive Attributes"})
        self.assertIn("Natural", {item.metric_label for item in olfactive})
        self.assertIn("Long-lasting", {item.metric_label for item in olfactive})
        self.assertEqual({item.section for item in olfactive}, {"FRAGRANCE CHARACTERISTICS"})
        self.assertEqual({item.display_label for item in freshness}, {"Type of Freshness"})
        self.assertIn("Lavender, Aromatic", {item.metric_label for item in freshness})
        self.assertLess(max(item.order for item in olfactive), min(item.order for item in freshness))
        self.assertLess(max(item.order for item in freshness), min(item.order for item in colors))
        wet_fit = next(item for item in info.questions if item.question_id == "Q-09-WET-Yellow fit")
        neat_ofl = next(item for item in info.questions if item.question_id == "Q-02-NEAT-OFL")
        self.assertLess(max(item.order for item in colors), wet_fit.order)
        self.assertLess(wet_fit.order, neat_ofl.order)

    def test_short_named_clt_excludes_maxdiff_and_groups_binary_batteries(self):
        info = inspect_smart_package([(path.name, path) for path in SKIP_CLT_FILES])
        self.assertEqual(info.study_format, "CLT")
        self.assertEqual(
            [item.split_name for item in info.inputs],
            ["TOTAL", "18-40 YO", "41-64 YO", "ARIEL MO", "SKIP MO"],
        )
        questions = {item.question_id: item for item in info.questions}
        favorite = questions["Q-08-Favorite"]
        self.assertFalse(favorite.keep)
        self.assertEqual(favorite.question_type, "Delete")
        self.assertEqual(favorite.section, "HORS TOPLINES")
        self.assertEqual(proposal_to_row(favorite)["CMI note"], "Question technique laissée de côté")

        benefits = [item for item in info.questions if "-Benefits-" in item.question_id and item.keep]
        attributes = [item for item in info.questions if "-Attributes-" in item.question_id]
        freshness = [item for item in info.questions if "Freshness type" in item.question_id]
        self.assertEqual({item.display_label for item in benefits}, {"Benefits"})
        self.assertEqual({item.section for item in benefits}, {"PRODUCT BENEFITS"})
        self.assertIn("Hygiene, deep clean", {item.metric_label for item in benefits})
        self.assertEqual({item.display_label for item in attributes}, {"Fragrance Attributes"})
        self.assertEqual({item.section for item in attributes}, {"FRAGRANCE CHARACTERISTICS"})
        self.assertIn("Aggressive", {item.metric_label for item in attributes})
        self.assertLess(questions["Q-02-OFO"].order, questions["Q-03-Intensity"].order)
        self.assertLess(questions["Q-03-Intensity"].order, questions["Q-04-1-Statements-The scent of this washing powder meets my expectations (in terms of freshness, fragrance, cleanliness, etc.)"].order)
        self.assertLess(max(item.order for item in benefits), min(item.order for item in freshness))
        self.assertLess(max(item.order for item in freshness), min(item.order for item in attributes))

    def test_detects_in_use_shape_as_hut(self):
        info = inspect_smart_package([(RAW_FILES[0].name, RAW_FILES[0])])
        self.assertEqual(info.study_format, "HUT / in-use")

    def test_hut_strength_and_sections_are_recognized(self):
        path = UPLOAD / "TOTAL.xlsx"
        info = inspect_smart_package([(path.name, path)])
        questions = {item.question_id: item for item in info.questions}
        self.assertEqual(info.study_format, "HUT / in-use")
        self.assertEqual(questions["Q-15-SCENT STRENGTH"].question_type, "Strength")
        self.assertEqual(questions["Q-16-1-While in the bottle / opening the bottle"].question_type, "Strength")
        self.assertEqual(questions["Q-14-1-While in the bottle / opening the bottle"].section, "FRAGRANCE JOURNEY - LIKING")
        self.assertEqual(questions["Q-18-1-Edible / yummy"].section, "FRAGRANCE CHARACTERISTICS")
        self.assertEqual(questions["Q-20-1-Relaxed"].section, "FRAGRANCE CHARACTERISTICS")
        self.assertEqual(questions["Q-18-1-Edible / yummy"].display_label, "Olfactive Space")
        self.assertEqual(questions["Q-18-1-Edible / yummy"].metric_label, "Edible / yummy")
        self.assertEqual(questions["Q-20-1-Relaxed"].display_label, "Emotions")
        self.assertEqual(questions["Q-20-1-Relaxed"].metric_label, "Relaxed")

    def test_hut_core_labels_and_cmi_review_notes_are_non_blocking(self):
        path = UPLOAD / "TOTAL.xlsx"
        info = inspect_smart_package([(path.name, path)])
        rows = {item.question_id: proposal_to_row(item) for item in info.questions}
        self.assertEqual(rows["Q-21-PURCHASE"]["Display label"], "Purchase Intent")
        self.assertEqual(rows["Q-7-Expectations"]["Section"], "PRODUCT EVALUATION")
        self.assertEqual(rows["Q-8-Softness"]["Display label"], "Overall Softness Opinion")
        self.assertEqual(rows["Q-8-Softness"]["CMI role"], "Indicateur principal")
        self.assertEqual(rows["Q-22-WHITE BOTTLE SUITABILITY"]["Section"], "FRAGRANCE BENEFITS")

        yumos = UPLOAD / "DA1981~1.XLS"
        yumos_info = inspect_smart_package([(yumos.name, yumos)])
        comparison = next(item for item in yumos_info.questions if item.question_id == "Q-10-Comparison to current")
        comparison_row = proposal_to_row(comparison)
        self.assertTrue(comparison_row["Keep"])
        self.assertIn("À confirmer", comparison_row["CMI note"])

    def test_clean_labels_only_remove_the_technical_question_prefix(self):
        self.assertEqual(
            humanize_question("Q-02P-OVERALL PRODUCT OPINION", "Overall Product Liking"),
            "OVERALL PRODUCT OPINION",
        )
        self.assertEqual(
            humanize_question("Q-14-1-While in the bottle / opening the bottle", "Touchpoint"),
            "While in the bottle / opening the bottle",
        )
        self.assertEqual(humanize_question("Q7-This is a high-quality product", "High Quality"), "This is a high-quality product")

    def test_intelligent_order_starts_with_product_then_fragrance_and_finishes_with_colors(self):
        path = UPLOAD / "TOTAL.xlsx"
        info = inspect_smart_package([(path.name, path)])
        kept = [item for item in info.questions if item.keep]
        sections = [item.section for item in kept]
        self.assertEqual(kept[0].question_id, "Q-5-Overall Product Opinion")
        self.assertLess(sections.index("PRODUCT BENEFITS"), sections.index("FRAGRANCE EVALUATION"))
        self.assertLess(sections.index("FRAGRANCE EVALUATION"), sections.index("FRAGRANCE STRENGTH"))
        self.assertEqual(sections[-1], "FRAGRANCE CHARACTERISTICS")

    def test_paired_hut_order_follows_the_cmi_storyline(self):
        path = UPLOAD / "DataViz_Raw Data Einstein LC_2026-06-15 14_27_01.xlsx"
        info = inspect_smart_package([(path.name, path)])
        order = {item.question_id: item.order for item in info.questions}
        storyline = [
            "Q-02P-OVERALL PRODUCT OPINION",
            "Q-03P-PURCHASE INTENT",
            "Q-04-RCR",
            "Q-05-EVALUATION VS EXPECTATIONS",
            "Q8-Scent experience overall",
            "Q-06-NEW & DIFFERENT",
            "Q7-This product gives me a benefit that I cannot get from other products",
            "Q-14-Overall Fragrance Rating",
            "Q-15-1-When opening the pack",
            "Q-16-Intensity",
            "Q-18-1-When opening the pack",
            "Q-19-Characteristics-Aggressive",
        ]
        self.assertEqual([order[item] for item in storyline], sorted(order[item] for item in storyline))
        technical = next(item for item in info.questions if item.question_id == "Q-44-Concept malodor-1st product tested")
        self.assertEqual(technical.section, "USAGE & CONTEXT")
        self.assertFalse(technical.keep)
        malodor_benefit = next(
            item for item in info.questions
            if item.question_id == "Q11-Preventing bad odours appearing when laundry is left in the washing machine"
        )
        self.assertEqual(malodor_benefit.section, "PRODUCT BENEFITS")

    def test_colors_are_grouped_under_one_variable_with_codes_as_metrics(self):
        path = UPLOAD / "TOTAL.xlsx"
        info = inspect_smart_package([(path.name, path)])
        colors = [item for item in info.questions if "COLOUR" in item.question_id][:3]
        rows = [proposal_to_row(item) for item in colors]
        self.assertEqual([row["Display label"] for row in rows], ["Color", "Color", "Color"])
        self.assertEqual([row["Metric label"] for row in rows], ["A1", "A2", "A4"])
        data, _ = build_smart_toplines(
            [(path.name, path)],
            rows,
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Top Box",),
            include_screeners=False,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        sheet = workbook["TOTAL"]
        self.assertEqual((sheet["A7"].value, sheet["B7"].value), ("Color", "A1"))
        self.assertEqual((sheet["A8"].value, sheet["B8"].value), (None, "A2"))
        self.assertEqual((sheet["A9"].value, sheet["B9"].value), (None, "A4"))

    def test_hut_benchmark_configurations_are_not_consumer_stages(self):
        files = [
            ("18-41 YO.xlsx", UPLOAD / "18-41 YO.xlsx"),
            ("AGE 18-40 YO.xlsx", UPLOAD / "AGE 18-40 YO.xlsx"),
            ("41-65 YO.xlsx", UPLOAD / "41-65 YO.xlsx"),
            ("AGE 41-65YO.xlsx", UPLOAD / "AGE 41-65YO.xlsx"),
        ]
        info = inspect_smart_package(files)
        self.assertEqual(info.study_format, "HUT / in-use")
        self.assertEqual(info.suggested_benchmark_count, 2)
        self.assertEqual({item.split_name for item in info.inputs}, {"18-40 YO", "41-64 YO"})
        self.assertEqual(sum(item.role == "Résultats" for item in info.inputs), 2)
        self.assertTrue(any("pas des touchpoints" in warning for warning in info.warnings))

    def test_hut_strength_recipe_and_two_benchmark_gaps(self):
        files = [
            ("AGE 18-40 YO.xlsx", UPLOAD / "AGE 18-40 YO.xlsx"),
            ("AGE 41-65YO.xlsx", UPLOAD / "AGE 41-65YO.xlsx"),
        ]
        info = inspect_smart_package(files)
        chosen = next(item for item in info.questions if item.question_id == "Q-15-SCENT STRENGTH")
        data, report = build_smart_toplines(
            files,
            [proposal_to_row(chosen)],
            split_names=["18-40 YO", "41-65 YO"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean",),
            include_screeners=False,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        sheet = workbook["18-40 YO"]
        self.assertEqual([sheet[f"B{row}"].value for row in (7, 8, 9)], ["Too Weak", "Just about right", "Too Strong"])
        self.assertEqual(sheet["H7"].value, "=G7-C7")
        self.assertEqual(sheet["I7"].value, "=G7-E7")
        self.assertEqual(report["products"], 6)

    def test_multi_benchmark_can_create_one_sheet_per_benchmark(self):
        path = UPLOAD / "DataViz_RAW DATA with main_boost split_2026-08-19 18_14_23.xlsx"
        if not path.exists():
            self.skipTest("Nouvel exemple multi-benchmarks absent.")
        info = inspect_smart_package([(path.name, path)])
        self.assertEqual(info.suggested_benchmark_count, 2)
        chosen = next(item for item in info.questions if item.question_id == "Q-5-OPO")
        data, report = build_smart_toplines(
            [(path.name, path)],
            [proposal_to_row(chosen)],
            split_names=["MAIN 41-64"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="separate",
            benchmark_labels=("YUMOS", "VERNEL"),
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["MAIN 41-64 vs YUMOS", "MAIN 41-64 vs VERNEL"])
        yumos = workbook["MAIN 41-64 vs YUMOS"]
        vernel = workbook["MAIN 41-64 vs VERNEL"]
        self.assertEqual(yumos["F7"].value, "=E7-C7")
        self.assertEqual(vernel["I7"].value, "=H7-C7")
        # DELTA marks Vernel's -8.85 gap in red, but 2_TAILED does not mark it
        # significant. The topline must preserve the selected 2_TAILED view.
        self.assertFalse(yumos["F7"].fill.patternType)
        self.assertEqual(yumos["O7"].fill.fgColor.rgb, "00800000")
        self.assertEqual(vernel["I7"].fill.fgColor.rgb, "00ccffcc")
        self.assertEqual(report["benchmark_sheet_mode"], "separate")

    def test_selected_two_tailed_colours_are_preserved_instead_of_delta_palette(self):
        if not DATAVIZ_TEST.exists():
            self.skipTest("DataViz de contrôle des couleurs absent.")
        info = inspect_smart_package([(DATAVIZ_TEST.name, DATAVIZ_TEST)])
        chosen = next(item for item in info.questions if item.question_id == "Q-1-1-Overall opinion")
        row = proposal_to_row(chosen)
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric in {"Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"}
        data, _ = build_smart_toplines(
            [(DATAVIZ_TEST.name, DATAVIZ_TEST)],
            [row],
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_labels=("Comfort Blue skies",),
            show_monadic_gaps=True,
            summary_scope="none",
        )
        sheet = load_workbook(BytesIO(data), data_only=False).active
        # Source 2_TAILED Q31 is bright green; DELTA Q31 is medium green.
        self.assertEqual((sheet["N7"].fill.fgColor.rgb or "").upper(), "0000FF00")
        self.assertEqual((sheet["O7"].fill.fgColor.rgb or "").upper(), "0000FF00")
        self.assertNotEqual((sheet["O7"].fill.fgColor.rgb or "").upper(), "004BCB4F")
        # Source 2_TAILED K31 is white, even though DELTA can colour other rows.
        self.assertFalse(sheet["I7"].fill.patternType)

    def test_hiding_monadic_gaps_preserves_selected_two_tailed_colours_on_values(self):
        if not DATAVIZ_TEST.exists():
            self.skipTest("DataViz de contrôle des couleurs absent.")
        info = inspect_smart_package([(DATAVIZ_TEST.name, DATAVIZ_TEST)])
        chosen = next(item for item in info.questions if item.question_id == "Q-1-1-Overall opinion")
        row = proposal_to_row(chosen)
        for metric in (
            "Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes",
            "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes",
        ):
            row[metric] = metric in {"Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"}
        common = dict(
            raw_files=[(DATAVIZ_TEST.name, DATAVIZ_TEST)],
            question_rows=[row],
            split_names=["TOTAL"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_labels=("Comfort Blue skies",),
            summary_scope="none",
        )
        with_gaps, _ = build_smart_toplines(show_monadic_gaps=True, **common)
        without_gaps, report = build_smart_toplines(show_monadic_gaps=False, **common)
        visible = load_workbook(BytesIO(with_gaps), data_only=False).active
        hidden = load_workbook(BytesIO(without_gaps), data_only=False).active

        visible_value_colours = [
            (cell.fill.patternType, (cell.fill.fgColor.rgb or "").upper())
            for cell in visible[7]
            if isinstance(cell.value, (int, float))
        ]
        hidden_value_colours = [
            (cell.fill.patternType, (cell.fill.fgColor.rgb or "").upper())
            for cell in hidden[7]
            if isinstance(cell.value, (int, float))
        ]
        self.assertEqual(hidden_value_colours, visible_value_colours)
        self.assertTrue(any(isinstance(cell.value, str) and cell.value.startswith("=") for cell in visible[7]))
        self.assertFalse(any(isinstance(cell.value, str) and cell.value.startswith("=") for cell in hidden[7]))
        self.assertFalse(report["show_monadic_gaps"])

    def test_automatic_exports_consolidate_same_split_with_permuted_products(self):
        paths = [
            UPLOAD / "DataViz_RAW DATA with main_boost split_2026-08-19 18_31_58.xlsx",
            UPLOAD / "DataViz_RAW DATA with main_boost split_2026-08-19 18_14_23(2).xlsx",
        ]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemples versus Vernel/Yumos absents.")
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        self.assertEqual({item.split_name for item in info.inputs}, {"18-40 YO", "41-64 YO"})
        chosen = next(item for item in info.questions if item.question_id == "Q-5-OPO")
        data, report = build_smart_toplines(
            files,
            [proposal_to_row(chosen)],
            split_names=["MAIN", "MAIN"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("VERNEL", "YUMOS"),
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["MAIN vs VERNEL", "MAIN vs YUMOS"])
        self.assertIn("Vernel", workbook["MAIN vs VERNEL"]["C3"].value)
        self.assertIn("Yumos", workbook["MAIN vs YUMOS"]["C3"].value)
        self.assertEqual(workbook["MAIN vs YUMOS"]["F7"].value, "=E7-C7")
        self.assertEqual(report["missing_benchmark_exports"], [])
        self.assertEqual(report["benchmark_sheet_mode"], "auto_exports")

    def test_yumos_hut_two_benchmarks_uses_reference_order_and_flexible_rendering(self):
        names = [
            "DATAVI~1.XLS", "DATAVI~2.XLS", "DATAVI~3.XLS", "DATAVI~4.XLS",
            "DAD68F~1.XLS", "DABA5C~1.XLS", "DA9C69~1.XLS", "DA1981~1.XLS",
            "DAC581~1.XLS", "DAAE49~1.XLS", "DA3D9E~1.XLS", "DA6FB3~1.XLS",
            "DA26AE~1.XLS", "DA873C~1.XLS", "DAF0E7~1.XLS", "DA8983~1.XLS",
        ]
        paths = [UPLOAD / name for name in names]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemple Yumos Pink HUT absent.")
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        self.assertEqual(
            {item.split_name for item in info.inputs},
            {"TOTAL", "YUMOS ORKIDE MO", "YUMOS MO", "OTHER BRANDS", "ISTANBUL", "ANKARA", "18-40 YO", "41-64 YO"},
        )
        self.assertEqual(info.suggested_benchmark_count, 2)
        specials = {
            item.question_id: item.suggested_splits
            for item in info.questions
            if item.question_id in {
                "Q-25-F-Comparison smell with Yumos Orkide",
                "Q-26-F-Feeling about the new fragrance",
            }
        }
        self.assertEqual(set(specials.values()), {("YUMOS ORKIDE MO",)})

        chosen_ids = {
            "Q-5-OPO",
            "Q-25-F-Comparison smell with Yumos Orkide",
            "Q-26-F-Feeling about the new fragrance",
        }
        question_rows = [proposal_to_row(item) for item in info.questions if item.question_id in chosen_ids]
        labels = ["Yumos Orchid", "Vernel Aromatherapy Fresh Rose", "ELLA", "MYSTIQUE", "CHEERFUL", "MYSTICAL"]
        subtitles = ["", "", "FORMULA D6A", "FORMULA E4L", "FORMULA C5O", "FORMULA M3Q"]
        data, report = build_smart_toplines(
            files,
            question_rows,
            split_names=[item.split_name for item in info.inputs],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Yumos", "Vernel"),
            output_sheet_order="benchmark_first",
            show_monadic_gaps=False,
            product_labels=labels,
            product_subtitles=subtitles,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(len(workbook.sheetnames), 16)
        self.assertTrue(all(name.endswith("vs Yumos") for name in workbook.sheetnames[:8]))
        self.assertTrue(all(name.endswith("vs Vernel") for name in workbook.sheetnames[8:]))
        total = workbook["TOTAL vs Yumos"]
        orkide = workbook["YUMOS ORKIDE MO vs Yumos"]
        self.assertEqual((total["C3"].value, total["E3"].value, total["G3"].value), tuple(labels[:3]))
        self.assertEqual(total["G5"].value, "FORMULA D6A")
        self.assertFalse(any("Gap vs" in str(cell.value) for cell in total[5]))
        self.assertNotIn("Comparison smell with Yumos Orkide", [cell.value for cell in total["A"]])
        self.assertIn("Comparison smell with Yumos Orkide", [cell.value for cell in orkide["A"]])
        self.assertEqual(report["output_sheet_order"], "benchmark_first")
        self.assertFalse(report["show_monadic_gaps"])

    def test_kpi_summary_suggestions_are_decisional_and_deduplicated(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DAC581~1.XLS"]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemple Yumos Pink HUT absent.")
        info = inspect_smart_package([(path.name, path) for path in paths])
        selected = [item for item in info.questions if item.summary_keep]
        self.assertGreaterEqual(len(selected), 6)
        self.assertLessEqual(len(selected), 10)
        self.assertEqual(len({item.summary_label.casefold() for item in selected}), len(selected))
        self.assertTrue(any("OPO" in item.question_id for item in selected))
        self.assertTrue(any("PI" in item.question_id for item in selected))
        self.assertFalse(any(item.section in {"USAGE & CONTEXT", "FRAGRANCE CHARACTERISTICS"} for item in selected))
        self.assertFalse(any(item.question_type == "Strength" for item in selected))

    def test_kpi_summary_total_creates_one_sheet_per_benchmark(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DAC581~1.XLS"]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemple Yumos Pink HUT absent.")
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        rows = [proposal_to_row(item) for item in info.questions if item.summary_keep]
        data, report = build_smart_toplines(
            files,
            rows,
            split_names=["TOTAL", "TOTAL"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Yumos", "Vernel"),
            show_monadic_gaps=False,
            summary_scope="total",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames[:2], ["KPI Summary vs Yumos", "KPI Summary vs Vernel"])
        self.assertEqual(report["summary_sheets"], workbook.sheetnames[:2])
        summary = workbook["KPI Summary vs Yumos"]
        self.assertEqual(summary["A1"].value, "TOTAL SAMPLE VS. YUMOS")
        symbols = {
            summary.cell(row, col).value
            for row in range(5, 10)
            for col in range(2, summary.max_column + 1)
        }
        self.assertTrue(symbols.issubset({"▲", "△", "=", "▽", "▼"}))
        self.assertTrue(symbols - {"="})
        self.assertEqual(report["summary_scope"], "total")

    def test_kpi_summary_reconciles_direction_with_values_for_each_benchmark(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DAC581~1.XLS"]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemple Yumos Pink HUT absent.")
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        rows = [proposal_to_row(item) for item in info.questions if item.summary_keep]
        data, _ = build_smart_toplines(
            files,
            rows,
            split_names=["TOTAL", "TOTAL"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Yumos", "Vernel"),
            show_monadic_gaps=True,
            summary_scope="total",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        vs_yumos = workbook["KPI Summary vs Yumos"]
        vs_vernel = workbook["KPI Summary vs Vernel"]
        # MIX...4108 is significantly above both controls on OPO: it must be a win.
        self.assertGreater(workbook["TOTAL vs Yumos"]["H7"].value, workbook["TOTAL vs Yumos"]["C7"].value)
        self.assertEqual(vs_yumos["B6"].value, "▲")
        self.assertGreater(workbook["TOTAL vs Vernel"]["H7"].value, workbook["TOTAL vs Vernel"]["C7"].value)
        self.assertEqual(vs_vernel["B6"].value, "▲")
        # Each benchmark sheet follows its own selected 2_TAILED export exactly.
        self.assertEqual(vs_yumos["B5"].value, "=")
        self.assertEqual(vs_vernel["B5"].value, "=")

    def test_input_audit_detects_missing_split_benchmark_combinations(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DAC581~1.XLS"]
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        complete = audit_input_plan(
            info, ["TOTAL", "TOTAL"], (0, 1), test_type="Monadic", automatic_exports=True
        )
        self.assertTrue(complete["ready"])
        self.assertEqual({row["Contrôle"] for row in complete["rows"]}, {"Prêt"})
        incomplete = audit_input_plan(
            info, ["TOTAL", "18-40 YO"], (0, 1), test_type="Monadic", automatic_exports=True
        )
        self.assertFalse(incomplete["ready"])
        self.assertTrue(any("Export manquant" in message for message in incomplete["blockers"]))

    def test_kpi_details_explain_values_and_consensus_conflicts(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DAC581~1.XLS"]
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        rows = [proposal_to_row(item) for item in info.questions if item.summary_keep]
        data, report = build_smart_toplines(
            files,
            rows,
            split_names=["TOTAL", "TOTAL"],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Yumos", "Vernel"),
            summary_scope="total",
            summary_metric_strategy="consensus",
            include_summary_details=True,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(report["summary_detail_sheet"], "KPI Details")
        details = workbook["KPI Details"]
        headers = [details.cell(1, col).value for col in range(1, details.max_column + 1)]
        self.assertEqual(headers[-5:], ["Valeur benchmark", "Écart", "Seuil", "Résultat", "Explication"])
        # DELTA used to create artificial conflicts here. The selected
        # 2_TAILED exports contain no supported opposite-direction conflict.
        self.assertFalse(any(details.cell(row, 13).value == "±" for row in range(2, details.max_row + 1)))
        self.assertFalse(any("directions opposées" in str(details.cell(row, 14).value) for row in range(2, details.max_row + 1)))

    def test_cmi_can_override_the_favourable_scale_direction(self):
        self.assertTrue(_summary_higher_is_better("Mean", {"Sens favorable": "Automatique"}))
        self.assertFalse(_summary_higher_is_better("Mean", {"Sens favorable": "Plus bas"}))
        self.assertTrue(_summary_higher_is_better("Just about right", {"Sens favorable": "Idéal au centre"}))
        self.assertFalse(_summary_higher_is_better("Too strong", {"Sens favorable": "Idéal au centre"}))
        self.assertIsNone(_summary_higher_is_better("Mean", {"Sens favorable": "Idéal au centre"}))
        self.assertIsNone(_summary_higher_is_better("Mean", {"Sens favorable": "Neutre"}))

    def test_kpi_summary_can_stack_every_split(self):
        paths = [UPLOAD / "DATAVI~1.XLS", UPLOAD / "DA9C69~1.XLS"]
        if not all(path.exists() for path in paths):
            self.skipTest("Exemple Yumos Pink HUT absent.")
        files = [(path.name, path) for path in paths]
        info = inspect_smart_package(files)
        chosen = next(item for item in info.questions if item.question_id == "Q-5-OPO")
        row = proposal_to_row(chosen)
        row["KPI Summary"] = True
        data, report = build_smart_toplines(
            files,
            [row],
            split_names=["TOTAL", "18-40 YO"],
            benchmark_positions=(0,),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes"),
            include_screeners=False,
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Yumos",),
            show_monadic_gaps=False,
            summary_scope="all",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        summary = workbook["KPI Summary"]
        titles = [summary.cell(row_index, 1).value for row_index in range(1, summary.max_row + 1)]
        self.assertIn("TOTAL SAMPLE VS. YUMOS", titles)
        self.assertIn("18-40 YO SAMPLE VS. YUMOS", titles)
        self.assertEqual(report["summary_scope"], "all")

    def test_question_union_handles_equal_priorities_across_split_exports(self):
        first = load_workbook(BytesIO(self._paired_example_bytes()))
        second = load_workbook(BytesIO(self._paired_example_bytes()))
        first_sheet = first.active
        second_sheet = second.active
        for row in range(8, 12):
            first_sheet.cell(row, 1).value = "Q-1-Overall product opinion"
            second_sheet.cell(row, 1).value = "Q-2-Overall product liking"
        first_buffer, second_buffer = BytesIO(), BytesIO()
        first.save(first_buffer)
        second.save(second_buffer)

        info = inspect_smart_package([
            ("split_a.xlsx", first_buffer.getvalue()),
            ("split_b.xlsx", second_buffer.getvalue()),
        ])

        self.assertEqual(
            {item.question_id for item in info.questions},
            {"Q-1-Overall product opinion", "Q-2-Overall product liking"},
        )

    def test_paired_mode_creates_one_delta_per_pair(self):
        raw = self._paired_example_bytes()
        info = inspect_smart_package([("paired.xlsx", raw)])
        self.assertTrue(info.suggested_test_type.startswith("Paired"))
        question_rows = [proposal_to_row(item) for item in info.questions]
        data, report = build_smart_toplines(
            [("paired.xlsx", raw)],
            question_rows,
            split_names=["TOTAL"],
            benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            test_type="Paired",
            mean_decimals=1,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        sheet = workbook["TOTAL"]
        self.assertEqual(sheet["E7"].value, "=D7-C7")
        self.assertEqual(sheet["I7"].value, "=H7-G7")
        self.assertEqual(sheet["C7"].number_format, "0.0")
        self.assertEqual(report["test_type"], "Paired")
        self.assertEqual(report["pairs"], 2)

    def test_paired_mode_can_create_a_kpi_summary(self):
        raw = self._paired_example_bytes()
        info = inspect_smart_package([("paired.xlsx", raw)])
        row = proposal_to_row(info.questions[0])
        row["KPI Summary"] = True
        data, report = build_smart_toplines(
            [("paired.xlsx", raw)],
            [row],
            split_names=["TOTAL"],
            benchmark_positions=(),
            standard_metrics=("Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"),
            include_screeners=False,
            test_type="Paired",
            summary_scope="total",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames[0], "KPI Summary")
        summary = workbook["KPI Summary"]
        self.assertEqual(summary["A1"].value, "TOTAL SAMPLE VS. PAIRED BENCHMARKS")
        self.assertIn("Candidate 1", summary["A5"].value)
        self.assertEqual(report["summary_sheets"], ["KPI Summary"])

    def test_paired_kpi_summary_uses_candidate_minus_paired_benchmark_direction(self):
        raw = self._paired_significant_example_bytes()
        info = inspect_smart_package([("paired_significant.xlsx", raw)])
        row = proposal_to_row(info.questions[0])
        row["KPI Summary"] = True
        data, _ = build_smart_toplines(
            [("paired_significant.xlsx", raw)],
            [row],
            split_names=["TOTAL"],
            benchmark_positions=(),
            standard_metrics=("Mean", "Bottom 2 Boxes"),
            include_screeners=False,
            test_type="Paired",
            summary_scope="total",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook["KPI Summary"]["B5"].value, "▲")

    def test_cmi_can_override_metrics_pair_direction_and_rendering(self):
        raw = self._paired_example_bytes()
        info = inspect_smart_package([("paired.xlsx", raw)])
        row = proposal_to_row(info.questions[0])
        row["Type"] = "Libre"
        row["Custom metrics"] = "Top Box; Bottom 2 Boxes"
        data, report = build_smart_toplines(
            [("paired.xlsx", raw)],
            [row],
            split_names=["RENDU CMI"],
            benchmark_positions=(),
            standard_metrics=(),
            include_screeners=False,
            test_type="Paired",
            paired_swaps={"paired.xlsx": [1]},
            include_deltas=False,
            include_sections=False,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        sheet = workbook["RENDU CMI"]
        self.assertEqual(sheet.max_column, 7)
        self.assertEqual((sheet["C6"].value, sheet["D6"].value), (0.3, 0.2))
        self.assertEqual(sheet["B6"].value, "Top Box")
        self.assertEqual(sheet["B7"].value, "Bottom 2 Boxes")
        self.assertFalse(report["include_deltas"])
        self.assertFalse(report["include_sections"])

    def test_real_einstein_paired_layout_is_detected(self):
        self.assertEqual(len(PAIRED_FILES), 4)
        workbook = load_workbook(PAIRED_FILES[0], read_only=True, data_only=False)
        sheet = detect_data_sheet(workbook)
        layout = detect_layout(sheet)
        active = _active_layout(sheet, layout)
        self.assertEqual(sheet.title, "Table_1_2_TAILED")
        self.assertEqual(
            (layout.header_row, layout.sample_row, layout.product_name_row, layout.product_header_row),
            (7, 5, 4, 6),
        )
        self.assertEqual(len(layout.product_cols), 18)
        self.assertEqual(len(active.product_cols), 6)

    def test_real_einstein_paired_accepts_variable_pair_counts(self):
        raw_files = [(path.name, path) for path in PAIRED_FILES]
        info = inspect_smart_package(raw_files)
        self.assertEqual(info.study_format, "HUT / in-use")
        self.assertTrue(info.suggested_test_type.startswith("Paired"))
        self.assertEqual(
            [item.split_name for item in info.inputs],
            [
                "ARIEL LOYALISTS MO SUFFERERS",
                "MALODOR SUFFERERS",
                "ARIEL LOYALISTS - NO MALODOR",
                "ARIEL LOYALISTS",
            ],
        )
        chosen = next(item for item in info.questions if item.question_id == "Q-02P-OVERALL PRODUCT OPINION")
        row = proposal_to_row(chosen)
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric == "Mean"
        data, report = build_smart_toplines(
            raw_files,
            [row],
            split_names=[item.split_name for item in info.inputs],
            benchmark_positions=(),
            standard_metrics=("Mean",),
            include_screeners=False,
            test_type="Paired",
            mean_decimals=2,
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        boost = workbook["ARIEL LOYALISTS MO SUFFERERS"]
        malodor = workbook["MALODOR SUFFERERS"]
        self.assertEqual(boost.max_column, 13)
        self.assertEqual(malodor.max_column, 25)
        self.assertEqual((boost["C7"].value, boost["D7"].value), (69.05, 67.86))
        self.assertEqual(boost["E7"].value, "=D7-C7")
        self.assertEqual(malodor["Y7"].value, "=X7-W7")
        self.assertEqual(report["pairs_by_split"]["MALODOR SUFFERERS"], 6)

    def test_real_einstein_paired_copies_direct_significance_fill(self):
        path = UPLOAD / "DataViz_Raw Data Einstein LC_2026-06-15 14_41_16.xlsx"
        info = inspect_smart_package([(path.name, path)])
        chosen = next(item for item in info.questions if item.question_id == "Q-02P-OVERALL PRODUCT OPINION")
        row = proposal_to_row(chosen)
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric == "Mean"
        data, _ = build_smart_toplines(
            [(path.name, path)],
            [row],
            split_names=["TOTAL"],
            benchmark_positions=(),
            standard_metrics=("Mean",),
            include_screeners=False,
            test_type="Paired",
        )
        sheet = load_workbook(BytesIO(data), data_only=False)["TOTAL"]
        self.assertEqual((sheet["G7"].value, sheet["H7"].value, sheet["I7"].value), (65.52, 59.48, "=H7-G7"))
        self.assertEqual((sheet["H7"].fill.fgColor.rgb or "")[-6:].upper(), "FF0000")
        self.assertEqual((sheet["I7"].fill.fgColor.rgb or "")[-6:].upper(), "FF0000")

    def test_real_consolidated_hut_builds_every_split_for_each_benchmark(self):
        info = inspect_smart_package([(WHITE_ALL_SPLITS.name, WHITE_ALL_SPLITS)])
        self.assertEqual(info.inputs[0].comparison_codes, ("E1L", "T2V"))
        row = proposal_to_row(info.questions[0])
        for metric in ("Mean", "Top Box", "Top 2 Boxes", "Top 3 Boxes", "Bottom Box", "Bottom 2 Boxes", "Bottom 3 Boxes"):
            row[metric] = metric == "Mean"
        data, report = build_smart_toplines(
            [(WHITE_ALL_SPLITS.name, WHITE_ALL_SPLITS)],
            [row],
            split_names=[info.inputs[0].split_name],
            benchmark_positions=(0, 1),
            standard_metrics=("Mean",),
            include_screeners=False,
            test_type="Monadic",
            benchmark_sheet_mode="auto_exports",
            benchmark_labels=("Comfort", "Fairy"),
            output_sheet_order="benchmark_first",
        )
        names = load_workbook(BytesIO(data), read_only=True).sheetnames
        self.assertEqual(len(names), 16)
        self.assertIn("TOTAL vs Comfort", names)
        self.assertIn("TOTAL vs Fairy", names)
        self.assertFalse(any("None" in name or "USE" in name or "REP vs" in name for name in names))
        self.assertEqual(report["missing_benchmark_exports"], [])

    def test_cmr_request_maps_yumos_products_to_clean_fantasy_names(self):
        matches = match_cmr_products(
            ("Y1M", "V2R", "D6A", "E4L", "C5O", "M3Q"),
            (
                "Yumos - <No Line> - Orchid - Türkiye - 2024 345494",
                "Vernel - Aromatherapy - Fresh Rose - Türkiye - 2026 357653",
                "MIX260004108 1.17% MIX260004108",
                "MIX260004063 1.22% MIX260004063",
                "MIX260004070 1.38% MIX260004070",
                "MIX260004057 1.38% MIX260004057",
            ),
            CMR_FILE,
        )
        self.assertEqual(
            [item.suggested_label for item in matches],
            [
                "Yumos Orchid",
                "Vernel Aromatherapy Fresh Rose",
                "ELLA G BLOOM",
                "MYSTIQUE G SOUK",
                "CHEERFUL G BLOOM",
                "CHEERFUL G MYSTICAL",
            ],
        )
        self.assertTrue(all(item.confidence == "Élevée" for item in matches))
        candidate = matches[2]
        self.assertEqual(cmr_product_label(candidate, "fantasy"), "ELLA G BLOOM")
        self.assertEqual(cmr_product_label(candidate, "formula"), "MIX260004108")
        self.assertEqual(cmr_product_label(candidate, "description"), "VSF040QVS@0.77 + VSI340AFF@0.4")
        self.assertEqual(cmr_product_label(candidate, "source"), "MIX260004108 1.17% MIX260004108")

    def test_cmr_matching_can_fall_back_to_formula_code(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Candidates & Bench"
        sheet.append(["CMR code", "Formula code", "Fantasy name", "Formula description"])
        sheet.append(["ABC", "MIX260009999", "STARLIGHT EB 2@0.9% + PLANETCAPS @0.2%", "FORMULA-X"])
        sheet.append(["DEF", "", "MOONLIGHT @0.8%", "ZAF708DZE@0.8 + CAPS@0.2"])
        buffer = BytesIO()
        workbook.save(buffer)
        match = match_cmr_products(("UNKNOWN",), ("MIX260009999 1.10%",), buffer.getvalue())[0]
        self.assertEqual(match.matched_by, "Formula code exact")
        self.assertEqual(match.suggested_label, "STARLIGHT")
        self.assertEqual(match.score, 98)
        description_match = match_cmr_products(
            ("UNKNOWN-2",), ("ZAF708DZE 0.80%",), buffer.getvalue()
        )[0]
        self.assertEqual(description_match.matched_by, "Code dans Formula description")
        self.assertEqual(description_match.suggested_label, "MOONLIGHT")
