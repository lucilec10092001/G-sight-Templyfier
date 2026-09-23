import unittest

from openpyxl import Workbook

from templyfier.core import TemplyfierError, detect_data_sheet, detect_layout


def result_sheet(title="Results", *, header_row=40, metric_col=20, header="Metric", with_bases=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.cell(header_row, metric_col, header)
    product_cols = (metric_col + 1, metric_col + 3)
    if with_bases:
        sample_row = header_row - 8
        sheet.cell(sample_row - 1, product_cols[0], "Candidate A")
        sheet.cell(sample_row - 1, product_cols[1], "Benchmark B")
        sheet.cell(sample_row, product_cols[0], "120 - A")
        sheet.cell(sample_row, product_cols[1], "118 - B")
    sheet.cell(header_row, product_cols[0], "A")
    sheet.cell(header_row, product_cols[1], "B")
    for offset, metric in enumerate(("Mean", "Top Box", "Top 2 Boxes"), 1):
        row = header_row + offset
        sheet.cell(row, metric_col - 1, "Q-1-Overall liking")
        sheet.cell(row, metric_col, metric)
        sheet.cell(row, product_cols[0], 6.2 + offset / 10)
        sheet.cell(row, product_cols[1], 5.8 + offset / 10)
        sheet.cell(row, product_cols[0] + 1, "B")
    return workbook, sheet


class AdaptiveLayoutTests(unittest.TestCase):
    def test_metric_table_can_move_beyond_legacy_scan_area(self):
        _, sheet = result_sheet(header_row=40, metric_col=20)
        layout = detect_layout(sheet)
        self.assertEqual((layout.header_row, layout.metric_col), (40, 20))
        self.assertEqual(layout.product_cols, (21, 23))
        self.assertEqual(layout.sample_row, 32)

    def test_metric_alias_and_export_without_base_row_are_supported(self):
        _, sheet = result_sheet(header_row=55, metric_col=25, header="Statistics", with_bases=False)
        layout = detect_layout(sheet)
        self.assertEqual((layout.header_row, layout.metric_col), (55, 25))
        self.assertEqual(layout.product_cols, (26, 28))
        self.assertEqual(layout.sample_row, 54)

    def test_malformed_priority_sheet_does_not_hide_valid_result_sheet(self):
        workbook = Workbook()
        malformed = workbook.active
        malformed.title = "Table_1 2_TAILED"
        malformed["B5"] = "Metric"
        malformed["C4"] = "100 - A"
        malformed["C5"] = "A"
        malformed["B6"] = "Mean"
        malformed["C6"] = 6.1

        _, valid = result_sheet(title="Unexpected result name", header_row=36, metric_col=18)
        copied = workbook.create_sheet(valid.title)
        for row in valid.iter_rows():
            for cell in row:
                if cell.value is not None:
                    copied.cell(cell.row, cell.column, cell.value)
        self.assertEqual(detect_data_sheet(workbook).title, "Unexpected result name")

    def test_narrative_metric_word_is_not_accepted_as_a_result_table(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Metric"
        sheet["A2"] = "This worksheet explains the metric definitions."
        with self.assertRaises(TemplyfierError):
            detect_layout(sheet)

    def test_unsupported_workbook_fails_closed_with_sheet_diagnostics(self):
        workbook = Workbook()
        workbook.active.title = "Overview"
        workbook.active["A1"] = "Project notes only"
        malformed = workbook.create_sheet("Possible results")
        malformed["B8"] = "Metric"
        malformed["B9"] = "Mean"
        malformed["C9"] = 6.4
        with self.assertRaises(TemplyfierError) as caught:
            detect_data_sheet(workbook)
        message = str(caught.exception)
        self.assertIn("No output was generated", message)
        self.assertIn("Possible results", message)


if __name__ == "__main__":
    unittest.main()
