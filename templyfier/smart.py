from __future__ import annotations

from copy import copy
from dataclasses import asdict, dataclass, replace
from io import BytesIO
import json
from pathlib import Path
import re
import unicodedata
from typing import BinaryIO, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import (
    TemplyfierError,
    _count_from_header,
    _default_split_name,
    _product_names,
    _read_bytes,
    _safe_sheet_name,
    _signature,
    _source_significance_col,
    _text,
    _workbook,
    detect_data_sheet,
    detect_layout,
)


STANDARD_METRICS = (
    "Mean",
    "Top Box",
    "Top 2 Boxes",
    "Top 3 Boxes",
    "Bottom Box",
    "Bottom 2 Boxes",
    "Bottom 3 Boxes",
)


@dataclass(frozen=True)
class QuestionProposal:
    keep: bool
    order: int
    section: str
    question_id: str
    mapped_kpi: str
    display_label: str
    question_type: str
    confidence: str
    first_row: int
    metrics: tuple[str, ...]
    metric_label: str = ""
    availability: tuple[str, ...] = ()
    suggested_splits: tuple[str, ...] = ()
    summary_keep: bool = False
    summary_label: str = ""
    metric_availability: tuple[tuple[str, tuple[str, ...]], ...] = ()
    stages: tuple[str, ...] = ()


@dataclass(frozen=True)
class SmartPackageInfo:
    reference_filename: str
    source_sheet: str
    product_names: tuple[str, ...]
    counts: tuple[int | None, ...]
    questions: tuple[QuestionProposal, ...]
    inputs: tuple["SmartInputInfo", ...]
    study_format: str
    stages: tuple[str, ...]
    embedded_splits: tuple[str, ...]
    warnings: tuple[str, ...]
    suggested_benchmark_count: int
    suggested_test_type: str


@dataclass(frozen=True)
class SmartInputInfo:
    filename: str
    split_name: str
    counts: tuple[int | None, ...]
    product_names: tuple[str, ...]
    role: str = "Résultats"
    stage: str = ""
    embedded_split_count: int = 1
    comparison_codes: tuple[str, ...] = ()
    split_detection: str = "High confidence"
    product_keys: tuple[str, ...] = ()


def audit_input_plan(
    info: SmartPackageInfo,
    split_names: Sequence[str],
    benchmark_positions: Sequence[int],
    *,
    test_type: str,
    automatic_exports: bool,
) -> dict:
    """Reconcile files, splits, benchmarks and product plans before export."""
    inputs = [item for item in info.inputs if item.role == "Résultats"] or list(info.inputs)
    names = [_text(value) for value in split_names]
    blockers: list[str] = []
    warnings: list[str] = []
    rows: list[dict] = []
    if len(names) != len(inputs):
        blockers.append("Le nombre de noms de splits ne correspond pas au nombre d’exports de résultats.")
        names = [item.split_name for item in inputs]
    if not inputs:
        return {"rows": rows, "blockers": ["Aucun export de résultats détecté."], "warnings": [], "ready": False}

    canonical = inputs[0].product_keys
    paired = _normal(test_type).startswith("paired")
    for item in inputs:
        same_plan = (
            len(item.product_keys) == len(canonical)
            and (set(item.product_keys) == set(canonical) if automatic_exports and not paired else item.product_keys == canonical)
        )
        if not same_plan:
            blockers.append(f"Plan produits différent dans {item.filename}.")

    if paired:
        for item, split in zip(inputs, names):
            valid = len(item.product_keys) >= 2 and len(item.product_keys) % 2 == 0
            if not valid:
                blockers.append(f"{split} ne contient pas un nombre pair de produits actifs.")
            rows.append({
                "Split": split,
                "Benchmark": "Paires internes",
                "Export": item.filename,
                "Produits": len(item.product_keys),
                "Contrôle": "Prêt" if valid else "À corriger",
            })
    elif automatic_exports:
        benchmark_keys = [canonical[position] for position in benchmark_positions if 0 <= position < len(canonical)]
        if not benchmark_keys:
            blockers.append("Aucun benchmark valide n’a été identifié.")
        split_order = list(dict.fromkeys(names))
        for split in split_order:
            indexes = [index for index, value in enumerate(names) if _normal(value) == _normal(split)]
            for benchmark_key in benchmark_keys:
                matches = [index for index in indexes if benchmark_key in inputs[index].comparison_codes]
                if len(matches) == 1:
                    status = "Prêt"
                    export_name = inputs[matches[0]].filename
                elif not matches:
                    status = "Export manquant"
                    export_name = "—"
                    blockers.append(f"Export manquant : {split} vs {benchmark_key}.")
                else:
                    status = "Doublon"
                    export_name = " · ".join(inputs[index].filename for index in matches)
                    blockers.append(f"Plusieurs exports détectés pour {split} vs {benchmark_key}.")
                rows.append({
                    "Split": split,
                    "Benchmark": benchmark_key,
                    "Export": export_name,
                    "Produits": len(canonical),
                    "Contrôle": status,
                })
        unused = [
            item.filename for item in inputs
            if not any(code in {canonical[position] for position in benchmark_positions if 0 <= position < len(canonical)}
                       for code in item.comparison_codes)
        ]
        if unused:
            warnings.append("Benchmark non rapproché dans : " + " · ".join(unused))
    else:
        duplicate_names = {
            name for name in names
            if sum(_normal(other) == _normal(name) for other in names) > 1
        }
        if duplicate_names:
            blockers.append("Noms de splits en double : " + " · ".join(sorted(duplicate_names)))
        benchmark_keys = [canonical[position] for position in benchmark_positions if 0 <= position < len(canonical)]
        for item, split in zip(inputs, names):
            missing = [key for key in benchmark_keys if key not in item.comparison_codes]
            if missing:
                warnings.append(f"{split} : significativité lue via les lettres G-Sight pour " + ", ".join(missing))
            rows.append({
                "Split": split,
                "Benchmark": " · ".join(benchmark_keys) or "—",
                "Export": item.filename,
                "Produits": len(item.product_keys),
                "Contrôle": "Prêt",
            })
    return {
        "rows": rows,
        "blockers": list(dict.fromkeys(blockers)),
        "warnings": list(dict.fromkeys(warnings)),
        "ready": not blockers,
    }


@dataclass(frozen=True)
class CmrProduct:
    cmr_code: str
    formula_code: str
    fantasy_name: str
    formula_description: str
    formula_type: str = ""
    fr_land_id: str = ""
    brand: str = ""
    line_variant: str = ""


@dataclass(frozen=True)
class CmrProductMatch:
    product_key: str
    source_name: str
    cmr_code: str = ""
    formula_code: str = ""
    formula_description: str = ""
    fantasy_name: str = ""
    suggested_label: str = ""
    matched_by: str = "Aucun rapprochement fiable"
    confidence: str = "—"
    score: int = 0


def cmr_product_label(match: CmrProductMatch, mode: str = "fantasy") -> str:
    """Return the CMI-selected CMR field, with safe fallbacks for missing cells."""
    mode = _text(mode).casefold()
    if mode == "source":
        return match.source_name
    if mode == "formula":
        return _text(match.formula_code) or _text(match.formula_description) or match.suggested_label or match.source_name
    if mode == "description":
        description = re.sub(r"\s+", " ", _text(match.formula_description).replace("\n", " ")).strip()
        return description or _text(match.formula_code) or match.suggested_label or match.source_name
    return match.suggested_label or _text(match.fantasy_name) or _text(match.formula_description) or match.source_name


def _cmr_header_key(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _normal(_text(value))).strip()


def _clean_fantasy_name(value: str) -> str:
    """Turn a verbose CMR fantasy-name cell into a concise, editable suggestion."""
    clean = re.sub(r"\s+", " ", _text(value).replace("\n", " ")).strip(" +")
    if not clean:
        return ""
    parts = [part.strip() for part in re.split(r"\s*\+\s*", clean) if part.strip()]
    preferred = next(
        (part for part in parts if not re.search(r"caps|encaps|planetcaps|techno", part, flags=re.I)),
        parts[0],
    )
    preferred = re.sub(r"\s*@\s*\d+(?:[.,]\d+)?\s*%?.*$", "", preferred).strip()
    preferred = re.sub(r"\s+EB\s*[-_.]?\s*\d+\s*$", "", preferred, flags=re.I).strip()
    return preferred or clean


def _cmr_benchmark_label(product: CmrProduct) -> str:
    parts = [product.brand, product.line_variant]
    clean_parts = []
    for part in parts:
        clean = re.sub(r"<\s*No Line\s*>\s*-?", "", _text(part), flags=re.I)
        clean = re.sub(r"\s+-\s+", " ", clean).strip(" -")
        if clean and _normal(clean) not in {_normal(item) for item in clean_parts}:
            clean_parts.append(clean)
    if clean_parts:
        return " ".join(clean_parts)
    description = _text(product.formula_description)
    description = re.sub(r"\s+-\s+(?:Fs\s+)?(?:Liquid|Powder).*$", "", description, flags=re.I)
    return description.strip(" -")


def parse_cmr_products(source) -> tuple[CmrProduct, ...]:
    """Read the candidate catalogue from a CMR request with flexible header locations."""
    workbook = _workbook(source, data_only=True)
    products = []
    aliases = {
        "cmr code": "cmr_code",
        "cmr codes": "cmr_code",
        "formula code": "formula_code",
        "fantasy name": "fantasy_name",
        "formula description": "formula_description",
        "formula type": "formula_type",
        "fr land id": "fr_land_id",
        "brand": "brand",
        "line variant": "line_variant",
    }
    for sheet in workbook.worksheets:
        header_row = None
        header_map = {}
        for row in range(1, min(sheet.max_row, 30) + 1):
            candidate = {}
            for col in range(1, sheet.max_column + 1):
                key = aliases.get(_cmr_header_key(sheet.cell(row, col).value))
                if key:
                    candidate[key] = col
            if "cmr_code" in candidate and ({"fantasy_name", "formula_description"} & set(candidate)):
                header_row, header_map = row, candidate
                break
        if header_row is None:
            continue
        for row in range(header_row + 1, sheet.max_row + 1):
            values = {
                field: _text(sheet.cell(row, col).value)
                for field, col in header_map.items()
            }
            if not any(values.get(field) for field in ("cmr_code", "formula_code", "fantasy_name", "formula_description")):
                continue
            products.append(CmrProduct(**{
                field: values.get(field, "")
                for field in CmrProduct.__dataclass_fields__
            }))
    if not products:
        raise TemplyfierError(
            "La CMR ne contient pas de tableau avec les colonnes CMR code et Fantasy name / Formula description."
        )
    return tuple(products)


def _cmr_match_score(product_key: str, source_name: str, product: CmrProduct) -> tuple[int, str]:
    key = re.sub(r"[^A-Z0-9]", "", _text(product_key).upper())
    source = _normal(source_name)
    compact_source = re.sub(r"[^a-z0-9]", "", source)
    cmr_code = re.sub(r"[^A-Z0-9]", "", product.cmr_code.upper())
    formula_code = re.sub(r"[^A-Z0-9]", "", product.formula_code.upper())
    if key and cmr_code and key == cmr_code:
        return 100, "Code CMR exact"
    if formula_code and formula_code.lower() in compact_source:
        return 98, "Formula code exact"
    fr_land_digits = re.sub(r"\D", "", product.fr_land_id)
    if len(fr_land_digits) >= 4 and fr_land_digits in re.sub(r"\D", "", source_name):
        return 95, "Fr-Land ID exact"
    description_codes = {
        re.sub(r"[^A-Z0-9]", "", token)
        for token in re.findall(r"\b[A-Z][A-Z0-9-]{5,}\b", product.formula_description.upper())
    }
    if any(code.lower() in compact_source for code in description_codes if len(code) >= 6):
        return 94, "Code dans Formula description"
    description = _normal(product.formula_description)
    if description and (description in source or source in description):
        return 92, "Formula description"
    fantasy = _normal(_clean_fantasy_name(product.fantasy_name))
    fantasy_words = {word for word in fantasy.split() if len(word) >= 4}
    if fantasy_words:
        overlap = len(fantasy_words & set(source.split())) / len(fantasy_words)
        if overlap >= 0.75:
            return 84, "Fantasy name"
    brand_words = {
        word for word in _normal(f"{product.brand} {product.line_variant}").split()
        if len(word) >= 4 and word not in {"line", "variant"}
    }
    if brand_words and len(brand_words & set(source.split())) / len(brand_words) >= 0.6:
        return 80, "Marque / variante benchmark"
    return 0, "Aucun rapprochement fiable"


def match_cmr_products(
    product_keys: Sequence[str],
    product_names: Sequence[str],
    cmr_source,
) -> tuple[CmrProductMatch, ...]:
    """Create one-to-one CMR suggestions, prioritising stable technical identifiers."""
    catalog = parse_cmr_products(cmr_source)
    edges = []
    for product_index, (key, name) in enumerate(zip(product_keys, product_names)):
        for cmr_index, product in enumerate(catalog):
            score, matched_by = _cmr_match_score(key, name, product)
            if score:
                edges.append((score, product_index, cmr_index, matched_by))
    assigned_products = {}
    used_cmr = set()
    for score, product_index, cmr_index, matched_by in sorted(edges, reverse=True):
        if product_index in assigned_products or cmr_index in used_cmr:
            continue
        assigned_products[product_index] = (catalog[cmr_index], score, matched_by)
        used_cmr.add(cmr_index)

    matches = []
    for product_index, (key, name) in enumerate(zip(product_keys, product_names)):
        assignment = assigned_products.get(product_index)
        if not assignment:
            matches.append(CmrProductMatch(product_key=key, source_name=name))
            continue
        product, score, matched_by = assignment
        suggestion = _clean_fantasy_name(product.fantasy_name) or _cmr_benchmark_label(product)
        confidence = "Élevée" if score >= 95 else "Moyenne" if score >= 80 else "Faible"
        matches.append(CmrProductMatch(
            product_key=key,
            source_name=name,
            cmr_code=product.cmr_code,
            formula_code=product.formula_code,
            formula_description=product.formula_description,
            fantasy_name=product.fantasy_name,
            suggested_label=suggestion,
            matched_by=matched_by,
            confidence=confidence,
            score=score,
        ))
    return tuple(matches)


def _stage_name(sheet) -> str:
    # G-Sight usually writes Stage in A1, but translated exports and workbook
    # conversions can move it within the header. Prefer an explicit label and
    # keep the historical A1 fallback for older exports.
    for row in range(1, min(sheet.max_row, 20) + 1):
        for col in range(1, min(sheet.max_column, 20) + 1):
            value = _text(sheet.cell(row, col).value).strip()
            match = re.match(r"^(?:Stage|Stade)\s*:\s*(.+)$", value, flags=re.I)
            if match and match.group(1).strip():
                return match.group(1).strip()
    value = _text(sheet.cell(1, 1).value).strip()
    return value or "Non précisé"


def _embedded_data_sheets(workbook):
    sheets = []
    for sheet in workbook.worksheets:
        if re.match(r"^Table_\d+(?:\s+|_)2_TAILED$", sheet.title, flags=re.I):
            try:
                detect_layout(sheet)
                sheets.append(sheet)
            except TemplyfierError:
                pass
    if sheets:
        return sheets
    for sheet in workbook.worksheets:
        if any(word in sheet.title.casefold() for word in ("legend", "screener")):
            continue
        try:
            detect_layout(sheet)
            sheets.append(sheet)
        except TemplyfierError:
            pass
    return sheets or [detect_data_sheet(workbook)]


def _active_layout(sheet, layout):
    """CLT exports often contain paired N/W columns; only one is active per stage."""
    active = tuple(
        col for col in layout.product_cols
        if (_count_from_header(sheet.cell(layout.sample_row, col).value) or 0) > 0
    )
    if len(active) >= 2 and len(active) < len(layout.product_cols):
        return type(layout)(
            layout.header_row,
            layout.metric_col,
            active,
            sample_row_index=layout.sample_row,
            product_name_row_index=layout.product_name_row,
            product_header_row_index=layout.product_header_row,
        )
    return layout


def _split_label(sheet) -> str:
    # Most G-Sight exports use A2, but layouts, translations and Excel
    # conversions can move the filter summary within the header area.
    prefix = r"(?:Search|Recherche|Filter(?:s)?|Filtre(?:s)?)"
    values=[]
    for row in range(1,min(sheet.max_row,20)+1):
        for col in range(1,min(sheet.max_column,20)+1):
            value=_text(sheet.cell(row,col).value).strip()
            if re.match(rf"^{prefix}\s*:",value,flags=re.I):
                values.append(value)
    if not values:return "TOTAL"
    stripped = [re.sub(rf"^{prefix}\s*:\s*", "", value, flags=re.I).strip() for value in values]
    stripped = list(dict.fromkeys(value for value in stripped if value))
    if not stripped:return "TOTAL"
    if len(stripped) == 1:return stripped[0]
    # A converted workbook can repeat a short filter beside its complete
    # summary. Keep the complete summary. If the filters are genuinely
    # separate, combine them so no split dimension is silently discarded.
    longest = max(stripped, key=len)
    if all(_normal(value) in _normal(longest) for value in stripped):
        return longest
    return "; Search: ".join(stripped)


def _has_split_filter(sheet) -> bool:
    prefix = r"(?:Search|Recherche|Filter(?:s)?|Filtre(?:s)?)"
    return any(re.match(rf"^{prefix}\s*:",_text(sheet.cell(row,col).value).strip(),flags=re.I)
        for row in range(1,min(sheet.max_row,20)+1)
        for col in range(1,min(sheet.max_column,20)+1))


def _friendly_split_label(sheet) -> str:
    search = _normal(_split_label(sheet))
    if "sample split" in search and "boost" in search and any(
        word in search for word in ("orkide", "orchid")
    ):
        return "YUMOS ORKIDE MO"
    if "city" in search and "istanbul" in search:
        return "ISTANBUL"
    if "city" in search and "ankara" in search:
        return "ANKARA"
    if "fabcon brand mo" in search and "yumos" in search:
        return "YUMOS MO"
    if "fabcon brand mo" in search:
        return "OTHER BRANDS"
    if "sample type" in search and "boost" in search and "ariel loyalist" not in search.split("boost", 1)[0]:
        return "ARIEL LOYALISTS MO SUFFERERS"
    if "sample type" in search and "ariel loyalist" in search and "boost" in search:
        return "MALODOR SUFFERERS"
    if "sample type" in search and "ariel loyalist" in search:
        if "malodor recode: no" in search or "malodor recode no" in search:
            return "ARIEL LOYALISTS - NO MALODOR"
        return "ARIEL LOYALISTS"
    if "rep or boost: boost" in search or "rep or boost boost" in search:
        return "BOOST"
    if ("rep or boost: rep" in search or "rep or boost rep" in search) and ", search:" not in search:
        return "TOTAL"
    if "brand fs mo: comfort" in search or "brand fs mo comfort" in search:
        return "COMFORT MO"
    if "brand fs mo:" in search or "brand fs mo" in search:
        return "OTHER BRANDS MO"
    mo_brands = re.search(r"\bmo brands?\s*:?\s*(.+)", search)
    if mo_brands:
        brands = [part.strip() for part in mo_brands.group(1).split(",") if part.strip()]
        brand_set = set(brands)
        retailer_brands = {
            "aldi almat", "asda", "lidl formil", "marks & spencer (m&s)",
            "morrisons", "sainsbury's", "tesco", "waitrose & partners",
        }
        if brand_set and brand_set.issubset(retailer_brands):
            return "RETAILER BRANDS MO"
        if len(brands) <= 2 and brands[0] in {"surf", "bold", "ariel"}:
            return f"{brands[0].upper()} MO"
        other = any(value.startswith("other brand") for value in brands)
        if other and len(brands) >= 20:
            return "OTHER BRANDS MO"
        if other:
            return "NATIONAL BRANDS MO"
        return f"{brands[0].upper()} MO"
    mo_brand = re.search(r"\bmo brand\s*:?\s*([^,]+)", search)
    if mo_brand:
        return f"{mo_brand.group(1).strip().upper()} MO"
    if "colour mo: white" in search or "colour mo white" in search:
        return "WHITE MO"
    if "colour mo:" in search or "colour mo" in search:
        return "OTHER COLORS"
    if "age" in search and any(token in search for token in ("under 18", "18 30", "31 40", "18 40")):
        return "18-40 YO"
    if "age" in search and any(token in search for token in ("41 50", "51 64", "41 64")):
        return "41+ YO" if "65 years" in search else "41-64 YO"
    if "age" in search and "65 years" in search:
        return "41-65 YO"
    if "sample split" in search and re.search(r"\bmain\b", search):
        return "TOTAL"
    return "TOTAL"


def _generic_split_label(sheet) -> str:
    """Extract arbitrary G-Sight search values when no known CMI split matches."""
    raw = _split_label(sheet)
    if not raw or _normal(raw) == "total":
        return "TOTAL"
    parts = re.split(r"(?:,|;|\r?\n)\s*(?:Search|Recherche|Filter(?:s)?|Filtre(?:s)?)\s*:\s*", raw, flags=re.I)
    values = []
    for part in parts:
        if ":" not in part:
            continue
        value = part.split(":", 1)[1].strip()
        value = re.sub(r"^\d+\s*[-.)]\s*", "", value).strip()
        if not value or _normal(value) in {"all", "total", "total population"}:
            continue
        clean = re.sub(r"\s+", " ", value).upper()
        if clean not in values:
            values.append(clean)
    return " · ".join(values) or "TOTAL"


def _split_detection(sheet, filename: str) -> tuple[str,str]:
    """Return a conservative label and visible confidence for CMI review."""
    has_filter=_has_split_filter(sheet)
    raw=_split_label(sheet)
    detected=_detected_split_label(sheet)
    normalized=_normal(raw)
    explicit_total=(normalized in {'','total','all','total population'}
        or ('sample split' in normalized and bool(re.search(r'\bmain\b',normalized)))
        or 'rep or boost rep' in normalized)
    if detected!='TOTAL':
        return detected,'High confidence'
    if explicit_total:
        return 'TOTAL','High confidence'
    if not has_filter and _normal(Path(filename).stem) in {'total','total sample'}:
        return 'TOTAL','High confidence'
    # A non-empty filter that cannot be parsed must never silently become TOTAL.
    return _default_split_name(filename),'Review recommended'


def _detected_split_label(sheet) -> str:
    friendly = _friendly_split_label(sheet)
    search = _normal(_split_label(sheet))
    if friendly == "TOTAL":
        is_main = "sample split" in search and re.search(r"\bmain\b", search)
        is_plain_rep = (
            ("rep or boost: rep" in search or "rep or boost rep" in search)
            and (", search:" not in search or bool(re.search(r", search: [^:]+:\s*$", search)))
        )
        if is_main or is_plain_rep:
            return "TOTAL"
    return friendly if friendly != "TOTAL" else _generic_split_label(sheet)


def _comparison_codes(sheet, layout) -> tuple[str, ...]:
    product_codes = set()
    for col in layout.product_cols:
        match = re.search(r"-\s*([A-Z0-9]+)\s*$", _text(sheet.cell(layout.sample_row, col).value), flags=re.I)
        if match:
            product_codes.add(match.group(1).upper())
    found = []
    for col in range(layout.metric_col + 1, sheet.max_column + 1):
        value = _text(sheet.cell(layout.product_header_row, col).value).upper()
        if value in product_codes and value not in found:
            found.append(value)
    return tuple(found)


def _suggest_test_type(inputs: Sequence["SmartInputInfo"]) -> str:
    if any(item.comparison_codes for item in inputs):
        return "Monadic"
    product_counts = {len(item.product_keys) for item in inputs if item.product_keys}
    if product_counts and all(count % 2 == 0 for count in product_counts):
        return "Paired (à confirmer)"
    return "Monadic (à confirmer)"


def _suggest_benchmark_count(inputs: Sequence["SmartInputInfo"]) -> int:
    result_inputs = [item for item in inputs if item.role == "Résultats"] or list(inputs)
    explicit = max((len(item.comparison_codes) for item in result_inputs), default=1) or 1
    names = result_inputs[0].product_names if result_inputs else ()
    # A frequent G-Sight plan starts with two named market controls followed by
    # coded prototypes. Treat this as a suggestion only; the CMI still confirms it.
    if len(names) >= 4:
        coded = [bool(re.search(r"\b(?:MIX|TEST|PROTO)[-_ ]?\d", name, flags=re.I)) for name in names]
        if not coded[0] and not coded[1] and sum(coded[2:]) >= max(2, len(coded[2:]) - 1):
            explicit = max(explicit, 2)
    return max(1, explicit)


def _product_keys(sheet, layout) -> tuple[str, ...]:
    keys = []
    paired_stage = _normal(_stage_name(sheet)) in {"neat", "damp wet", "wet"}
    for position, col in enumerate(layout.product_cols):
        sample = _text(sheet.cell(layout.sample_row, col).value)
        match = re.search(r"-\s*([A-Z0-9]+)\s*$", sample, flags=re.I)
        if match:
            code = match.group(1).upper()
            if paired_stage and code.endswith(("N", "W")):
                code = code[:-1]
            keys.append(code)
        else:
            keys.append(_normal(_product_names(sheet, layout)[position]))
    return tuple(keys)


def _study_format(filenames: Sequence[str], stages: Sequence[str], max_rows: int = 0) -> str:
    haystack = " ".join(filenames).casefold()
    if "clt" in haystack:
        return "CLT"
    if "hut" in haystack or any("in use" in stage.casefold() for stage in stages):
        return "HUT / in-use"
    normalized_stages = {_normal(stage) for stage in stages}
    if normalized_stages and normalized_stages.issubset({"neat", "damp wet", "wet"}):
        return "CLT"
    if max_rows >= 700 and normalized_stages.issubset({"none", "use", "non precise"}):
        return "HUT / in-use"
    return "À confirmer"


from .language import analysis_normal as _analysis_normal


def _normal(text: str) -> str:
    text = unicodedata.normalize("NFKD", _text(text)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text.casefold().replace("_", " ").replace("-", " ")).strip()


def _metric_key(metric: str) -> str:
    from .language import french_box_key
    norm = _normal(metric)
    french_box = french_box_key(norm)
    if french_box:
        return french_box
    if norm in {"mean", "moyenne", "promedio", "media", "mittelwert"}:
        return "mean"
    # A numbered answer mentioning 'top', 'bottom', 'haut' or 'bas' is
    # still an individual answer, never an aggregate box.
    if re.match(r"^\d+\s", norm):
        return norm
    direction = "top" if any(word in norm for word in ("top", "haut", "superieur")) else "bottom" if any(word in norm for word in ("bottom", "bas", "inferieur")) else ""
    if direction:
        number = "3" if re.search(r"\b3\b", norm) else "2" if re.search(r"\b2\b", norm) else "1"
        return f"{direction}_{number}"
    return norm


def _looks_like_code(value: str) -> bool:
    return bool(re.match(r"^q\s*[-_]?\s*\d+", _text(value), flags=re.I))


def _is_maxdiff_question(question_id: str, label: str = "", mapped_kpi: str = "") -> bool:
    combined = _normal(f"{question_id} | {label} | {mapped_kpi}")
    return (
        "max diff" in combined
        or "maxdiff" in combined
        or _normal(label) == "favorite"
        or bool(re.search(r"\bq\s*0?8\s+favorite\b", combined))
    )


def humanize_question(question_id: str, mapped_kpi: str = "") -> str:
    """Remove only the technical Q-number prefix and preserve the questionnaire wording."""
    label = re.sub(
        r"^Q\s*[-_]?\s*\d+[A-Z]?(?:\s*[-_]\s*\d+[A-Z]?)*\s*[-_:=]?\s*",
        "",
        _text(question_id),
        flags=re.I,
    )
    label = re.sub(r"_+", " ", label)
    label = re.sub(r"\s+", " ", label).strip(" -_:=")
    normalized_label = _normal(label)
    if normalized_label in {"f opo", "ofo", "ofl", "ofr", "overall fragrance rating"}:
        return "Overall Fragrance Opinion"
    if normalized_label in {"opo", "overall product liking"}:
        return "Overall Product Opinion"
    if normalized_label == "pi":
        return "Purchase Intent"
    if normalized_label == "purchase":
        return "Purchase Intent"
    if normalized_label == "expectations":
        return "Evaluation vs. Expectations"
    if normalized_label in {"softness", "overall softness"}:
        return "Overall Softness Opinion"
    label = re.sub(r"^F(?:\s*[-_]\s*\d+)?\s*[-_:]\s*", "", label, flags=re.I)
    label = re.sub(
        r"^(?:Statements?|Product benefit|Fragrance statement|Opinion by stage|Strength by stage|Perception Mapping)\s*[-_:]\s*",
        "",
        label,
        flags=re.I,
    )
    if label:
        return label
    if mapped_kpi and not _looks_like_code(mapped_kpi):
        return re.sub(r"\s+", " ", mapped_kpi.replace("_", " ")).strip()
    return _text(question_id)


def _color_code(question_id: str, label: str = "") -> str:
    for value in (question_id, label):
        match = re.search(r"(?:colour|color)[\s_-]*([A-Z]\d+)\s*$", value, flags=re.I)
        if match:
            return match.group(1).upper()
    return ""


def _suggest_section(question_id: str, label: str, question_type: str, mapped_kpi: str = "") -> str:
    combined = _analysis_normal(f"{question_id} | {label} | {mapped_kpi}")
    qid = _analysis_normal(question_id)
    if _is_maxdiff_question(question_id, label, mapped_kpi):
        return "HORS TOPLINES"
    if question_type == "Delete" and re.search(r"\b(?:other\s*\(?\s*precise|autre\s+preciser)\b", combined):
        return "OPEN ENDS — HORS TOPLINES"
    if _color_code(question_id, label):
        return "FRAGRANCE CHARACTERISTICS"
    if question_type == "Attribute" and re.search(r"\bbenefits?\b", qid):
        return "PRODUCT BENEFITS"
    # Attribute batteries must be interpreted as a whole. An item called
    # "Long-lasting" inside "Olfactive attributes" is a perception item, not
    # a standalone fragrance-benefit KPI.
    if question_type == "Attribute" and any(word in combined for word in (
        "olfactive attribute", "olfactory attribute", "freshness type", "type of freshness",
        "consumer perception", "perception mapping", "characteristic", "emotion", "feeling", " attributes ",
    )):
        return "FRAGRANCE CHARACTERISTICS"
    if any(word in combined for word in (
        "quantity used", "dosing", "wash load", "wash temp", "wash cycle", "dry method",
        "detergent format", "detergent brand", "scent booster used", "scent booster brand",
        "scent booster variant", "my own laundry", "laundry of my", "usage", "product tested",
        "per wash", "amount of product", "comparison amount",
    )) or re.match(r"q\s*34(?:\s|$)", qid):
        return "USAGE & CONTEXT"
    if any(word in combined for word in ("emotion", "good mood", "relaxed", "reassur", "comforting")):
        return "FRAGRANCE CHARACTERISTICS"
    if "scent experience overall" in combined:
        return "PRODUCT EVALUATION"
    if "expectation" in combined:
        return "PRODUCT EVALUATION"
    if _analysis_normal(label) in {"softness", "overall softness", "overall softness opinion"}:
        return "PRODUCT EVALUATION"
    if "fit to brand" in combined:
        return "PRODUCT EVALUATION"
    if any(word in combined for word in (
        "bottle suitability", "variant suitability", "universe suitability", "comfort pure suitability",
    )):
        return "FRAGRANCE BENEFITS"
    if "fit to" in combined or re.search(r"\bfit\b", combined):
        return "FIT"
    if re.match(r"q\s*11(?:\s|$)", qid) and any(
        word in combined for word in ("odor", "odour", "smell", "fresh", "hygien")
    ):
        return "PRODUCT BENEFITS"
    if re.match(r"q\s*13(?:\s|$)", qid):
        return "FRAGRANCE BENEFITS"
    if re.match(r"q\s*14(?:\s+[a-z])?\s+\d+(?:\s|$)", qid) and any(
        word in combined for word in ("opinion by stage", "opening", "machine", "drying", "wearing", "wardrobe", "rubbing")
    ):
        return "FRAGRANCE JOURNEY - LIKING"
    if any(word in combined for word in ("overall fragrance", "overall scent", "f opo", "ofo", "ofl", "ofr")):
        return "FRAGRANCE EVALUATION"
    if re.match(r"q\s*18\s+f(?:\s|$)", qid) and "perception mapping" in combined:
        return "FRAGRANCE CHARACTERISTICS"
    if question_type == "Strength" or any(word in combined for word in ("intensity", "strength", "too weak", "too strong")):
        if re.match(r"q\s*1[568](?:\s+[a-z])?\s+\d+(?:\s|$)", qid) or any(
            word in combined for word in ("opening", "machine", "drying", "folding", "wearing", "wardrobe", "storage", "rubbing")
        ):
            return "FRAGRANCE JOURNEY - STRENGTH"
        return "FRAGRANCE STRENGTH"
    if re.match(r"q\s*18(?:\s|$)", qid) and question_type == "Attribute":
        return "FRAGRANCE CHARACTERISTICS"
    if re.match(r"q\s*19(?:\s|$)", qid) or "characteristic" in combined:
        return "FRAGRANCE CHARACTERISTICS"
    if re.match(r"q\s*20(?:\s|$)", qid):
        return "FRAGRANCE CHARACTERISTICS"
    if any(word in combined for word in ("fragrance benefit", "scent benefit", "long lasting freshness", "long lastingness", "smell great")):
        return "FRAGRANCE BENEFITS"
    if combined.strip().endswith("long lasting"):
        return "FRAGRANCE BENEFITS"
    if (
        (re.match(r"q\s*1[45](?:\s+[a-z])?(?:\s+\d+)?(?:\s|$)", qid) and question_type == "Standard")
        or any(word in combined for word in (
            "when opening", "opening the bottle", "washing machine", "removing wet clothes",
            "while drying", "hanging clothes", "folding", "when wearing", "wardrobe",
        ))
    ):
        return "FRAGRANCE JOURNEY - LIKING"
    if re.match(r"q\s*24\s+f(?:\s|$)", qid) and "comparison to current" in combined:
        return "FRAGRANCE EVALUATION"
    if re.match(r"q\s*10(?:\s|$)", qid) and "comparison to current" in combined:
        return "PRODUCT EVALUATION"
    if any(word in combined for word in (
        "product benefit", "cleaning", "stain", "whiteness", "residue", "laundry result", "safe on",
        "malodor", "malodour", "bad odor", "bad odour", "hygienically clean",
    )):
        return "PRODUCT BENEFITS"
    if re.match(r"q\s*[789](?:\s|$)", qid):
        return "PRODUCT BENEFITS"
    if question_type == "Attribute":
        if any(word in combined for word in (
            "benefit", "perception", "fresh", "fraich", "parfum", "fragrance", "scent", "odeur", "smell",
        )):
            return "FRAGRANCE CHARACTERISTICS"
        return "ATTRIBUTES"
    if any(word in combined for word in ("parfum", "fragrance", "scent", "odeur", "smell")):
        return "FRAGRANCE EVALUATION"
    return "PRODUCT EVALUATION"


def _question_priority(question_id: str, label: str, section: str, first_row: int) -> tuple[int, int]:
    combined = _analysis_normal(f"{question_id} | {label}")
    qid = _analysis_normal(question_id)
    if any(word in combined for word in ("overall fragrance", "overall scent", "f opo", "ofo", "ofr")):
        rank = 220
    elif any(word in combined for word in ("overall product opinion", "overall product liking", "opo")):
        rank = 100
    elif re.match(r"q\s*8(?:\s|$)", qid) and "expectation" in combined:
        rank = 105
    elif "fit to brand" in combined:
        rank = 108
    elif section == "PRODUCT EVALUATION" and any(word in combined for word in ("purchase intent", "purchase", " pi ")):
        rank = 110
    elif re.search(r"\brcr\b", combined):
        rank = 120
    elif section == "PRODUCT EVALUATION" and "expectation" in combined:
        rank = 125
    elif "scent experience overall" in combined:
        rank = 130
    elif section == "PRODUCT EVALUATION" and any(word in combined for word in ("new different", "new & different")):
        rank = 135
    elif section == "FRAGRANCE STRENGTH" and re.match(r"q\s*16(?:\s|$)", qid) and "intensity" in combined:
        # Common Paired HUT flow: OFO → liking touchpoint → overall intensity → strength touchpoints.
        rank = 245
    elif section == "FIT" and re.search(r"\b(?:neat|wet|damp)\b", combined):
        rank = 430
    elif section == "FRAGRANCE CHARACTERISTICS":
        if _color_code(question_id, label):
            rank = 380
        elif re.match(r"q\s*18(?:\s|$)", _analysis_normal(question_id)) or "olfactive" in combined:
            rank = 320
        elif re.match(r"q\s*19(?:\s|$)", _analysis_normal(question_id)) or any(
            word in combined for word in ("characteristic", "perception")
        ):
            rank = 340
        elif re.match(r"q\s*20(?:\s|$)", _analysis_normal(question_id)) or "emotion" in combined:
            rank = 350
        elif any(word in combined for word in ("freshness type", "type of freshness")):
            rank = 315
        else:
            rank = 360
    else:
        ranks = {
            "PRODUCT EVALUATION": 140,
            "PRODUCT BENEFITS": 180,
            "FRAGRANCE EVALUATION": 225,
            "FRAGRANCE STRENGTH": 230,
            "FRAGRANCE JOURNEY - LIKING": 240,
            "FRAGRANCE JOURNEY - STRENGTH": 250,
            "FRAGRANCE BENEFITS": 260,
            "FIT": 270,
            "FRAGRANCE CHARACTERISTICS": 330,
            "ATTRIBUTES": 360,
            "USAGE & CONTEXT": 420,
        }
        rank = ranks.get(section, 500)
    return rank, first_row


def _study_question_priority(
    study_format: str,
    question_id: str,
    label: str,
    section: str,
    first_row: int,
    study_stages: Sequence[str] = (),
) -> tuple[int, int]:
    """Adapt the proposed narrative to the study design without hiding questions."""
    if study_format != "CLT":
        return _question_priority(question_id, label, section, first_row)

    combined = _analysis_normal(f"{question_id} | {label}")
    normalized_stages = {_normal(value) for value in study_stages if _normal(value)}
    wet_only = bool(normalized_stages) and normalized_stages.issubset({"wet", "damp wet"})
    stage = (
        "wet" if re.search(r"\b(?:damp wet|wet)\b", combined)
        else "neat" if re.search(r"\bneat\b", combined)
        else "wet" if wet_only
        else ""
    )
    if stage == "wet":
        if wet_only:
            if any(word in combined for word in ("overall fragrance", "overall scent", "f opo", "ofo", "ofl", "ofr")):
                return 100, first_row
            if section == "FRAGRANCE STRENGTH":
                return 110, first_row
            return 120, first_row
        if section == "FRAGRANCE CHARACTERISTICS":
            if "olfactive attribute" in combined or "olfactory attribute" in combined:
                return 130, first_row
            if "freshness type" in combined or "type of freshness" in combined:
                return 131, first_row
            if _color_code(question_id, label):
                return 132, first_row
        ranks = {
            "FRAGRANCE EVALUATION": 100,
            "FRAGRANCE STRENGTH": 110,
            "PRODUCT BENEFITS": 120,
            "FRAGRANCE BENEFITS": 125,
            "FRAGRANCE CHARACTERISTICS": 130,
            "FIT": 140,
        }
        return ranks.get(section, 145), first_row
    if stage == "neat":
        ranks = {
            "FRAGRANCE EVALUATION": 150,
            "FRAGRANCE STRENGTH": 160,
            "FRAGRANCE CHARACTERISTICS": 170,
            "FIT": 180,
        }
        return ranks.get(section, 175), first_row
    # Some CLT exports carry the stage only in A1 and have shortened Windows
    # filenames. In that case, use the validated scent-test storyline.
    if any(word in combined for word in ("overall product", "purchase intent", " rcr", " opo")):
        return 50, first_row
    if section == "FRAGRANCE EVALUATION" and any(
        word in combined for word in ("overall fragrance", "overall scent", "ofo", "ofr")
    ):
        return 100, first_row
    if section == "FRAGRANCE STRENGTH":
        return 110, first_row
    if section in {"PRODUCT EVALUATION", "FRAGRANCE EVALUATION"}:
        return 120, first_row
    if section in {"PRODUCT BENEFITS", "FRAGRANCE BENEFITS"}:
        return 130, first_row
    if section == "FRAGRANCE CHARACTERISTICS":
        if "freshness type" in combined or "type of freshness" in combined:
            return 140, first_row
        return 150, first_row
    return _question_priority(question_id, label, section, first_row)


def _summary_label(label: str) -> str:
    """Create a short presentation label while keeping it editable in the UI."""
    clean = _text(label)
    normalized = _normal(clean)
    replacements = {
        "overall product opinion": "Overall Opinion",
        "overall product liking": "Overall Opinion",
        "overall fragrance opinion": "Fragrance Opinion",
        "comparison with expectations": "Better vs. Expectations",
        "evaluation vs expectations": "Better vs. Expectations",
        "purchase intent": "Purchase Intent",
        "fit to brand": "Fit to Brand",
    }
    if normalized in replacements:
        return replacements[normalized]
    contains_replacements = (
        (("high quality",), "High Quality"),
        (("cleaning performance", "cleaning efficacy"), "Cleaning Performance"),
        (("bad smell", "bad odor", "bad odour", "malodor", "malodour"), "Effective vs. Bad Smells"),
        (("long lasting freshness and fragrance",), "Long-lasting Freshness"),
        (("long lasting fragrance", "long lastingness"), "Long-lasting Fragrance"),
        (("fresh for longer",), "Freshness Longevity"),
        (("freshness i notice when wearing",), "Freshness When Worn"),
        (("fragrance that i like",), "Fragrance Liking"),
        (("stands out",), "Distinctive Fragrance"),
        (("room smelling great",), "Room Fragrance"),
        (("keeps clothes soft", "softness"), "Softness"),
    )
    for terms, replacement in contains_replacements:
        if any(term in normalized for term in terms):
            return replacement
    clean = re.sub(r"^(?:This product|This fragrance|This smell)\s+", "", clean, flags=re.I)
    clean = re.sub(r"^(?:is|has|gives|leaves|keeps)\s+", "", clean, flags=re.I)
    return clean[:52].strip() or _text(label)


def _summary_score(proposal: QuestionProposal) -> int:
    """Rank decision-making KPIs; low-level batteries and technical items stay opt-in."""
    if not proposal.keep or proposal.question_type in {"Delete", "Strength"}:
        return -1
    if proposal.section in {
        "USAGE & CONTEXT",
        "FRAGRANCE JOURNEY - LIKING",
        "FRAGRANCE JOURNEY - STRENGTH",
        "FRAGRANCE STRENGTH",
        "FRAGRANCE CHARACTERISTICS",
        "ATTRIBUTES",
    }:
        return -1
    if proposal.question_type == "Attribute":
        return -1
    combined = _analysis_normal(f"{proposal.question_id} | {proposal.display_label} | {proposal.mapped_kpi}")
    priorities = (
        (1000, ("overall product opinion", "overall product liking", " opo")),
        (980, ("purchase intent", " pi")),
        (965, ("expectation",)),
        (950, ("overall satisfaction", "scent experience overall")),
        (940, ("rcr", "fit to brand")),
        (925, ("cleaning performance", "cleaning efficacy", "stain removal")),
        (915, ("overall fragrance", "overall scent", "f opo", "ofo", "ofr")),
        (905, ("long lasting fragrance", "long lastingness", "long lasting freshness")),
        (895, ("bad smell", "bad odor", "bad odour", "malodor", "malodour")),
        (885, ("high quality", "quality product")),
        (875, ("fresh for longer", "freshness")),
        (865, ("softness", "keeps clothes soft", "soft clothes")),
        (855, ("stands out", "new and different", "new different")),
        (845, ("smelling great", "fragrance that i like")),
    )
    for score, terms in priorities:
        if any(term in combined for term in terms):
            return score
    section_scores = {
        "PRODUCT EVALUATION": 800,
        "PRODUCT BENEFITS": 700,
        "FRAGRANCE EVALUATION": 650,
        "FRAGRANCE BENEFITS": 600,
        "FIT": 550,
    }
    return section_scores.get(proposal.section, -1)


def _suggest_question_splits(
    question_id: str,
    label: str,
    split_names: Sequence[str],
) -> tuple[str, ...]:
    """Suggest a split restriction only when the wording names it clearly."""
    combined = _normal(f"{question_id} {label}")
    matches: list[tuple[int, str]] = []
    for split_name in split_names:
        split_key = _normal(split_name)
        if split_key == "total":
            continue
        meaningful = re.sub(r"\b(?:mo|yo|years?|old)\b", " ", split_key)
        meaningful = re.sub(r"\s+", " ", meaningful).strip()
        if len(meaningful) >= 5 and meaningful in combined:
            matches.append((len(meaningful), split_name))
    if not matches:
        return ()
    longest = max(length for length, _ in matches)
    return tuple(split_name for length, split_name in matches if length == longest)


def _attribute_display_labels(
    question_id: str,
    label: str,
    question_type: str,
    section: str,
) -> tuple[str, str]:
    color_code = _color_code(question_id, label)
    if color_code:
        return "Color", color_code
    if question_type != "Attribute" or section not in {"FRAGRANCE CHARACTERISTICS", "PRODUCT BENEFITS"}:
        return label, ""
    qid = _normal(question_id)
    battery_source = re.sub(r"^Q\s*[-_]?\s*\d+[A-Z]?(?:\s*[-_]\s*\d+[A-Z]?)*\s*[-_:=]?\s*", "", question_id, flags=re.I)
    battery_patterns = (
        (r"(?:olfactive|olfactory)\s+attributes?\s*[-_:]\s*(.+)$", "Olfactive Attributes"),
        (r"(?:freshness\s+type|type\s+of\s+freshness)\s*[-_:]\s*(.+)$", "Type of Freshness"),
        (r"(?:^|[-_:])benefits?\s*[-_:]\s*(.+)$", "Benefits"),
        (r"(?:^|[-_:])attributes?\s*[-_:]\s*(.+)$", "Fragrance Attributes"),
    )
    for pattern, display_label in battery_patterns:
        match = re.search(pattern, battery_source, flags=re.I)
        if match:
            return display_label, re.sub(r"\s+", " ", match.group(1)).strip(" -_:=")
    metric_label = re.sub(
        r"^(?:characters?|characteristics?|consumer perceptions?|emotions?|feelings?)\s*[-_:]?\s*",
        "",
        label,
        flags=re.I,
    ).strip() or label
    if re.match(r"q\s*18(?:\s|$)", qid) or "olfactive" in _normal(label):
        return "Olfactive Space", metric_label
    if re.match(r"q\s*20(?:\s|$)", qid) or "emotion" in _normal(label):
        return "Emotions", metric_label
    if "characteristic" in _normal(f"{question_id} {label}"):
        return "Fragrance Characteristics", metric_label
    if re.match(r"q\s*19(?:\s|$)", qid) or "perception" in _normal(label):
        return "Consumer Perceptions", metric_label
    return "Fragrance Characteristics", metric_label


def _cmi_role(proposal: QuestionProposal) -> str:
    combined = _normal(f"{proposal.question_id} | {proposal.display_label} | {proposal.mapped_kpi}")
    if proposal.section == "USAGE & CONTEXT" or proposal.question_type == "Delete":
        return "Question technique"
    if proposal.section in {"FRAGRANCE CHARACTERISTICS", "ATTRIBUTES"} or proposal.question_type == "Attribute":
        return "Liste d’attributs"
    if proposal.section in {"PRODUCT BENEFITS", "FRAGRANCE BENEFITS"}:
        return "Bénéfice évalué"
    if proposal.section in {"FRAGRANCE JOURNEY - LIKING", "FRAGRANCE JOURNEY - STRENGTH", "FRAGRANCE STRENGTH"}:
        return "Mesure complémentaire"
    if any(word in combined for word in (
        "overall product", "overall fragrance", "purchase intent", "expectation", "overall satisfaction",
        "overall softness", "scent experience overall", "fit to brand", " rcr", " opo", " ofo", " ofr",
    )):
        return "Indicateur principal"
    return "Mesure complémentaire"


def _cmi_note(proposal: QuestionProposal) -> str:
    combined = _analysis_normal(f"{proposal.question_id} | {proposal.display_label} | {proposal.mapped_kpi}")
    if _is_maxdiff_question(proposal.question_id, proposal.display_label, proposal.mapped_kpi):
        return "Question technique laissée de côté"
    if proposal.section == "USAGE & CONTEXT" or proposal.question_type == "Delete":
        return "Hors toplines proposé ; disponible si le projet l’exige"
    if any(word in combined for word in (
        "comparison to current", "comparison with current", "comparison smell", "feeling about new",
        "feeling about the new",
        "preference", "perception mapping",
    )):
        return "À confirmer selon l’objectif du projet"
    return ""


def _comparative_question(question_id, mapped_kpi, metrics):
    wording = _analysis_normal(f"{question_id} {mapped_kpi}")
    if re.search(r"\b(?:pref\w*|most|strongest|which|compar\w*)\b", wording):
        return True
    choices = " ".join(_analysis_normal(m) for m in metrics if re.match(r"^\d+\s", _analysis_normal(m)))
    return bool(re.search(r"\b(?:prefer\w*|usual|habituel\w*|current|actuel\w*)\b", choices))


def _preference_question(question_id, mapped_kpi, metrics):
    """Recognise explicit preference without treating every comparison alike."""
    wording = _analysis_normal(f"{question_id} {mapped_kpi}")
    direct = re.search(
        r"\b(?:pref(?:er(?:red|ence)?)?|favorite|favourite|favori\w*|prefere\w*)\b|"
        r"\b(?:like|liked)\b.{0,24}\bmost\b|\bwhich\b.{0,24}\bprefer\w*\b|\bthe most\b|\bmost\b",
        wording,
    )
    choices = " ".join(_analysis_normal(metric) for metric in metrics if re.match(r"^\d+\s", _analysis_normal(metric)))
    return bool(direct or re.search(r"\b(?:most|i prefer|je prefere|preferred)\b", choices))


def _bipolar_question(question_id, mapped_kpi, metrics):
    wording = _analysis_normal(f"{question_id} {mapped_kpi}")
    if "bipolar" in wording or "bipolaire" in wording or "semantic differential" in wording:
        return True
    choices = " ".join(_analysis_normal(m) for m in metrics if re.match(r"^\d+\s", _analysis_normal(m)))
    opposites = (("feminine", "masculine"), ("feminin", "masculin"),
                 ("traditional", "modern"), ("natural", "artificial"),
                 ("warm", "cool"), ("sweet", "fresh"))
    return any(a in choices and b in choices for a, b in opposites)


def _classify(question_id: str, mapped_kpi: str, metrics: Sequence[str]) -> tuple[str, str]:
    normalized_metrics = [_analysis_normal(metric) for metric in metrics]
    combined = _analysis_normal(f"{question_id} {mapped_kpi}")
    if _is_maxdiff_question(question_id, humanize_question(question_id, mapped_kpi), mapped_kpi):
        return "Delete", "Élevée"
    if re.search(r"\bother\s*\(?\s*precise\b", combined) or "autre preciser" in combined:
        return "Delete", "Moyenne"
    has_mean = any(_metric_key(metric) == "mean" for metric in metrics)
    has_box = any(re.fullmatch(r"(?:top|bottom)_[123]", _metric_key(metric)) or "box" in _normal(metric) or "boite" in _normal(metric) for metric in metrics)
    response_codes = {
        int(match.group(1))
        for source,metric in zip(metrics,normalized_metrics)
        if not re.fullmatch(r'(?:top|bottom)_[123]',_metric_key(source))
        if (match := re.match(r"^(\d+)\s", metric))
    }
    jar_words = ("juste", "just right", "just about right", "ideal", "parfait", "perfect")
    jar = any(
        re.match(r"^3\s", metric)
        and any(word in metric for word in jar_words)
        for metric in normalized_metrics
    )
    direct_strength = (
        any("too weak" in metric or "trop faible" in metric for metric in normalized_metrics)
        and any("just about right" in metric or "just right" in metric or "juste" in metric for metric in normalized_metrics)
        and any("too strong" in metric or "trop fort" in metric for metric in normalized_metrics)
    )
    strength_word = any(word in combined for word in ("intens", "strength", "puissance", "forte", "strong"))

    if _preference_question(question_id, mapped_kpi, metrics):
        return "Preference", "Élevée"
    # Comparative choices must not be collapsed into one CATA response.
    if _comparative_question(question_id, mapped_kpi, metrics):
        return "Autres", "Moyenne"
    if _bipolar_question(question_id, mapped_kpi, metrics):
        return "Bipolaire", "Moyenne"
    if response_codes == {1, 2} and any(re.match(r"^1\s+(?:no|non)(?:\s|$)", m) for m in normalized_metrics):
        return "Attribute", "Élevée"
    if jar or direct_strength:
        return "Strength", "Élevée"
    if strength_word and has_box:
        return "Strength", "Moyenne"
    if response_codes and (not has_mean or any(word in combined for word in ("listing", "list of", "select all", "which of"))):
        return "Listing", "Moyenne"
    if has_mean and has_box:
        return "Standard", "Élevée"
    if has_box:
        return "Standard", "Moyenne"
    return "Autres", "Faible"


def _analyze_question_sheet(sheet) -> tuple:
    """Analyse every question in one validated G-Sight result sheet."""
    layout = detect_layout(sheet)
    source_stage = _stage_name(sheet)
    groups: dict[str, dict] = {}
    for row in range(layout.header_row + 1, sheet.max_row + 1):
        question_id = _text(sheet.cell(row, layout.metric_col - 1).value)
        mapped_kpi = _text(sheet.cell(row, max(1, layout.metric_col - 2)).value)
        metric = _text(sheet.cell(row, layout.metric_col).value)
        if not question_id or not metric:
            continue
        key = question_id.casefold()
        if key not in groups:
            groups[key] = {
                "question_id": question_id,
                "mapped_kpi": mapped_kpi,
                "first_row": row,
                "metrics": [],
            }
        if metric not in groups[key]["metrics"]:
            groups[key]["metrics"].append(metric)

    prepared = []
    for group in groups.values():
        label = humanize_question(group["question_id"], group["mapped_kpi"])
        question_type, confidence = _classify(group["question_id"], group["mapped_kpi"], group["metrics"])
        color_code = _color_code(group["question_id"], label)
        if color_code:
            label = "Color"
        section = _suggest_section(
            group["question_id"],
            label,
            question_type,
            group["mapped_kpi"],
        )
        label, metric_label = _attribute_display_labels(
            group["question_id"], label, question_type, section
        )
        prepared.append((
            _question_priority(group["question_id"], label, section, group["first_row"]),
            group,
            label,
            metric_label,
            question_type,
            confidence,
            section,
        ))

    proposals = []
    ordered_prepared = sorted(
        prepared,
        key=lambda item: (item[0], _normal(item[1]["question_id"])),
    )
    for order, (_, group, label, metric_label, question_type, confidence, section) in enumerate(ordered_prepared, 1):
        combined = _normal(f"{group['question_id']} {group['mapped_kpi']} {label} {section}")
        explicit_stages = tuple(
            stage for stage in ("NEAT", "WET", "DRY", "DAMP")
            if re.search(rf"\b{_normal(stage)}\b", combined)
        )
        fallback_stage = () if _normal(source_stage) in {"", "none", "not specified", "non precise"} else (source_stage,)
        proposals.append(
            QuestionProposal(
                keep=question_type != "Delete" and section != "USAGE & CONTEXT",
                order=order,
                section=section,
                question_id=group["question_id"],
                mapped_kpi=group["mapped_kpi"],
                display_label=label,
                question_type=question_type,
                confidence=confidence,
                first_row=group["first_row"],
                metrics=tuple(group["metrics"]),
                metric_label=metric_label,
                stages=explicit_stages or fallback_stage,
            )
        )
    return sheet.title, layout, tuple(proposals)


def analyze_questions(source) -> tuple:
    """Analyse the primary result sheet for backward-compatible callers."""
    workbook = _workbook(source)
    return _analyze_question_sheet(detect_data_sheet(workbook))


def inspect_smart_package(raw_files: Sequence[tuple[str, object]]) -> SmartPackageInfo:
    if not raw_files:
        raise TemplyfierError("Ajoute au moins un export G-Sight.")
    candidates = []
    all_stages: list[str] = []
    all_splits: list[str] = []
    warnings: list[str] = []
    max_rows = 0
    for filename, source in raw_files:
        workbook = _workbook(source)
        sheet = detect_data_sheet(workbook)
        layout = _active_layout(sheet, detect_layout(sheet))
        score = sum(count or 0 for count in _signature(sheet, layout))
        embedded = _embedded_data_sheets(workbook)
        stage = _stage_name(sheet)
        max_rows = max(max_rows, sheet.max_row)
        role = "Comparaison benchmark" if re.search(r"compar(?:aison|ison)?\s+bench", filename, flags=re.I) else "Résultats"
        if re.search(r"all\s+splits", filename, flags=re.I) and len(raw_files) > 1:
            role = "Consolidé de contrôle"
        embedded_comparison_codes = tuple(dict.fromkeys(
            code
            for embedded_sheet in embedded
            for code in _comparison_codes(
                embedded_sheet,
                _active_layout(embedded_sheet, detect_layout(embedded_sheet)),
            )
        ))
        all_stages.extend(_stage_name(item) for item in embedded)
        all_splits.extend(_split_label(item) for item in embedded)
        candidates.append((score, filename, source, sheet, layout, role, stage, len(embedded), embedded_comparison_codes))

    study_format = _study_format([name for name, _ in raw_files], tuple(dict.fromkeys(all_stages)), max_rows)
    if study_format == "HUT / in-use":
        by_split: dict[str, list[int]] = {}
        for index, item in enumerate(candidates):
            if item[5] == "Résultats":
                by_split.setdefault(_detected_split_label(item[3]), []).append(index)
        for indexes in by_split.values():
            if len(indexes) < 2:
                continue
            distinct_codes = {
                code for index in indexes for code in candidates[index][8]
            }
            if len(distinct_codes) > max(len(candidates[index][8]) for index in indexes):
                # Same consumer split exported once per benchmark: all files are useful.
                continue
            richest = max(indexes, key=lambda index: len(candidates[index][8]))
            for index in indexes:
                if index != richest:
                    row = list(candidates[index])
                    row[5] = "Comparaison benchmark"
                    candidates[index] = tuple(row)
    result_candidates = [item for item in candidates if item[5] == "Résultats"] or candidates
    _, filename, source, sheet, layout, _, _, _, _ = max(result_candidates, key=lambda item: item[0])
    inputs = []
    for score, item_filename, _, item_sheet, item_layout, role, stage, embedded_count, comparison_codes in candidates:
        split_name,split_detection=_split_detection(item_sheet,item_filename)
        inputs.append(
            SmartInputInfo(
                filename=item_filename,
                split_name=split_name,
                counts=_signature(item_sheet, item_layout),
                product_names=_product_names(item_sheet, item_layout),
                role=role,
                stage=stage,
                embedded_split_count=embedded_count,
                comparison_codes=comparison_codes,
                product_keys=_product_keys(item_sheet, item_layout),
                split_detection=split_detection,
            )
        )
    unique_stages = tuple(dict.fromkeys(all_stages))
    unique_splits = tuple(dict.fromkeys(all_splits))
    if len(unique_stages) > 1 and study_format != "HUT / in-use":
        warnings.append("Plusieurs stages détectés : ils doivent être assemblés, pas traités comme des splits consommateurs.")
    if study_format == "HUT / in-use" and len(unique_stages) > 1:
        warnings.append("HUT reconnu : les valeurs Stage None/USE sont des configurations de comparaison benchmark, pas des touchpoints.")
    if max(item.embedded_split_count for item in inputs) > 1:
        warnings.append("Plusieurs tables/splits sont intégrés dans chaque export G-Sight.")
    if any(item.role == "Comparaison benchmark" for item in inputs):
        warnings.append("Les fichiers de comparaison serviront aux couleurs de significativité, pas aux valeurs principales.")
    uncertain_splits=[item.filename for item in inputs if item.split_detection=='Review recommended']
    if uncertain_splits:
        warnings.append("Review the detected split for: " + " · ".join(uncertain_splits) + ". The proposed name comes from the filename and remains editable.")
    # Union of the questions across result exports: a boost-only question must
    # be proposed even when it does not exist in the TOTAL export.
    merged_questions: dict[str, QuestionProposal] = {}
    availability: dict[str, set[str]] = {}
    metric_availability: dict[str, dict[str, set[str]]] = {}
    question_candidates = [
        (candidate, input_info)
        for candidate, input_info in zip(candidates, inputs)
        if input_info.role == "Résultats"
    ] or list(zip(candidates, inputs))
    source_sheet = ""
    for candidate, input_info in question_candidates:
        _, item_filename, item_source, _, _, _, _, _, _ = candidate
        item_workbook = _workbook(item_source)
        # A single consolidated G-Sight workbook may contain one valid result
        # sheet per stage (for example NEAT and WET). The former implementation
        # analysed only detect_data_sheet(), so questions exclusive to every
        # other stage disappeared from the review and therefore from Excel.
        for item_sheet in _embedded_data_sheets(item_workbook):
            item_layout = _active_layout(item_sheet, detect_layout(item_sheet))
            item_source_sheet, _, item_questions = _analyze_question_sheet(item_sheet)
            if item_filename == filename and not source_sheet:
                source_sheet = item_source_sheet
            indexed_metrics: dict[str, dict[str, str]] = {}
            for source_row in range(item_layout.header_row + 1, item_sheet.max_row + 1):
                qid=_text(item_sheet.cell(source_row,item_layout.metric_col-1).value)
                metric=_text(item_sheet.cell(source_row,item_layout.metric_col).value)
                has_value=any(isinstance(item_sheet.cell(source_row,col).value,(int,float))
                    and not isinstance(item_sheet.cell(source_row,col).value,bool) for col in item_layout.product_cols)
                if qid and metric and has_value:
                    indexed_metrics.setdefault(qid.casefold(),{})[_metric_key(metric)]=metric
            for proposal in item_questions:
                key = proposal.question_id.casefold()
                if key not in indexed_metrics:
                    continue
                availability.setdefault(key, set()).add(input_info.split_name)
                for metric_key in indexed_metrics[key]:
                    metric_availability.setdefault(key,{}).setdefault(metric_key,set()).add(item_filename)
                current = merged_questions.get(key)
                if current is None:
                    merged_questions[key] = proposal
                else:
                    # Keep every available metric, even when stages or splits
                    # expose equally long but different metric lists. Values
                    # remain attached to their original result sheet.
                    merged_questions[key] = replace(
                        current,
                        metrics=tuple(dict.fromkeys((*current.metrics, *proposal.metrics))),
                        stages=tuple(dict.fromkeys((*current.stages, *proposal.stages))),
                    )

    prepared_questions = []
    for key, proposal in merged_questions.items():
        presence=metric_availability.get(key,{})
        proposal=replace(proposal,metric_availability=tuple(
            (metric,tuple(item.filename for item in inputs if item.filename in presence.get(_metric_key(metric),set())))
            for metric in proposal.metrics
        ))
        prepared_questions.append((
            _study_question_priority(
                study_format,
                proposal.question_id,
                proposal.display_label,
                proposal.section,
                proposal.first_row,
                unique_stages,
            ),
            proposal,
            tuple(sorted(availability.get(key, ()), key=_normal)),
        ))
    # Several split exports can introduce different questions on the exact same
    # source row and with the same CMI priority. Never let Python fall through
    # to comparing QuestionProposal objects, which are intentionally not
    # orderable; use the stable question identifier as the tie-breaker.
    sorted_questions = sorted(
        prepared_questions,
        key=lambda item: (item[0], _normal(item[1].question_id)),
    )
    result_split_names = tuple(dict.fromkeys(item.split_name for item in inputs if item.role == "Résultats"))
    summary_candidates = sorted(
        (
            (_summary_score(proposal), order_index, proposal.question_id.casefold())
            for order_index, (_, proposal, _) in enumerate(sorted_questions)
            if _summary_score(proposal) >= 0
        ),
        key=lambda item: (-item[0], item[1]),
    )
    proposal_by_id = {proposal.question_id.casefold(): proposal for _, proposal, _ in sorted_questions}
    suggested_summary_ids = set()
    seen_summary_labels = set()
    for _, _, question_id in summary_candidates:
        proposal = proposal_by_id[question_id]
        label_key = _normal(_summary_label(proposal.display_label))
        if label_key in seen_summary_labels:
            continue
        suggested_summary_ids.add(question_id)
        seen_summary_labels.add(label_key)
        if len(suggested_summary_ids) == 10:
            break
    split_suggestions: dict[str, tuple[str, ...]] = {}
    prior_named: tuple[int, str, int, tuple[str, ...]] | None = None
    for _, proposal, _ in sorted_questions:
        suggestion = _suggest_question_splits(
            proposal.question_id,
            proposal.display_label,
            result_split_names,
        )
        qmatch = re.match(r"q\s*[-_]?\s*(\d+)(?:\s*[-_]?\s*([a-z]))?", proposal.question_id, flags=re.I)
        if not suggestion and prior_named and qmatch:
            qnumber, qletter = int(qmatch.group(1)), (qmatch.group(2) or "").casefold()
            prior_number, prior_letter, prior_row, prior_suggestion = prior_named
            if qnumber == prior_number + 1 and qletter == prior_letter and proposal.first_row - prior_row <= 25:
                suggestion = prior_suggestion
        if suggestion and qmatch:
            prior_named = (
                int(qmatch.group(1)),
                (qmatch.group(2) or "").casefold(),
                proposal.first_row,
                suggestion,
            )
        split_suggestions[proposal.question_id.casefold()] = suggestion

    questions = tuple(
        QuestionProposal(
            keep=proposal.keep,
            order=order,
            section=proposal.section,
            question_id=proposal.question_id,
            mapped_kpi=proposal.mapped_kpi,
            display_label=proposal.display_label,
            question_type=proposal.question_type,
            confidence=proposal.confidence,
            first_row=proposal.first_row,
            metrics=proposal.metrics,
            metric_label=proposal.metric_label,
            availability=item_availability,
            suggested_splits=split_suggestions.get(proposal.question_id.casefold(), ()),
            summary_keep=proposal.question_id.casefold() in suggested_summary_ids,
            summary_label=_summary_label(proposal.display_label),
            metric_availability=proposal.metric_availability,
            stages=proposal.stages,
        )
        for order, (_, proposal, item_availability) in enumerate(sorted_questions, 1)
    )
    if not source_sheet:
        source_sheet, _, _ = analyze_questions(source)
    suggested_benchmark_count = _suggest_benchmark_count(inputs)
    return SmartPackageInfo(
        reference_filename=filename,
        source_sheet=source_sheet,
        product_names=next((item.product_names for item in inputs if item.role == "Résultats"), inputs[0].product_names),
        counts=_signature(sheet, layout),
        questions=questions,
        inputs=tuple(inputs),
        study_format=study_format,
        stages=unique_stages,
        embedded_splits=unique_splits,
        warnings=tuple(warnings),
        suggested_benchmark_count=suggested_benchmark_count,
        suggested_test_type=_suggest_test_type(inputs),
    )


def profile_to_json(question_rows: Sequence[dict], settings: dict) -> bytes:
    payload = {"version": 1, "settings": settings, "questions": list(question_rows)}
    return json.dumps(
        payload,
        ensure_ascii=False,
        indent=2,
        default=lambda value: value.item() if hasattr(value, "item") else str(value),
    ).encode("utf-8")


def profile_from_json(source: bytes | BinaryIO) -> dict:
    from .review import validate_profile
    raw = source if isinstance(source, bytes) else source.read()
    try:
        return validate_profile(json.loads(raw.decode("utf-8-sig")))
    except (ValueError, UnicodeDecodeError) as exc:
        raise TemplyfierError(f"Impossible de lire les réglages : {exc}") from exc


def apply_profile(proposals: Sequence[QuestionProposal], profile: dict) -> list[dict]:
    saved = {str(row.get("Question ID", "")).casefold(): row for row in profile.get("questions", [])}
    result = []
    for proposal in proposals:
        base = proposal_to_row(proposal)
        prior = saved.get(proposal.question_id.casefold())
        if prior:
            for field in (
                "Keep", "Order", "Section", "Display label", "Metric label", "Type", "Custom metrics",
                "Included splits", "Stage", "KPI Summary", "Summary label", "Sens favorable",
                "Selected metrics", "Metric labels", "Selection type", "Group ID",
                "Grouping choice", "Dismissed groups", "Ungrouped labels", *STANDARD_METRICS
            ):
                if field in prior:
                    base[field] = prior[field]
            # Correct the historic Preference-as-Attribute profile error, but
            # preserve intentional selections made with the new editor.
            if "Selected metrics" not in prior and proposal.question_type in {"Autres", "Preference"} and base["Type"] == "Attribute":
                base["Type"] = proposal.question_type
                base["Custom metrics"] = ""
            base["Confidence"] = "Profil"
        result.append(base)
    return result


def proposal_to_row(proposal: QuestionProposal) -> dict:
    row = {
        "Keep": proposal.keep,
        "Order": proposal.order,
        "Section": proposal.section,
        "Question ID": proposal.question_id,
        "Display label": proposal.display_label,
        "Metric label": proposal.metric_label,
        "Type": proposal.question_type,
        "Confidence": proposal.confidence,
        "CMI role": _cmi_role(proposal),
        "CMI note": _cmi_note(proposal),
        "Available metrics": " · ".join(proposal.metrics),
        "Availability": " · ".join(proposal.availability) or "Tous les splits",
        "Included splits": "; ".join(proposal.suggested_splits) or "Tous les splits",
        "Stage": " ; ".join(proposal.stages) or "Unassigned",
        "KPI Summary": proposal.summary_keep,
        "Summary label": proposal.summary_label or _summary_label(proposal.display_label),
        "Sens favorable": "Idéal au centre" if proposal.question_type == "Strength" else "Automatique",
        "Custom metrics": "",
    }
    defaults = {"Mean", "Top Box", "Top 2 Boxes", "Bottom 2 Boxes"}
    available = {_metric_key(metric) for metric in proposal.metrics}
    for metric in STANDARD_METRICS:
        row[metric] = proposal.question_type == "Standard" and _metric_key(metric) in available and metric in defaults
    return row


def _standard_metrics_for_question(config: dict, fallback: Sequence[str]) -> tuple[str, ...]:
    custom = _text(config.get("Custom metrics"))
    if custom:
        return tuple(value.strip() for value in re.split(r"[;\n|]+", custom) if value.strip())
    explicit = tuple(metric for metric in STANDARD_METRICS if bool(config.get(metric)))
    has_columns = any(metric in config for metric in STANDARD_METRICS)
    return explicit if has_columns else tuple(fallback)


def _metric_matches(metric: str, selected: Sequence[str]) -> bool:
    return _metric_key(metric) in {_metric_key(value) for value in selected}


def _strength_keep(metric: str) -> bool:
    norm = _normal(metric)
    if _metric_key(metric) in {'top_2','bottom_2'}:
        return True
    if norm in {"too weak", "too strong", "trop faible", "trop fort", "just about right", "just right", "juste comme il faut"}:
        return True
    return bool(re.match(r"^3\s", norm)) and any(
        word in norm for word in ("juste", "just right", "just about right", "ideal", "parfait", "perfect")
    )


def _attribute_keep(metric: str) -> bool:
    norm = _normal(metric)
    key=_metric_key(metric)
    if re.fullmatch(r'(?:top|bottom)_[123]',key):return key=='top_1'
    return bool(re.match(r"^2\s", norm))


def _find_source_rows(sheet, layout, question_id: str, question_type: str, standard_metrics: Sequence[str]):
    question_candidates = [
        row for row in range(layout.header_row + 1, sheet.max_row + 1)
        if _text(sheet.cell(row, layout.metric_col - 1).value).casefold() == question_id.casefold()
    ]
    if question_type == "Strength":
        direct_labels = {"too weak", "just about right", "too strong"}
        present = {_analysis_normal(sheet.cell(row, layout.metric_col).value) for row in question_candidates}
        if direct_labels.issubset(present):
            return [
                row for row in question_candidates
                if _analysis_normal(sheet.cell(row, layout.metric_col).value) in direct_labels
            ]
    rows = []
    for row in question_candidates:
        metric = _text(sheet.cell(row, layout.metric_col).value)
        keep = (
            _metric_matches(metric, standard_metrics) if question_type == "Standard"
            else _strength_keep(metric) if question_type == "Strength"
            else _attribute_keep(metric) if question_type == "Attribute"
            else _metric_matches(metric, standard_metrics) if question_type == "Libre" and standard_metrics
            else question_type == "Libre"
        )
        if keep:
            values = [sheet.cell(row, col).value for col in layout.product_cols]
            if any(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
                rows.append(row)
    if question_type == "Attribute" and len(rows) > 1:
        preferred = [row for row in rows if re.match(r"^2\s", _normal(sheet.cell(row, layout.metric_col).value))]
        rows = preferred[:1] or rows[:1]
    return rows


def default_metric_selection(question_type, available):
    """Suggest only metrics actually exported by G-Sight; never fabricate boxes."""
    metrics = list(available)
    if question_type == "Standard":
        wanted = {"mean", "top_1", "top_2", "bottom_2"}
        return [m for m in metrics if _metric_key(m) in wanted]
    if question_type == "Strength":
        direct = {"too weak", "just about right", "too strong"}
        if direct.issubset({_analysis_normal(m) for m in metrics}):
            return [m for m in metrics if _analysis_normal(m) in direct]
        return [m for m in metrics if _strength_keep(m)]
    if question_type in {"CATA", "Attribute"}:
        positive = [m for m in metrics if re.match(r"^2\s", _normal(m)) and not re.fullmatch(r'(?:top|bottom)_[123]',_metric_key(m))]
        return positive or [m for m in metrics if _metric_key(m) == "top_1"]
    if question_type == "Delete":
        return []
    numbered = [m for m in metrics if re.match(r"^\d+\s", _normal(m)) and not re.fullmatch(r'(?:top|bottom)_[123]',_metric_key(m))]
    return numbered or [m for m in metrics if _metric_key(m) not in {"mean", "top_1", "top_2", "top_3", "bottom_1", "bottom_2", "bottom_3"}]


def configured_source_rows(sheet, layout, config, fallback):
    question_id = _text(config.get("Question ID"))
    selected = config.get("Selected metrics")
    if isinstance(selected, (list, tuple)):
        # Explicit CMI selection overrides every type-specific default, including
        # strength and CATA. Empty means empty, not 'all metrics'.
        keys = {_metric_key(m) for m in selected}
        return [r for r in range(layout.header_row + 1, sheet.max_row + 1)
                if _text(sheet.cell(r, layout.metric_col - 1).value).casefold() == question_id.casefold()
                and _metric_key(_text(sheet.cell(r, layout.metric_col).value)) in keys
                and any(isinstance(sheet.cell(r, c).value, (int, float))
                        and not isinstance(sheet.cell(r, c).value, bool) for c in layout.product_cols)]
    question_type = _text(config.get("Type"))
    if question_type in {"Autres", "Bipolaire", "Listing", "Preference"}:
        available = [_text(sheet.cell(r, layout.metric_col).value)
                     for r in range(layout.header_row + 1, sheet.max_row + 1)
                     if _text(sheet.cell(r, layout.metric_col - 1).value).casefold() == question_id.casefold()]
        custom = _text(config.get("Custom metrics"))
        selected = (_standard_metrics_for_question(config, fallback) if custom
                    else default_metric_selection(question_type, available))
        return configured_source_rows(sheet, layout, {**config, "Selected metrics": selected}, fallback)
    return _find_source_rows(sheet, layout, question_id, "Attribute" if question_type == "CATA" else question_type,
                             _standard_metrics_for_question(config, fallback))


def clean_metric_label(config, source_metric):
    labels = config.get("Metric labels", {})
    if isinstance(labels, dict):
        for source, label in labels.items():
            if _metric_key(source) == _metric_key(source_metric) and _text(label):
                return _text(label)
    return source_metric


def grouped_metric_label(config, source_metric, count):
    item = _text(config.get('Metric label')) or _text(config.get('Display label'))
    label = clean_metric_label(config, source_metric)
    return item if count == 1 and label == source_metric else f'{item} · {label}'


def _copy_sheet(source, target):
    for row in source.iter_rows():
        for cell in row:
            new = target[cell.coordinate]
            new.value = cell.value
            if cell.has_style:
                new._style = copy(cell._style)
            if cell.number_format:
                new.number_format = cell.number_format
    for key, dim in source.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
        target.column_dimensions[key].hidden = dim.hidden
    for key, dim in source.row_dimensions.items():
        target.row_dimensions[key].height = dim.height
        target.row_dimensions[key].hidden = dim.hidden
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def _style_sheet(sheet, max_col: int):
    dark = "#3E356B"
    light = "#EDEAF7"
    grid = Side(style="thin", color="B8B8C3")
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C6"
    sheet.column_dimensions["A"].width = 62
    sheet.column_dimensions["B"].width = 25
    for col in range(3, max_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 12
    sheet.row_dimensions[3].height = 86
    sheet.row_dimensions[4].height = 22
    sheet.row_dimensions[5].height = 24
    sheet.cell(5, 1).value = "Variable"
    sheet.cell(5, 2).value = "Metric"
    header = sheet.cell(5, 1).parent.cell
    for row in (3, 4, 5):
        for col in range(1, max_col + 1):
            cell = sheet.cell(row, col)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=grid, top=grid, left=grid, right=grid)
    for col in range(1, max_col + 1):
        sheet.cell(5, col).fill = PatternFill("solid", fgColor=dark.replace("#", ""))
        sheet.cell(5, col).font = Font(color="FFFFFF", bold=True)


def _format_variable_label(sheet, row: int) -> None:
    cell = sheet.cell(row, 1)
    if not cell.value:
        return
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    length = len(_text(cell.value))
    if length > 115:
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, 42)
    elif length > 70:
        sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, 30)


def _mean_number_format(decimals: int) -> str:
    decimals = max(0, min(2, int(decimals)))
    return "0" if decimals == 0 else "0." + ("0" * decimals)


def _selected_question_rows(question_rows: Sequence[dict]) -> list[dict]:
    selected = [row for row in question_rows if bool(row.get("Keep")) and row.get("Type") != "Delete"]
    selected.sort(key=lambda row: (int(row.get("Order", 9999)), str(row.get("Question ID", ""))))
    if not selected:
        raise TemplyfierError("Aucune question n’est sélectionnée.")
    return selected


def _question_applies_to_split(config: dict, split_name: str) -> bool:
    raw = _text(config.get("Included splits"))
    if not raw or _normal(raw) in {"all", "all splits", "tous", "tous les splits", "*"}:
        return True
    requested = {
        _normal(value)
        for value in re.split(r"[;|\n]+", raw)
        if _text(value)
    }
    return _normal(split_name) in requested


def _source_metric_row_count(sheet, layout) -> int:
    return sum(
        1 for row in range(layout.header_row + 1, sheet.max_row + 1)
        if any(isinstance(sheet.cell(row, col).value, (int, float)) and not isinstance(sheet.cell(row, col).value, bool) for col in layout.product_cols)
    )


def _build_paired_toplines(
    raw_files: Sequence[tuple[str, object]],
    question_rows: Sequence[dict],
    *,
    split_names: Sequence[str],
    standard_metrics: Sequence[str],
    include_screeners: bool,
    mean_decimals: int,
    paired_swaps: dict[str, Sequence[int]] | None,
    include_deltas: bool,
    include_sections: bool,
    summary_scope: str,
    summary_metric_strategy: str,
    include_summary_details: bool,
) -> tuple[bytes, dict]:
    workbooks = [_workbook(source) for _, source in raw_files]
    data_sheets = [detect_data_sheet(workbook) for workbook in workbooks]
    layouts = [_active_layout(sheet, detect_layout(sheet)) for sheet in data_sheets]
    if len(split_names) != len(raw_files):
        raise TemplyfierError("Le nombre de noms de splits est incorrect.")
    product_keys = [_product_keys(sheet, layout) for sheet, layout in zip(data_sheets, layouts)]
    product_counts = [len(keys) for keys in product_keys]
    if any(count < 2 or count % 2 for count in product_counts):
        raise TemplyfierError(
            "Un test Paired doit contenir un nombre pair de produits, ordonnés Benchmark puis Candidat pour chaque paire."
        )
    selected_questions = _selected_question_rows(question_rows)
    output = Workbook()
    output.remove(output.active)
    used_names: set[str] = set()
    written_data_rows = 0
    skipped_questions = 0
    mean_format = _mean_number_format(mean_decimals)
    pairs_by_split: dict[str, int] = {}
    summary_records = []

    for file_index, ((filename, _), source_sheet, layout) in enumerate(zip(raw_files, data_sheets, layouts)):
        sheet_name = _safe_sheet_name(split_names[file_index], used_names)
        target = output.create_sheet(sheet_name)
        pairs = [(pos, pos + 1) for pos in range(0, len(layout.product_cols), 2)]
        swaps = {int(value) for value in (paired_swaps or {}).get(filename, ())}
        blocks = []
        current_col = 3
        for pair_number, (benchmark_pos, candidate_pos) in enumerate(pairs, 1):
            if pair_number in swaps:
                benchmark_pos, candidate_pos = candidate_pos, benchmark_pos
            blocks.append({
                "benchmark_pos": benchmark_pos,
                "candidate_pos": candidate_pos,
                "benchmark": current_col,
                "candidate": current_col + 1,
                "delta": current_col + 2 if include_deltas else None,
                "spacer": current_col + (3 if include_deltas else 2),
            })
            current_col += 4 if include_deltas else 3
        # A narrow spacer separates two paired blocks, but the sheet should not
        # end with an empty technical column.
        max_col = current_col - 2
        pairs_by_split[sheet_name] = len(pairs)
        _style_sheet(target, max_col)
        for pair_number, block in enumerate(blocks, 1):
            for role, position in (("benchmark", block["benchmark_pos"]), ("candidate", block["candidate_pos"])):
                raw_col = layout.product_cols[position]
                out_col = block[role]
                target.cell(3, out_col).value = source_sheet.cell(layout.product_name_row, raw_col).value
                target.cell(4, out_col).value = source_sheet.cell(layout.sample_row, raw_col).value
                target.cell(5, out_col).value = source_sheet.cell(layout.product_header_row, raw_col).value
            if block["delta"] is not None:
                target.cell(3, block["delta"]).value = f"Paire {pair_number}"
                target.cell(5, block["delta"]).value = "Delta Candidat - Benchmark"
            if pair_number < len(blocks):
                target.column_dimensions[get_column_letter(block["spacer"])].width = 3

        output_row = 6
        current_section = None
        previous_variable_label = None
        summary_question_rows: dict[str, list[tuple[int, str]]] = {}
        for config in selected_questions:
            if not _question_applies_to_split(config, split_names[file_index]):
                continue
            question_id = _text(config.get("Question ID"))
            question_type = _text(config.get("Type"))
            question_metrics = _standard_metrics_for_question(config, standard_metrics)
            source_rows = configured_source_rows(source_sheet, layout, config, standard_metrics)
            if not source_rows:
                skipped_questions += 1
                continue
            section = _text(config.get("Section")) or "TOPLINES"
            if include_sections and section != current_section:
                target.cell(output_row, 1).value = section
                target.merge_cells(start_row=output_row, start_column=1, end_row=output_row, end_column=max_col)
                cell = target.cell(output_row, 1)
                cell.fill = PatternFill("solid", fgColor="D9D3EA")
                cell.font = Font(bold=True, color="2F2952")
                output_row += 1
                current_section = section
                previous_variable_label = None

            for metric_index, source_row in enumerate(source_rows):
                source_metric = _text(source_sheet.cell(source_row, layout.metric_col).value)
                display_label = _text(config.get("Display label"))
                metric_label = _text(config.get("Metric label"))
                if metric_label:
                    group_identity = (_text(config.get("Group ID")) or display_label, section, display_label)
                    target.cell(output_row, 1).value = display_label if group_identity != previous_variable_label else None
                    target.cell(output_row, 2).value = (
                        grouped_metric_label(config, source_metric, len(source_rows))
                    )
                    previous_variable_label = group_identity
                elif question_type in {"Attribute", "CATA"}:
                    target.cell(output_row, 2).value = (display_label if len(source_rows) == 1 else f"{display_label} · {clean_metric_label(config, source_metric)}")
                    previous_variable_label = None
                else:
                    target.cell(output_row, 1).value = display_label if metric_index == 0 else None
                    target.cell(output_row, 2).value = clean_metric_label(config, source_metric)
                    previous_variable_label = None
                number_format = mean_format if _metric_key(source_metric) == "mean" else "0%"
                for block in blocks:
                    bench_raw_col = layout.product_cols[block["benchmark_pos"]]
                    cand_raw_col = layout.product_cols[block["candidate_pos"]]
                    target.cell(output_row, block["benchmark"]).value = source_sheet.cell(source_row, bench_raw_col).value
                    target.cell(output_row, block["candidate"]).value = source_sheet.cell(source_row, cand_raw_col).value
                    bench_letter = get_column_letter(block["benchmark"])
                    cand_letter = get_column_letter(block["candidate"])
                    if block["delta"] is not None:
                        target.cell(output_row, block["delta"]).value = f"={cand_letter}{output_row}-{bench_letter}{output_row}"
                    value_columns = [block["benchmark"], block["candidate"]]
                    if block["delta"] is not None:
                        value_columns.append(block["delta"])
                    for col in value_columns:
                        target.cell(output_row, col).number_format = number_format
                        target.cell(output_row, col).alignment = Alignment(horizontal="right")
                    fill = _benchmark_fill(
                        workbooks[file_index],
                        source_sheet,
                        layout,
                        source_row,
                        cand_raw_col,
                        block["benchmark_pos"],
                        product_keys[file_index],
                        (),
                        direct_value_fill=True,
                    )
                    if fill is not None:
                        target.cell(output_row, block["candidate"]).fill = fill
                        if block["delta"] is not None:
                            target.cell(output_row, block["delta"]).fill = copy(fill)
                _format_variable_label(target, output_row)
                summary_question_rows.setdefault(question_id.casefold(), []).append((output_row, source_metric))
                output_row += 1
                written_data_rows += 1
        target.auto_filter.ref = f"A5:{get_column_letter(max_col)}{output_row - 1}"
        fake_blocks = {
            pair_index: {
                "value": block["candidate"],
                "benchmark": block["benchmark"],
                "gaps": [],
                "spacer": None,
            }
            for pair_index, block in enumerate(blocks)
        }
        paired_product_names = {}
        for pair_index, block in enumerate(blocks):
            candidate_name = _text(target.cell(3, block["candidate"]).value)
            benchmark_name = _text(target.cell(3, block["benchmark"]).value)
            paired_product_names[pair_index] = f"{candidate_name}\nvs. {benchmark_name}"
        summary_records.append({
            "target": target,
            "split": split_names[file_index],
            "active_benchmarks": (-1,),
            "benchmark_labels": {-1: "Paired benchmarks"},
            "product_order": tuple(range(len(blocks))),
            "column_blocks": fake_blocks,
            "question_rows": summary_question_rows,
            "source_workbook": workbooks[file_index],
            "product_names": paired_product_names,
        })

    total_index = max(range(len(data_sheets)), key=lambda i: sum(count or 0 for count in _signature(data_sheets[i], layouts[i])))
    summary_sheet_names, summary_detail_sheet = _create_paired_summary(
        output,
        used_names,
        summary_records,
        question_rows,
        summary_scope,
        split_names[total_index],
        summary_metric_strategy,
        include_summary_details,
    )
    if include_screeners:
        screener = next((sheet for sheet in workbooks[total_index].worksheets if "screener" in sheet.title.casefold()), None)
        if screener is not None:
            _copy_sheet(screener, output.create_sheet("Screeners"))
    output.calculation.fullCalcOnLoad = True
    output.calculation.forceFullCalc = True
    output.calculation.calcMode = "auto"
    buffer = BytesIO()
    from .english_output import finalize
    finalize(output, include_deltas)
    output.save(buffer)
    scanned = sum(_source_metric_row_count(sheet, layout) for sheet, layout in zip(data_sheets, layouts))
    return buffer.getvalue(), {
        "mode": "Nouveau projet intelligent",
        "test_type": "Paired",
        "splits": list(split_names),
        "questions": len(selected_questions),
        "questions_deleted": sum(not bool(row.get("Keep")) or row.get("Type") == "Delete" for row in question_rows),
        "data_rows_written": written_data_rows,
        "source_metric_rows_scanned": scanned,
        "metric_rows_filtered": max(0, scanned - written_data_rows),
        "technical_columns_ignored": sum(max(0, sheet.max_column - 2 - len(layout.product_cols)) for sheet, layout in zip(data_sheets, layouts)),
        "products": max(product_counts),
        "products_by_split": dict(zip(split_names, product_counts)),
        "pairs": max(product_counts) // 2,
        "pairs_by_split": pairs_by_split,
        "benchmarks": "premier produit de chaque paire",
        "paired_swaps": {name: list(values) for name, values in (paired_swaps or {}).items()},
        "include_deltas": bool(include_deltas),
        "include_sections": bool(include_sections),
        "summary_scope": _normal(summary_scope) or "none",
        "summary_sheets": summary_sheet_names,
        "summary_detail_sheet": summary_detail_sheet,
        "summary_metric_strategy": summary_metric_strategy,
        "summary_kpis": len(_selected_summary_questions(question_rows)),
        "mean_decimals": int(mean_decimals),
        "reference_export": raw_files[total_index][0],
        "skipped_stage_questions": skipped_questions,
    }


def _legend_significance_fill(source_workbook, positive: bool, confidence_row: int):
    legend = next((sheet for sheet in source_workbook.worksheets if "legend" in sheet.title.casefold()), None)
    if legend is None:
        colours = {
            (True, 4): "00FF00",
            (False, 4): "FF0000",
            (True, 5): "CCFFCC",
            (False, 5): "FF9A00",
        }
        return PatternFill("solid", fgColor=colours[(positive, confidence_row)])
    return copy(legend.cell(confidence_row, 2 if positive else 3).fill)


def _benchmark_fill(
    source_workbook,
    source_sheet,
    layout,
    source_row: int,
    candidate_raw_col: int,
    benchmark_position: int,
    product_keys: Sequence[str],
    comparison_codes: Sequence[str],
    direct_value_fill: bool = False,
):
    """Return the significance fill for one candidate-vs-benchmark comparison.

    Preserve the colours of the selected G-Sight analysis sheet (normally
    2_TAILED). Never substitute the palette from the DELTA sheet: those colours
    describe gap thresholds, not the selected statistical-significance view.
    For an additional benchmark without a dedicated colour column, interpret
    G-Sight's case-sensitive comparison letters and use the matching legend.
    """
    if direct_value_fill:
        source_fill = source_sheet.cell(source_row, candidate_raw_col).fill
        if _is_significance_fill(source_workbook, source_sheet, source_fill):
            return copy(source_fill)

    benchmark_key = product_keys[benchmark_position]
    if benchmark_key in comparison_codes:
        source_sig_col = _source_significance_col(candidate_raw_col, comparison_codes.index(benchmark_key))
        if source_sig_col <= source_sheet.max_column and source_row <= source_sheet.max_row:
            source_fill = source_sheet.cell(source_row, source_sig_col).fill
            if _is_significance_fill(source_workbook, source_sheet, source_fill):
                return copy(source_fill)
            # The selected G-Sight sheet explicitly says this comparison is not
            # significant. Do not override it with letters or DELTA thresholds.
            return None

    benchmark_raw_col = layout.product_cols[benchmark_position]
    benchmark_code = _text(source_sheet.cell(layout.product_header_row, benchmark_raw_col).value)
    candidate_code = _text(source_sheet.cell(layout.product_header_row, candidate_raw_col).value)
    candidate_letters = _text(source_sheet.cell(source_row, candidate_raw_col + 1).value)
    benchmark_letters = _text(source_sheet.cell(source_row, benchmark_raw_col + 1).value)
    exact = bool(benchmark_code and benchmark_code in candidate_letters)
    lower = bool(benchmark_code and benchmark_code.casefold() in candidate_letters.casefold())
    if not lower:
        # G-Sight stores each pair once. When the benchmark appears after the
        # candidate, the comparison letter is attached to the benchmark cell.
        exact = bool(candidate_code and candidate_code in benchmark_letters)
        lower = bool(candidate_code and candidate_code.casefold() in benchmark_letters.casefold())
    if not exact and not lower:
        return None
    candidate_value = source_sheet.cell(source_row, candidate_raw_col).value
    benchmark_value = source_sheet.cell(source_row, benchmark_raw_col).value
    if not isinstance(candidate_value, (int, float)) or not isinstance(benchmark_value, (int, float)):
        return None
    # G-Sight uses uppercase comparison letters for 95% and lowercase for 90%.
    confidence_row = 4 if exact else 5
    return _legend_significance_fill(source_workbook, candidate_value > benchmark_value, confidence_row)


def _fill_key(fill) -> tuple:
    colour = fill.fgColor
    return (
        fill.patternType,
        colour.type,
        colour.rgb,
        colour.indexed,
        colour.theme,
        colour.tint,
    )


def _is_significance_fill(source_workbook, source_sheet, fill) -> bool:
    """True only for colours belonging to the selected G-Sight analysis."""
    if fill is None or not fill.patternType:
        return False
    legend = next((sheet for sheet in source_workbook.worksheets if "legend" in sheet.title.casefold()), None)
    title = _normal(source_sheet.title)
    if "1 tailed" in title:
        legend_rows = (9,)
    elif "delta" in title:
        legend_rows = (13, 14)
    else:
        legend_rows = (3, 4, 5)
    if legend is not None:
        return any(
            legend.cell(row, col).fill.patternType
            and _fill_key(fill) == _fill_key(legend.cell(row, col).fill)
            for row in legend_rows
            for col in (2, 3)
        )
    rgb = (fill.fgColor.rgb or "").upper()[-6:]
    palettes = {
        "1_tailed": {"21A625", "A62521"},
        "delta": {"21A625", "A62521", "4BCB4F", "EB3933"},
        "2_tailed": {"008000", "800000", "00FF00", "FF0000", "CCFFCC", "FF9A00"},
    }
    palette = palettes["1_tailed" if "1 tailed" in title else "delta" if "delta" in title else "2_tailed"]
    return rgb in palette


def _summary_confidence(source_workbook, fill) -> int | None:
    """Return the significance confidence without inferring win/loss direction.

    G-Sight can use either the 2-tailed colours (legend rows 3-5) or the
    dedicated DELTA colours (rows 13-14). The colour tells us whether the
    difference is significant and at which level; the candidate and benchmark
    values themselves must decide whether the KPI is a win or a loss.
    """
    if fill is None or not fill.patternType:
        return None
    legend = next((sheet for sheet in source_workbook.worksheets if "legend" in sheet.title.casefold()), None)
    if legend is not None:
        # 99% and 95% both use the solid 95% summary arrow. DELTA exposes
        # equivalent p-value thresholds at 5% and 10% on rows 14 and 13.
        for confidence, rows in ((95, (3, 4, 14)), (90, (5, 13))):
            for row in rows:
                for col in (2, 3):
                    legend_fill = legend.cell(row, col).fill
                    if legend_fill.patternType and _fill_key(fill) == _fill_key(legend_fill):
                        return confidence
    rgb = (fill.fgColor.rgb or "").upper()[-6:]
    return {
        "008000": 95,
        "800000": 95,
        "00FF00": 95,
        "FF0000": 95,
        "4BCB4F": 95,
        "EB3933": 95,
        "97E39C": 95,
        "F5B7AD": 95,
        "CCFFCC": 90,
        "FF9A00": 90,
        "FF9900": 90,
        "21A625": 90,
        "A62521": 90,
        "C7F0CC": 90,
        "FADBD4": 90,
    }.get(rgb)


def _summary_higher_is_better(metric: str, config: dict | None = None) -> bool | None:
    """Return the favourable direction after applying the CMI scale rule."""
    key = _normal(metric)
    mode = _normal((config or {}).get("Sens favorable", "Automatique"))
    if mode in {"neutre", "neutral", "aucun"}:
        return None
    is_pole = any(value in key for value in ("too weak", "too strong", "trop faible", "trop fort"))
    is_jar = "just" in key or "ideal" in key or "parfait" in key or bool(re.match(r"^3\s", key))
    if mode in {"ideal au centre", "centre", "center", "ideal"}:
        if is_jar:
            return True
        if is_pole:
            return False
        return None
    higher = not (
        "bottom" in key
        or is_pole
    )
    if mode in {"plus bas", "lower", "low"}:
        return not higher
    return higher


def _summary_signal_from_values(
    source_workbook,
    fill,
    metric: str,
    candidate_value,
    benchmark_value,
    config: dict | None = None,
) -> str | None:
    confidence = _summary_confidence(source_workbook, fill)
    if confidence is None:
        return None
    if (
        not isinstance(candidate_value, (int, float))
        or isinstance(candidate_value, bool)
        or not isinstance(benchmark_value, (int, float))
        or isinstance(benchmark_value, bool)
        or candidate_value == benchmark_value
    ):
        return None
    higher_is_better = _summary_higher_is_better(metric, config)
    if higher_is_better is None:
        return None
    favourable = candidate_value > benchmark_value
    if not higher_is_better:
        favourable = not favourable
    if favourable:
        return "▲" if confidence == 95 else "△"
    return "▼" if confidence == 95 else "▽"


def _summary_metric_rank(metric: str) -> int:
    key = _metric_key(metric)
    if key == "mean":
        return 0
    if key == "top_2":
        return 1
    if key == "top_1":
        return 2
    if "just" in key or re.match(r"^3\s", _normal(metric)):
        return 2
    if key == "bottom_2":
        return 3
    return 4


def _pick_summary_signal(signals: Sequence[tuple[str, str]]) -> str:
    if not signals:
        return "="
    ordered = sorted(signals, key=lambda item: (_summary_metric_rank(item[0]), item[1] in {"△", "▽"}))
    directions = {"up" if symbol in {"▲", "△"} else "down" for _, symbol in ordered}
    if len(directions) > 1:
        # Conflicting metrics: use the most decision-relevant retained metric
        # (Mean, then T2B, Top Box, JAR/B2B) instead of overstating consensus.
        return ordered[0][1]
    solid = next((symbol for _, symbol in ordered if symbol in {"▲", "▼"}), None)
    return solid or ordered[0][1]


def _pick_summary_decision(entries: Sequence[dict], strategy: str) -> dict:
    """Choose one transparent summary decision from all retained metrics."""
    ordered_all = sorted(entries, key=lambda item: _summary_metric_rank(item["metric"]))
    if ordered_all and not any(item.get("direction_defined", True) for item in ordered_all):
        return {"symbol": "—", "chosen": ordered_all[0], "conflict": False,
                "reason": "Aucun sens favorable n’est défini pour cette question"}
    significant = [item for item in ordered_all if item.get("signal")]
    normalized_strategy = _normal(strategy).replace(" ", "_")
    if normalized_strategy in {"primary", "primary_metric", "metric_principale"}:
        chosen = ordered_all[0] if ordered_all else None
        return {
            "symbol": chosen.get("signal") if chosen and chosen.get("signal") else "=",
            "chosen": chosen,
            "conflict": False,
            "reason": "Métrique principale uniquement",
        }
    if not significant:
        return {"symbol": "=", "chosen": ordered_all[0] if ordered_all else None, "conflict": False,
                "reason": "Aucune différence significative sur les métriques retenues"}
    directions = {"win" if item["signal"] in {"▲", "△"} else "loss" for item in significant}
    if normalized_strategy in {"consensus", "accord"} and len(directions) > 1:
        return {"symbol": "±", "chosen": None, "conflict": True,
                "reason": "Les métriques significatives donnent des directions opposées"}
    if len(directions) > 1:
        chosen = significant[0]
        return {"symbol": chosen["signal"], "chosen": chosen, "conflict": True,
                "reason": f"Résultats partagés — priorité donnée à {chosen['metric']}"}
    solid = next((item for item in significant if item["signal"] in {"▲", "▼"}), None)
    chosen = solid or significant[0]
    return {"symbol": chosen["signal"], "chosen": chosen, "conflict": False,
            "reason": f"Résultat significatif retenu sur {chosen['metric']}"}


def _selected_summary_questions(question_rows: Sequence[dict]) -> list[dict]:
    selected = [
        row for row in question_rows
        if bool(row.get("Keep"))
        and row.get("Type") != "Delete"
        and bool(row.get("KPI Summary"))
    ]
    selected.sort(key=lambda row: (int(row.get("Order", 9999)), _text(row.get("Question ID"))))
    return selected


def _summary_product_name(record: dict, position: int) -> str:
    if position in record.get("product_names", {}):
        return record["product_names"][position]
    target = record["target"]
    value_col = record["column_blocks"][position]["value"]
    name = _text(target.cell(3, value_col).value) or f"Product {position + 1}"
    subtitle = _text(target.cell(5, value_col).value)
    return f"{name}\n{subtitle}" if subtitle else name


def _write_summary_block(
    sheet,
    start_row: int,
    record: dict,
    benchmark_position: int,
    kpis: Sequence[dict],
    summary_metric_strategy: str,
    detail_rows: list[dict],
) -> int:
    max_col = 1 + len(kpis)
    title_fill = PatternFill("solid", fgColor="4F8ED8")
    subtitle_fill = PatternFill("solid", fgColor="E7E7E7")
    header_fill = PatternFill("solid", fgColor="000000")
    neutral_fill = PatternFill("solid", fgColor="EFEFEF")
    fills = {
        "▲": PatternFill("solid", fgColor="97E39C"),
        "△": PatternFill("solid", fgColor="C7F0CC"),
        "▼": PatternFill("solid", fgColor="F5B7AD"),
        "▽": PatternFill("solid", fgColor="FADBD4"),
        "±": PatternFill("solid", fgColor="FFF2CC"),
        "—": neutral_fill,
        "=": neutral_fill,
    }
    fonts = {
        "▲": Font(color="008A16", bold=True, size=14),
        "△": Font(color="008A16", bold=True, size=14),
        "▼": Font(color="C00000", bold=True, size=14),
        "▽": Font(color="C00000", bold=True, size=14),
        "±": Font(color="9C6500", bold=True, size=14),
        "—": Font(color="7A7A7A", bold=True, size=12),
        "=": Font(color="7A7A7A", bold=True, size=12),
    }
    thin = Side(style="thin", color="C9C9C9")

    benchmark_label = record["benchmark_labels"].get(benchmark_position, f"Benchmark {benchmark_position + 1}")
    sheet.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=max_col)
    title = sheet.cell(start_row, 1)
    title.value = f"{record['split'].upper()} SAMPLE VS. {benchmark_label.upper()}"
    title.fill = title_fill
    title.font = Font(color="FFFFFF", bold=True, italic=True, size=15)
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[start_row].height = 24

    product_positions = [position for position in record["product_order"] if position not in record["active_benchmarks"]]
    counts = []
    for position in product_positions:
        value_col = record["column_blocks"][position]["value"]
        count = _count_from_header(record["target"].cell(4, value_col).value)
        if count is not None:
            counts.append(count)
    base_note = ""
    if counts:
        base_note = f" | n~{min(counts)}" + (f"-{max(counts)}" if min(counts) != max(counts) else "") + " per product"
    sheet.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=max_col)
    subtitle = sheet.cell(start_row + 1, 1)
    retained_metric_labels = []
    for metric in STANDARD_METRICS:
        if any(bool(config.get(metric)) for config in kpis):
            retained_metric_labels.append(metric)
    metric_note = " / ".join(retained_metric_labels[:4]) or "selected metrics"
    subtitle.value = f"Significant result across retained metrics ({metric_note})" + base_note
    subtitle.fill = subtitle_fill
    subtitle.font = Font(color="666666", italic=True, size=10)
    subtitle.alignment = Alignment(vertical="center")

    header_row = start_row + 3
    headers = ["Product"] + [
        _text(config.get("Summary label")) or _summary_label(_text(config.get("Display label")))
        for config in kpis
    ]
    for col, value in enumerate(headers, 1):
        cell = sheet.cell(header_row, col)
        cell.value = value
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.row_dimensions[header_row].height = 48

    question_rows = record["question_rows"]
    active_benchmarks = list(record["active_benchmarks"])
    benchmark_offset = active_benchmarks.index(benchmark_position)
    for row_offset, position in enumerate(product_positions, 1):
        row = header_row + row_offset
        product_cell = sheet.cell(row, 1)
        product_cell.value = _summary_product_name(record, position)
        product_cell.fill = PatternFill("solid", fgColor="FFFFFF")
        product_cell.font = Font(bold=True, size=9)
        product_cell.alignment = Alignment(vertical="center", wrap_text=True)
        product_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.row_dimensions[row].height = 34
        block = record["column_blocks"][position]
        signal_col = block["gaps"][benchmark_offset] if block["gaps"] else block["value"]
        benchmark_value_col = (
            block.get("benchmark")
            if benchmark_position == -1
            else record["column_blocks"][benchmark_position]["value"]
        )
        for kpi_offset, config in enumerate(kpis, 2):
            question_id = _text(config.get("Question ID")).casefold()
            entries = []
            for output_row, metric in question_rows.get(question_id, ()):
                candidate_value = record["target"].cell(output_row, block["value"]).value
                benchmark_value = record["target"].cell(output_row, benchmark_value_col).value
                confidence = _summary_confidence(
                    record["source_workbook"],
                    record["target"].cell(output_row, signal_col).fill,
                )
                signal = _summary_signal_from_values(
                    record["source_workbook"],
                    record["target"].cell(output_row, signal_col).fill,
                    metric,
                    candidate_value,
                    benchmark_value,
                    config,
                )
                entries.append({
                    "metric": metric,
                    "signal": signal,
                    "confidence": confidence,
                    "candidate": candidate_value,
                    "benchmark": benchmark_value,
                    "direction_defined": _summary_higher_is_better(metric, config) is not None,
                })
            decision = _pick_summary_decision(entries, summary_metric_strategy)
            symbol = decision["symbol"]
            cell = sheet.cell(row, kpi_offset)
            cell.value = symbol
            cell.fill = fills[symbol]
            cell.font = fonts[symbol]
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            chosen = decision.get("chosen")
            product_name = _summary_product_name(record, position).replace("\n", " · ")
            candidate_value = chosen.get("candidate") if chosen else None
            benchmark_value = chosen.get("benchmark") if chosen else None
            gap = (
                candidate_value - benchmark_value
                if isinstance(candidate_value, (int, float)) and isinstance(benchmark_value, (int, float))
                else None
            )
            significant_comparisons = " · ".join(
                f"{entry['metric']}: {entry['candidate']} vs {entry['benchmark']} ({entry['signal']})"
                for entry in entries if entry.get("signal")
            ) or "None"
            detail_rows.append({
                "Split": record["split"],
                "Benchmark": benchmark_label,
                "Produit": product_name,
                "KPI": _text(config.get("Summary label")) or _summary_label(_text(config.get("Display label"))),
                "Règle favorable": _text(config.get("Sens favorable")) or "Automatique",
                "Méthode": summary_metric_strategy,
                "Métrique retenue": chosen.get("metric") if chosen else "Multiple metrics",
                "Comparaisons significatives": significant_comparisons,
                "Valeur produit": candidate_value,
                "Valeur benchmark": benchmark_value,
                "Écart": gap,
                "Seuil": f"{chosen.get('confidence')}%" if chosen and chosen.get("confidence") else "Non significatif",
                "Résultat": symbol,
                "Explication": decision["reason"],
            })

    legend_row = header_row + len(product_positions) + 2
    sheet.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=max_col)
    legend = sheet.cell(legend_row, 1)
    legend.value = "▲ win 95%   △ win 90%   = parity   ± résultats partagés   — non interprété   ▽ loss 90%   ▼ loss 95%"
    legend.font = Font(color="666666", italic=True, size=9)
    legend.alignment = Alignment(horizontal="left")
    return legend_row + 3


def _create_summary_details(output, used_names: set[str], detail_rows: Sequence[dict]) -> str | None:
    if not detail_rows:
        return None
    sheet = output.create_sheet(_safe_sheet_name("KPI Details", used_names))
    headers = list(detail_rows[0])
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(1, col, header)
        cell.fill = PatternFill("solid", fgColor="3E356B")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, values in enumerate(detail_rows, 2):
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row_index, col, values.get(header))
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=header in {"Produit", "KPI", "Comparaisons significatives", "Explication"},
            )
            if header in {"Valeur produit", "Valeur benchmark", "Écart"}:
                cell.number_format = "0.00"
    widths = {
        "Split": 20, "Benchmark": 22, "Produit": 36, "KPI": 28, "Règle favorable": 18,
        "Méthode": 18, "Métrique retenue": 22, "Valeur produit": 14, "Valeur benchmark": 16,
        "Comparaisons significatives": 48, "Écart": 12, "Seuil": 15, "Résultat": 10, "Explication": 46,
    }
    for col, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(col)].width = widths.get(header, 18)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    return sheet.title


def _create_monadic_summaries(
    output,
    used_names: set[str],
    records: Sequence[dict],
    question_rows: Sequence[dict],
    summary_scope: str,
    reference_split_name: str,
    summary_metric_strategy: str,
    include_summary_details: bool,
) -> tuple[list[str], str | None]:
    normalized_scope = _normal(summary_scope)
    if normalized_scope in {"", "none", "no", "off"}:
        return [], None
    kpis = _selected_summary_questions(question_rows)
    if not kpis:
        return [], None
    if normalized_scope not in {"total", "all"}:
        raise TemplyfierError("Périmètre KPI Summary inconnu.")
    selected_records = list(records)
    if normalized_scope == "total":
        reference_key = _normal(reference_split_name)
        selected_records = [record for record in records if _normal(record["split"]) == reference_key]

    benchmark_order = []
    for record in selected_records:
        for benchmark_position in record["active_benchmarks"]:
            if benchmark_position not in benchmark_order:
                benchmark_order.append(benchmark_position)
    created = []
    created_sheets = []
    detail_rows: list[dict] = []
    for benchmark_position in benchmark_order:
        matching = [record for record in selected_records if benchmark_position in record["active_benchmarks"]]
        seen_splits = set()
        unique_matching = []
        for record in matching:
            split_key = _normal(record["split"])
            if split_key in seen_splits:
                continue
            seen_splits.add(split_key)
            unique_matching.append(record)
        matching = unique_matching
        if not matching:
            continue
        label = matching[0]["benchmark_labels"].get(benchmark_position, f"B{benchmark_position + 1}")
        requested_name = "KPI Summary" if len(benchmark_order) == 1 else f"KPI Summary vs {label}"
        sheet = output.create_sheet(_safe_sheet_name(requested_name, used_names))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B4"
        sheet.column_dimensions["A"].width = 38
        for col in range(2, 2 + len(kpis)):
            sheet.column_dimensions[get_column_letter(col)].width = 16
        next_row = 1
        for record in matching:
            next_row = _write_summary_block(
                sheet, next_row, record, benchmark_position, kpis,
                summary_metric_strategy, detail_rows,
            )
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        created.append(sheet.title)
        created_sheets.append(sheet)
    for sheet in reversed(created_sheets):
        output.move_sheet(sheet, offset=-output.index(sheet))
    detail_name = _create_summary_details(output, used_names, detail_rows) if include_summary_details else None
    if detail_name:
        detail_sheet = output[detail_name]
        output.move_sheet(detail_sheet, offset=-output.index(detail_sheet) + len(created))
    return created, detail_name


def _create_paired_summary(
    output,
    used_names: set[str],
    records: Sequence[dict],
    question_rows: Sequence[dict],
    summary_scope: str,
    reference_split_name: str,
    summary_metric_strategy: str,
    include_summary_details: bool,
) -> tuple[list[str], str | None]:
    normalized_scope = _normal(summary_scope)
    if normalized_scope in {"", "none", "no", "off"}:
        return [], None
    if normalized_scope not in {"total", "all"}:
        raise TemplyfierError("Périmètre KPI Summary inconnu.")
    kpis = _selected_summary_questions(question_rows)
    if not kpis:
        return [], None
    selected_records = list(records)
    if normalized_scope == "total":
        reference_key = _normal(reference_split_name)
        selected_records = [record for record in records if _normal(record["split"]) == reference_key]
    if not selected_records:
        return [], None
    sheet = output.create_sheet(_safe_sheet_name("KPI Summary", used_names))
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 42
    for col in range(2, 2 + len(kpis)):
        sheet.column_dimensions[get_column_letter(col)].width = 16
    next_row = 1
    detail_rows: list[dict] = []
    for record in selected_records:
        next_row = _write_summary_block(
            sheet, next_row, record, -1, kpis,
            summary_metric_strategy, detail_rows,
        )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output.move_sheet(sheet, offset=-output.index(sheet))
    detail_name = _create_summary_details(output, used_names, detail_rows) if include_summary_details else None
    if detail_name:
        detail_sheet = output[detail_name]
        output.move_sheet(detail_sheet, offset=-output.index(detail_sheet) + 1)
    return [sheet.title], detail_name


def build_smart_toplines(
    raw_files: Sequence[tuple[str, object]],
    question_rows: Sequence[dict],
    *,
    split_names: Sequence[str],
    benchmark_positions: Sequence[int],
    standard_metrics: Sequence[str],
    include_screeners: bool = True,
    test_type: str = "Monadic",
    mean_decimals: int = 2,
    paired_swaps: dict[str, Sequence[int]] | None = None,
    include_deltas: bool = True,
    include_sections: bool = True,
    benchmark_sheet_mode: str = "combined",
    benchmark_labels: Sequence[str] | None = None,
    output_sheet_order: str = "split_first",
    show_monadic_gaps: bool = True,
    product_labels: Sequence[str] | None = None,
    product_subtitles: Sequence[str] | None = None,
    highlight_benchmarks: bool = True,
    summary_scope: str = "none",
    summary_metric_strategy: str = "priority",
    include_summary_details: bool = True,
) -> tuple[bytes, dict]:
    if not raw_files:
        raise TemplyfierError("Ajoute au moins un export G-Sight.")
    normalized_test_type = _normal(test_type)
    columns_mode = benchmark_sheet_mode in {'auto_columns','benchmark_columns'}
    if columns_mode:
        benchmark_sheet_mode = 'auto_exports' if benchmark_sheet_mode == 'auto_columns' else 'separate'
    if normalized_test_type.startswith("paired"):
        return _build_paired_toplines(
            raw_files,
            question_rows,
            split_names=split_names,
            standard_metrics=standard_metrics,
            include_screeners=include_screeners,
            mean_decimals=mean_decimals,
            paired_swaps=paired_swaps,
            include_deltas=include_deltas,
            include_sections=include_sections,
            summary_scope=summary_scope,
            summary_metric_strategy=summary_metric_strategy,
            include_summary_details=include_summary_details,
        )
    if not normalized_test_type.startswith("monadic"):
        raise TemplyfierError("Type de test non reconnu : choisis Paired ou Monadic.")
    workbooks = [_workbook(source) for _, source in raw_files]
    virtual_inputs = []
    for file_index, ((filename, _), workbook) in enumerate(zip(raw_files, workbooks)):
        for sheet in _embedded_data_sheets(workbook):
            virtual_inputs.append((file_index, filename, sheet))

    if len(virtual_inputs) == len(raw_files):
        if len(split_names) != len(raw_files):
            raise TemplyfierError("Le nombre de noms de splits est incorrect.")
        output_split_names = list(split_names)
    elif len(split_names) == len(virtual_inputs):
        output_split_names = list(split_names)
    elif len(split_names) == len(raw_files):
        virtual_stages = tuple(dict.fromkeys(_stage_name(sheet) for _, _, sheet in virtual_inputs))
        virtual_study_format = _study_format(
            [filename for filename, _ in raw_files],
            virtual_stages,
            max((sheet.max_row for _, _, sheet in virtual_inputs), default=0),
        )
        if virtual_study_format == "HUT / in-use":
            # In consolidated HUT exports, Stage None/USE often represents the
            # benchmark configuration, not a consumer split or touchpoint.
            output_split_names = [_detected_split_label(sheet) for _, _, sheet in virtual_inputs]
        else:
            output_split_names = [
                f"{_stage_name(sheet)} · {_detected_split_label(sheet)}"
                for _, _, sheet in virtual_inputs
            ]
    else:
        raise TemplyfierError("Le nombre de noms de splits est incorrect.")

    data_sheets = [sheet for _, _, sheet in virtual_inputs]
    source_workbooks = [workbooks[file_index] for file_index, _, _ in virtual_inputs]
    source_filenames = [filename for _, filename, _ in virtual_inputs]
    layouts = [_active_layout(sheet, detect_layout(sheet)) for sheet in data_sheets]
    comparison_codes_by_sheet = [_comparison_codes(sheet, layout) for sheet, layout in zip(data_sheets, layouts)]
    product_names = [_product_names(sheet, layout) for sheet, layout in zip(data_sheets, layouts)]
    product_keys = [_product_keys(sheet, layout) for sheet, layout in zip(data_sheets, layouts)]
    normalized_sheet_mode = _normal(benchmark_sheet_mode).replace("_", " ")
    automatic_exports = normalized_sheet_mode in {"auto exports", "automatic exports", "auto"}
    if normalized_sheet_mode not in {"combined", "separate", "separate sheets", "auto exports", "automatic exports", "auto"}:
        raise TemplyfierError("Organisation des benchmarks inconnue.")
    canonical_keys = product_keys[0]
    same_product_sets = all(len(keys) == len(canonical_keys) and set(keys) == set(canonical_keys) for keys in product_keys)
    if not same_product_sets:
        raise TemplyfierError("Le plan produits n’est pas identique dans tous les exports.")
    if not automatic_exports and any(keys != canonical_keys for keys in product_keys[1:]):
        raise TemplyfierError(
            "L’ordre des produits diffère entre les exports. Active la détection automatique par fichier benchmark."
        )

    product_count = len(canonical_keys)
    benchmarks = tuple(int(pos) for pos in benchmark_positions)
    if automatic_exports and not benchmarks:
        detected_codes = list(dict.fromkeys(code for codes in comparison_codes_by_sheet for code in codes))
        benchmarks = tuple(canonical_keys.index(code) for code in detected_codes if code in canonical_keys)
    if not benchmarks or any(pos < 0 or pos >= product_count for pos in benchmarks):
        raise TemplyfierError("Benchmark invalide.")
    separate_benchmarks = automatic_exports or (len(benchmarks) > 1 and normalized_sheet_mode.startswith("separate"))
    short_benchmark_labels = [_text(label) for label in (benchmark_labels or ())]
    if short_benchmark_labels and len(short_benchmark_labels) != len(benchmarks):
        raise TemplyfierError("Le nombre de noms courts ne correspond pas au nombre de benchmarks.")
    clean_product_labels = [_text(label) for label in (product_labels or ())]
    clean_product_subtitles = [_text(label) for label in (product_subtitles or ())]
    if clean_product_labels and len(clean_product_labels) != product_count:
        raise TemplyfierError("Le nombre de noms produits ne correspond pas au plan produits.")
    if clean_product_subtitles and len(clean_product_subtitles) != product_count:
        raise TemplyfierError("Le nombre de sous-titres produits ne correspond pas au plan produits.")
    normalized_output_order = _normal(output_sheet_order).replace("_", " ")
    if normalized_output_order not in {"split first", "benchmark first"}:
        raise TemplyfierError("Ordre des onglets inconnu.")

    selected_questions = _selected_question_rows(question_rows)

    output = Workbook()
    output.remove(output.active)
    used_names: set[str] = set()
    total_index = max(range(len(data_sheets)), key=lambda i: sum(count or 0 for count in _signature(data_sheets[i], layouts[i])))
    written_data_rows = 0
    skipped_stage_questions = 0
    mean_format = _mean_number_format(mean_decimals)

    generated_sheet_names = []
    missing_benchmark_exports = []
    summary_records = []
    benchmark_label_map = {
        position: (
            short_benchmark_labels[offset]
            if offset < len(short_benchmark_labels) and short_benchmark_labels[offset]
            else clean_product_labels[position]
            if clean_product_labels
            else product_names[0][position]
        )
        for offset, position in enumerate(benchmarks)
    }
    output_jobs = []
    if automatic_exports:
        split_groups: dict[str, list[int]] = {}
        split_display_names: dict[str, str] = {}
        for index, split_name in enumerate(output_split_names):
            split_key = _normal(split_name)
            split_groups.setdefault(split_key, []).append(index)
            split_display_names.setdefault(split_key, split_name)
        combinations = (
            (
                (split_key, indexes, label_offset, benchmark_position)
                for label_offset, benchmark_position in enumerate(benchmarks)
                for split_key, indexes in split_groups.items()
            )
            if normalized_output_order == "benchmark first"
            else (
                (split_key, indexes, label_offset, benchmark_position)
                for split_key, indexes in split_groups.items()
                for label_offset, benchmark_position in enumerate(benchmarks)
            )
        )
        for split_key, indexes, label_offset, benchmark_position in combinations:
            benchmark_key = canonical_keys[benchmark_position]
            matches = [index for index in indexes if benchmark_key in comparison_codes_by_sheet[index]]
            if len(matches) > 1:
                raise TemplyfierError(
                    f"Plusieurs exports correspondent à {split_display_names[split_key]} vs {benchmark_key}. "
                    "Garde un seul fichier pour cette combinaison."
                )
            if not matches:
                missing_benchmark_exports.append(f"{split_display_names[split_key]} vs {benchmark_key}")
                continue
            output_jobs.append((matches[0], label_offset, (benchmark_position,), split_display_names[split_key]))
    else:
        if separate_benchmarks and normalized_output_order == "benchmark first":
            for label_offset, position in enumerate(benchmarks):
                for file_index in range(len(data_sheets)):
                    output_jobs.append((file_index, label_offset, (position,), output_split_names[file_index]))
        else:
            for file_index in range(len(data_sheets)):
                variants = [(None, benchmarks)]
                if separate_benchmarks:
                    variants = [(offset, (position,)) for offset, position in enumerate(benchmarks)]
                for label_offset, active_benchmarks in variants:
                    output_jobs.append((file_index, label_offset, active_benchmarks, output_split_names[file_index]))

    if not output_jobs:
        raise TemplyfierError("Aucune combinaison split × benchmark complète n’a été détectée.")

    for file_index, label_offset, active_benchmarks, split_display_name in output_jobs:
            filename = source_filenames[file_index]
            source_sheet = data_sheets[file_index]
            layout = layouts[file_index]
            source_positions = {
                canonical_position: product_keys[file_index].index(key)
                for canonical_position, key in enumerate(canonical_keys)
            }
            product_order = list(active_benchmarks) + [pos for pos in range(product_count) if pos not in active_benchmarks]
            column_blocks: dict[int, dict] = {}
            current_col = 3
            for pos in product_order:
                if pos in active_benchmarks:
                    column_blocks[pos] = {"value": current_col, "gaps": [], "spacer": current_col + 1}
                    current_col += 2
                elif not show_monadic_gaps:
                    column_blocks[pos] = {"value": current_col, "gaps": [], "spacer": current_col + 1}
                    current_col += 2
                else:
                    gaps = list(range(current_col + 1, current_col + 1 + len(active_benchmarks)))
                    column_blocks[pos] = {"value": current_col, "gaps": gaps, "spacer": gaps[-1] + 1}
                    current_col = gaps[-1] + 2
            max_col = current_col - 1

            requested_name = split_display_name
            if label_offset is not None:
                default_source_position = source_positions[active_benchmarks[0]]
                default_label = product_names[file_index][default_source_position]
                label = short_benchmark_labels[label_offset] if short_benchmark_labels else default_label
                requested_name = f"{requested_name} vs {label}"
            target = output.create_sheet(_safe_sheet_name(requested_name, used_names))
            generated_sheet_names.append(target.title)
            _style_sheet(target, max_col)
            for pos in product_order:
                block = column_blocks[pos]
                source_position = source_positions[pos]
                raw_col = layout.product_cols[source_position]
                value_col = block["value"]
                target.cell(3, value_col).value = (
                    clean_product_labels[pos]
                    if clean_product_labels
                    else source_sheet.cell(layout.product_name_row, raw_col).value
                )
                target.cell(4, value_col).value = source_sheet.cell(layout.sample_row, raw_col).value
                target.cell(5, value_col).value = (
                    clean_product_subtitles[pos]
                    if clean_product_subtitles
                    else source_sheet.cell(layout.product_header_row, raw_col).value
                )
                if highlight_benchmarks and pos in benchmarks:
                    for header_row in (3, 4):
                        header_cell = target.cell(header_row, value_col)
                        header_cell.fill = PatternFill("solid", fgColor="D6008F")
                        header_cell.font = Font(color="FFFFFF", bold=True)
                for bench_offset, gap_col in enumerate(block["gaps"]):
                    bench_name = source_sheet.cell(
                        layout.product_header_row,
                        layout.product_cols[source_positions[active_benchmarks[bench_offset]]],
                    ).value
                    target.cell(5, gap_col).value = f"Gap vs {bench_name}"
                target.column_dimensions[get_column_letter(block["spacer"])].width = 3

            output_row = 6
            current_section = None
            previous_variable_label = None
            summary_question_rows: dict[str, list[tuple[int, str]]] = {}
            for config in selected_questions:
                if not _question_applies_to_split(config, split_display_name):
                    continue
                section = _text(config.get("Section")) or "TOPLINES"
                if include_sections and section != current_section:
                    target.cell(output_row, 1).value = section
                    target.merge_cells(start_row=output_row, start_column=1, end_row=output_row, end_column=max_col)
                    section_cell = target.cell(output_row, 1)
                    section_cell.fill = PatternFill("solid", fgColor="D9D3EA")
                    section_cell.font = Font(bold=True, color="2F2952")
                    section_cell.alignment = Alignment(vertical="center")
                    target.row_dimensions[output_row].height = 22
                    output_row += 1
                    current_section = section
                    previous_variable_label = None

                question_id = _text(config.get("Question ID"))
                question_type = _text(config.get("Type"))
                question_metrics = _standard_metrics_for_question(config, standard_metrics)
                source_rows = configured_source_rows(source_sheet, layout, config, standard_metrics)
                if not source_rows:
                    skipped_stage_questions += 1
                    continue

                for metric_index, source_row in enumerate(source_rows):
                    source_metric = _text(source_sheet.cell(source_row, layout.metric_col).value)
                    display_label = _text(config.get("Display label"))
                    metric_label = _text(config.get("Metric label"))
                    if metric_label:
                        group_identity = (_text(config.get("Group ID")) or display_label, section, display_label)
                        target.cell(output_row, 1).value = display_label if group_identity != previous_variable_label else None
                        target.cell(output_row, 2).value = grouped_metric_label(config, source_metric, len(source_rows))
                        previous_variable_label = group_identity
                    elif question_type in {"Attribute", "CATA"}:
                        target.cell(output_row, 1).value = None
                        target.cell(output_row, 2).value = (display_label if len(source_rows) == 1 else f"{display_label} · {clean_metric_label(config, source_metric)}")
                        previous_variable_label = None
                    else:
                        target.cell(output_row, 1).value = display_label if metric_index == 0 else None
                        target.cell(output_row, 2).value = clean_metric_label(config, source_metric)
                        previous_variable_label = None

                    is_mean = _metric_key(source_metric) == "mean"
                    number_format = mean_format if is_mean else "0%"
                    for pos in product_order:
                        source_position = source_positions[pos]
                        raw_col = layout.product_cols[source_position]
                        block = column_blocks[pos]
                        value_cell = target.cell(output_row, block["value"])
                        value_cell.value = source_sheet.cell(source_row, raw_col).value
                        value_cell.number_format = number_format
                        value_cell.alignment = Alignment(horizontal="right")
                        if pos not in active_benchmarks:
                            for bench_offset, gap_col in enumerate(block["gaps"]):
                                benchmark_position = active_benchmarks[bench_offset]
                                benchmark_value_col = column_blocks[benchmark_position]["value"]
                                gap_cell = target.cell(output_row, gap_col)
                                gap_cell.value = (
                                    f"={get_column_letter(block['value'])}{output_row}-"
                                    f"{get_column_letter(benchmark_value_col)}{output_row}"
                                )
                                gap_cell.number_format = number_format
                                fill = _benchmark_fill(
                                    source_workbooks[file_index], source_sheet, layout, source_row, raw_col,
                                    source_positions[benchmark_position], product_keys[file_index],
                                    comparison_codes_by_sheet[file_index],
                                )
                                if fill is not None:
                                    value_cell.fill = fill
                                    gap_cell.fill = copy(fill)
                            if not block["gaps"] and active_benchmarks:
                                fill = _benchmark_fill(
                                    source_workbooks[file_index],
                                    source_sheet,
                                    layout,
                                    source_row,
                                    raw_col,
                                    source_positions[active_benchmarks[0]],
                                    product_keys[file_index],
                                    comparison_codes_by_sheet[file_index],
                                )
                                if fill is not None:
                                    value_cell.fill = fill

                    thin = Side(style="thin", color="B8B8C3")
                    spacer_columns = [block["spacer"] for block in column_blocks.values()]
                    for col in range(2, max_col + 1):
                        if col not in spacer_columns:
                            target.cell(output_row, col).border = Border(bottom=thin)
                    _format_variable_label(target, output_row)
                    summary_question_rows.setdefault(question_id.casefold(), []).append((output_row, source_metric))
                    output_row += 1
                    written_data_rows += 1

            target.auto_filter.ref = f"A5:{get_column_letter(max_col)}{output_row - 1}"
            summary_records.append({
                "target": target,
                "split": split_display_name,
                "file_index": file_index,
                "active_benchmarks": tuple(active_benchmarks),
                "benchmark_labels": benchmark_label_map,
                "product_order": tuple(product_order),
                "column_blocks": column_blocks,
                "question_rows": summary_question_rows,
                "source_workbook": source_workbooks[file_index],
            })

    summary_sheet_names, summary_detail_sheet = _create_monadic_summaries(
        output,
        used_names,
        summary_records,
        question_rows,
        summary_scope,
        output_split_names[total_index],
        summary_metric_strategy,
        include_summary_details,
    )

    if include_screeners:
        total_wb = source_workbooks[total_index]
        screener = next((sheet for sheet in total_wb.worksheets if "screener" in sheet.title.casefold()), None)
        if screener is not None:
            _copy_sheet(screener, output.create_sheet("Screeners"))

    panel_layouts, missing_panel_metrics = [], []
    if columns_mode:
        from .benchmark_columns import combine_readings
        generated_sheet_names, panel_layouts, missing_panel_metrics = combine_readings(output, summary_records, _metric_key)
    output.calculation.fullCalcOnLoad = True
    output.calculation.forceFullCalc = True
    output.calculation.calcMode = "auto"
    buffer = BytesIO()
    from .english_output import finalize
    finalize(output, show_monadic_gaps)
    output.save(buffer)
    scanned = sum(_source_metric_row_count(sheet, layout) for sheet, layout in zip(data_sheets, layouts))
    return buffer.getvalue(), {
        "mode": "Nouveau projet intelligent",
        "test_type": "Monadic",
        "splits": generated_sheet_names,
        "source_splits": output_split_names,
        "benchmark_sheet_mode": ('auto_columns' if automatic_exports else 'benchmark_columns') if columns_mode else ("auto_exports" if automatic_exports else ("separate" if separate_benchmarks else "combined")),
        "benchmark_column_panels": panel_layouts,
        "missing_panel_metrics": missing_panel_metrics,
        "missing_benchmark_exports": missing_benchmark_exports,
        "questions": len(selected_questions),
        "questions_deleted": sum(not bool(row.get("Keep")) or row.get("Type") == "Delete" for row in question_rows),
        "data_rows_written": written_data_rows,
        "source_metric_rows_scanned": scanned,
        "metric_rows_filtered": max(0, scanned - written_data_rows),
        "technical_columns_ignored": sum(max(0, sheet.max_column - 2 - len(layout.product_cols)) for sheet, layout in zip(data_sheets, layouts)),
        "products": product_count,
        "benchmarks": [pos + 1 for pos in benchmarks],
        "benchmark_labels": short_benchmark_labels,
        "output_sheet_order": (
            "benchmark_first" if normalized_output_order == "benchmark first" else "split_first"
        ),
        "show_monadic_gaps": bool(show_monadic_gaps),
        "product_labels": clean_product_labels,
        "product_subtitles": clean_product_subtitles,
        "highlight_benchmarks": bool(highlight_benchmarks),
        "summary_scope": _normal(summary_scope) or "none",
        "summary_sheets": summary_sheet_names,
        "summary_detail_sheet": summary_detail_sheet,
        "summary_metric_strategy": summary_metric_strategy,
        "summary_kpis": len(_selected_summary_questions(question_rows)),
        "reference_export": source_filenames[total_index],
        "skipped_stage_questions": skipped_stage_questions,
        "mean_decimals": int(mean_decimals),
        "include_sections": bool(include_sections),
    }
