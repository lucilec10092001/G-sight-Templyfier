"""Translate tool-owned Excel annotations, never survey or user labels."""
from openpyxl.utils import get_column_letter
from .english_catalog import TEXT


def finalize(output, show_gaps=True):
    for sheet in output.worksheets:
        if sheet.title=='KPI Details':
            headings={cell.column:cell.value for cell in sheet[1]}
            controlled={'Règle favorable','Méthode','Seuil','Explication'}
            for col,header in headings.items():
                sheet.cell(1,col).value=TEXT.get(header,header)
                if header in controlled:
                    for row in range(2,sheet.max_row+1):
                        cell=sheet.cell(row,col)
                        if isinstance(cell.value,str):
                            if cell.value in TEXT:cell.value=TEXT[cell.value]
                            else:
                                for original,english in sorted(TEXT.items(),key=lambda pair:-len(pair[0])):
                                    if len(original)>5 and cell.value.startswith(original):
                                        cell.value=english+cell.value[len(original):];break
            if not show_gaps:
                gap=next((col for col,header in headings.items() if header=='Écart'),None)
                if gap:
                    for col in range(gap,sheet.max_column):
                        sheet.column_dimensions[get_column_letter(col)].width=sheet.column_dimensions[get_column_letter(col+1)].width
                    sheet.delete_cols(gap)
                    sheet.auto_filter.ref=f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        elif sheet.title.startswith('KPI Summary'):
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,str) and 'win 95%' in cell.value and 'loss 95%' in cell.value:
                        cell.value=cell.value.replace('résultats partagés','mixed results').replace('non interprété','not interpreted')
