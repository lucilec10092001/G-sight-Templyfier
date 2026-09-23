"""Place benchmark-specific readings side by side without mixing their signals."""
from copy import copy
import re
from openpyxl.comments import Comment
from openpyxl.formula import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from .core import TemplyfierError, _safe_sheet_name


def _row_plan(record, metric_key):
    sheet=record['target']
    rows={}
    for qid, entries in record['question_rows'].items():
        occurrences={}
        for row, metric in entries:
            identity=(qid,metric_key(metric))
            occurrences[identity]=occurrences.get(identity,0)+1
            rows[row]=('metric',*identity,occurrences[identity])
    sections={}
    for row in range(6,sheet.max_row+1):
        if row not in rows:
            label=sheet.cell(row,1).value
            if label is not None:
                sections[label]=sections.get(label,0)+1
                rows[row]=('section',str(label),sections[label])
    return [(key,row) for row,key in sorted(rows.items())]


def _copy_cell(source, target, columns, rows):
    value=source.value
    if isinstance(value,str) and value.startswith('='):
        tokens=Tokenizer(value)
        for token in tokens.items:
            if token.type=='OPERAND' and token.subtype=='RANGE':
                if '!' in token.value:
                    raise TemplyfierError('A worksheet-qualified source formula cannot safely be moved into benchmark columns.')
                def replace(match):
                    col=column_index_from_string(match[2]); row=int(match[4])
                    if col not in columns or row not in rows:
                        raise TemplyfierError('A source formula references a row outside the retained benchmark reading.')
                    return match[1]+get_column_letter(columns[col])+match[3]+str(rows[row])
                token.value=re.sub(r'(\$?)([A-Z]{1,3})(\$?)([1-9][0-9]*)',replace,token.value)
        value='='+''.join(token.value for token in tokens.items)
    target.value=value
    target._style=copy(source._style)
    if source.comment: target.comment=copy(source.comment)
    if source.hyperlink: target.hyperlink=copy(source.hyperlink)


def combine_readings(output, records, metric_key):
    grouped={}
    for record in records: grouped.setdefault(record['split'],[]).append(record)
    generated=[]; layouts=[]; missing=[]
    old_names={record['target'].title for record in records}
    used={name for name in output.sheetnames if name not in old_names}
    for split, panels in grouped.items():
        benchmarks=[panel['active_benchmarks'][0] for panel in panels]
        if len(set(benchmarks))!=len(benchmarks):
            raise TemplyfierError('Duplicate readings for the same split and benchmark. Use automatic benchmark detection, or give distinct splits distinct names.')
        panel_plans=[_row_plan(panel,metric_key) for panel in panels]
        order=[]; prototypes={}
        # Ordered union: align by source question and full metric identity, not labels.
        for panel, plan in zip(panels,panel_plans):
            for i,(key,row) in enumerate(plan):
                prototypes.setdefault(key,(panel['target'],row))
                if key not in order:
                    following=next((later for later,_ in plan[i+1:] if later in order),None)
                    order.insert(order.index(following) if following else len(order),key)
        row_numbers={key:index+6 for index,key in enumerate(order)}
        insertion=min(output.index(panel['target']) for panel in panels)
        target=output.create_sheet(_safe_sheet_name(split,used),insertion)
        generated.append(target.title)
        target.freeze_panes='C6'
        first=panels[0]['target']
        for row in range(2,6):
            for col in (1,2):
                _copy_cell(first.cell(row,col),target.cell(row,col),{1:1,2:2},{row:row})
            target.row_dimensions[row].height=first.row_dimensions[row].height
        for col in (1,2):
            target.column_dimensions[get_column_letter(col)].width=first.column_dimensions[get_column_letter(col)].width
        for key,(source,row) in prototypes.items():
            destination=row_numbers[key]
            for col in (1,2):
                _copy_cell(source.cell(row,col),target.cell(destination,col),{1:1,2:2},{row:destination})
            target.row_dimensions[destination].height=source.row_dimensions[row].height
        current=3
        for panel,plan in zip(panels,panel_plans):
            source=panel['target']; width=source.max_column-2
            if current+width-1>16384:
                raise TemplyfierError('This reading exceeds Excel’s column limit. Use separate benchmark sheets.')
            columns={1:1,2:2,**{col:current+col-3 for col in range(3,source.max_column+1)}}
            rows={**{row:row for row in range(1,6)},**{row:row_numbers[key] for key,row in plan}}
            for row in range(2,6):
                for col in range(3,source.max_column+1):
                    _copy_cell(source.cell(row,col),target.cell(row,columns[col]),columns,rows)
            position=panel['active_benchmarks'][0]
            label=panel['benchmark_labels'][position]
            target.merge_cells(start_row=1,start_column=current,end_row=1,end_column=current+width-1)
            heading=target.cell(1,current,f'Comparison vs {label}')
            heading.font=Font(bold=True,color='FFFFFF',size=12)
            heading.fill=PatternFill('solid',fgColor='183D45')
            heading.alignment=Alignment(horizontal='center',vertical='center')
            target.row_dimensions[1].height=30
            present={key for key,_ in plan}
            for key,row in plan:
                if key[0]=='section':continue
                for col in range(3,source.max_column+1):
                    _copy_cell(source.cell(row,col),target.cell(row_numbers[key],columns[col]),columns,rows)
            for key in order:
                if key[0]=='metric' and key not in present:
                    cell=target.cell(row_numbers[key],current)
                    cell.comment=Comment('No matching source metric in this benchmark export. No value or gap has been calculated.','Templyfier')
                    cell.fill=PatternFill('solid',fgColor='F2F2F2')
                    missing.append({'split':split,'benchmark':label,'question':key[1],'metric':key[2]})
            for col in range(3,source.max_column+1):
                target.column_dimensions[get_column_letter(columns[col])].width=source.column_dimensions[get_column_letter(col)].width
            layouts.append({'sheet':target.title,'benchmark':label,'first_column':current,'last_column':current+width-1})
            current+=width+1
        for key in order:
            if key[0]=='section':
                row=row_numbers[key]
                target.merge_cells(start_row=row,start_column=1,end_row=row,end_column=current-2)
                target.cell(row,1).fill=PatternFill('solid',fgColor='E4EEEE')
        target.auto_filter.ref=f'A5:{get_column_letter(current-2)}{max(5,len(order)+5)}'
        target.sheet_view.showGridLines=False
        target.sheet_properties.pageSetUpPr.fitToPage=True
        target.page_setup.orientation='landscape'
        target.page_setup.fitToWidth=1
        target.page_setup.fitToHeight=0
        target.print_title_rows='1:5'
        target.print_title_cols='A:B'
        for panel in panels:output.remove(panel['target'])
    return generated,layouts,missing
